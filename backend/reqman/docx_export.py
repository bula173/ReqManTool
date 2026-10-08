"""Export a ReqMan project into a formatted Word (.docx) spec document.

Unlike ``csv_export.py``/``xlsx_export.py`` (flat tables), this reads as a
document: project title, one heading per module, and within each module the
requirement tree (respecting ``parent_id`` nesting and ``order``) as nested
headings with each requirement's attribute values underneath.

Rich text values of ``text`` attributes (see ``richtext.py``) keep their
formatting: each paragraph is a Word paragraph (the first one shares the
attribute's label paragraph), ``br`` a line break, ``strong``/``em``/``u``/
``s`` bold/italic/underline/strike runs, a span's ``color`` the run color and
its ``font-size`` the run size (px * 0.75 pt). Lists use Word's built-in
``List Bullet``/``List Number`` styles (``List Bullet 2``/``3``... for nested
lists; deeper than three levels, the level-3 style indented further), every
numbered list restarting at 1; further paragraphs of a list item use the
``List Continue`` styles. Plain values are written exactly as before.

Like the other export modules, this has no dependency on ``reqman.models``/
``reqman.service`` on purpose (duck typing) so it stays trivially testable
and side-effect free.
"""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import richtext

MAX_HEADING_LEVEL = 9

# Word's built-in list styles exist for three nesting levels.
MAX_LIST_STYLE_LEVEL = 3
# The indent per list level of those styles (and their hanging indent).
LIST_INDENT = Inches(0.25)

# Elements that follow ``w:shd`` inside ``w:pPr`` (ECMA-376 schema order).
_PPR_AFTER_SHD = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _row_style(requirement: Any, row_styles: Sequence[Any]) -> Optional[Any]:
    """The first rule whose attribute value the requirement has, or None."""
    for style in row_styles:
        if requirement.attributes.get(style.attribute_id) == style.value:
            return style
    return None


def _apply_style(paragraph: Any, style: Optional[Any]) -> None:
    """Color, shade and embolden a paragraph's runs as ``style`` says (None: unchanged)."""
    if style is None:
        return
    if style.background_color:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), style.background_color.lstrip("#").upper())
        paragraph._p.get_or_add_pPr().insert_element_before(shading, *_PPR_AFTER_SHD)
    for run in paragraph.runs:
        # A run's own color (a rich text span) wins over the rule's, as in the UI.
        if style.font_color and run.font.color.rgb is None:
            run.font.color.rgb = RGBColor.from_string(style.font_color.lstrip("#").upper())
        if style.bold:
            run.bold = True


def _add_inline(paragraph: Any, nodes: Sequence[Any], marks: frozenset = frozenset(), span: tuple = (None, None)) -> None:
    """Add rich inline ``nodes`` to ``paragraph`` as formatted runs.

    ``marks`` are the enclosing ``strong``/``em``/``u``/``s`` tags; ``span``
    the innermost enclosing span color and font size.
    """
    for node in nodes:
        if isinstance(node, str):
            run = paragraph.add_run(node)
            if "strong" in marks:
                run.bold = True
            if "em" in marks:
                run.italic = True
            if "u" in marks:
                run.underline = True
            if "s" in marks:
                run.font.strike = True
            color, size = span
            if color:
                run.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
            if size:
                run.font.size = Pt(float(size.removesuffix("px")) * 0.75)
        elif isinstance(node, richtext.LineBreak):
            paragraph.add_run().add_break()
        elif node.tag == "span":
            color, size = richtext.span_style(node.style)
            _add_inline(paragraph, node.children, marks, (color or span[0], size or span[1]))
        else:
            _add_inline(paragraph, node.children, marks | {node.tag}, span)


def _level_style(name: str, level: int) -> str:
    """``name`` for list level 1, ``name 2``/``name 3`` deeper (capped at three)."""
    level = min(level, MAX_LIST_STYLE_LEVEL)
    return name if level == 1 else f"{name} {level}"


def _indent_beyond_styles(paragraph: Any, level: int, hanging: bool) -> None:
    """Indent a paragraph nested deeper than the list styles go."""
    if level > MAX_LIST_STYLE_LEVEL:
        paragraph.paragraph_format.left_indent = LIST_INDENT * level
        if hanging:
            paragraph.paragraph_format.first_line_indent = -LIST_INDENT


def _restarted_numbering(document: Document, style_name: str) -> str:
    """A new numbering instance for ``style_name`` (a numbered list style) starting at 1."""
    numbering = document.part.numbering_part.element
    style_num_id = document.styles[style_name].element.pPr.numPr.numId.val
    num = numbering.add_num(numbering.num_having_numId(style_num_id).abstractNumId.val)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return str(num.numId)


def _add_blocks(document: Document, blocks: Sequence[Any], level: int, out: List[Any]) -> None:
    """Add rich ``blocks`` nested ``level`` lists deep as paragraphs, appending each to ``out``."""
    for block in blocks:
        if isinstance(block, richtext.Paragraph):
            if level == 0:
                paragraph = document.add_paragraph()
            else:
                paragraph = document.add_paragraph(style=_level_style("List Continue", level))
                _indent_beyond_styles(paragraph, level, hanging=False)
            _add_inline(paragraph, block.children)
            out.append(paragraph)
            continue
        item_level = level + 1
        style_name = _level_style("List Bullet" if block.tag == "ul" else "List Number", item_level)
        num_id = _restarted_numbering(document, style_name) if block.tag == "ol" else None
        for item in block.items:
            paragraph = document.add_paragraph(style=style_name)
            if num_id is not None:
                num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_numId().val = int(num_id)
            _indent_beyond_styles(paragraph, item_level, hanging=True)
            _add_inline(paragraph, item[0].children)
            out.append(paragraph)
            _add_blocks(document, item[1:], item_level, out)


def _add_rich_value(document: Document, label: str, value: str) -> List[Any]:
    """The paragraphs of a rich value: its first paragraph follows the bold label."""
    blocks = richtext.blocks(value)
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraphs = [paragraph]
    if blocks and isinstance(blocks[0], richtext.Paragraph):
        _add_inline(paragraph, blocks[0].children)
        blocks = blocks[1:]
    _add_blocks(document, blocks, 0, paragraphs)
    return paragraphs


def _add_requirement(
    document: Document, requirement: Any, project: Any, level: int, row_styles: Sequence[Any]
) -> None:
    style = _row_style(requirement, row_styles)
    _apply_style(document.add_heading(requirement.id, level=min(level, MAX_HEADING_LEVEL)), style)
    for attr in project.attributes:
        value = requirement.attributes.get(attr.id)
        if value is None or value == "":
            continue
        if getattr(attr, "type", None) == "text" and richtext.is_rich(value):
            paragraphs = _add_rich_value(document, attr.name, value)
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{attr.name}: ")
            run.bold = True
            # Without the characters XML 1.0 cannot hold (python-docx raises on them).
            paragraph.add_run(richtext.xml_text(str(value)))
            paragraphs = [paragraph]
        for paragraph in paragraphs:
            _apply_style(paragraph, style)


def export_project_to_docx(
    project: Any,
    modules: Sequence[Tuple[Any, Sequence[Any]]],
    row_styles: Sequence[Any] = (),
) -> bytes:
    """Export a project's modules/requirements to a Word document.

    Args:
        project: object with ``.name: str`` and ``.attributes:
            list[AttributeDefinition]`` where each has ``.id: str``,
            ``.name: str`` and ``.type: str`` (a rich value of a ``"text"``
            attribute is written formatted, see the module docstring).
        modules: list of ``(module, requirements)`` tuples, in the module
            order that should appear in the output. Each ``module`` has
            ``.name: str``. Each item of ``requirements`` has ``.id: str``,
            ``.attributes: dict[str, Any]`` (keyed by AttributeDefinition.id),
            ``.parent_id: str | None``, and ``.order: int``.
        row_styles: rules with ``.attribute_id``, ``.value``,
            ``.font_color``/``.background_color`` (``#rrggbb`` or None) and
            ``.bold``, already resolved against the project (see
            ``ProjectService._effective_row_styles``). A requirement takes
            the first rule whose value it has: its heading and attribute
            paragraphs get that font color (not on a run with its own
            color), background shading and bold.

    Returns:
        The bytes of a .docx document: the project name as the title
        (level 0), one level-1 heading per module, and within each module
        the requirement tree (root items at level 2, each level of nesting
        one level deeper, capped at level 9) — each requirement heading is
        its id, followed by a bold-label paragraph per non-empty attribute
        value.
    """
    document = Document()
    document.add_heading(project.name, level=0)

    for module, requirements in modules:
        document.add_heading(module.name, level=1)

        children_of: Dict[Optional[str], List[Any]] = {}
        for requirement in requirements:
            children_of.setdefault(requirement.parent_id, []).append(requirement)
        for children in children_of.values():
            children.sort(key=lambda r: r.order)

        def walk(parent_id: Optional[str], depth: int) -> None:
            for requirement in children_of.get(parent_id, []):
                _add_requirement(document, requirement, project, depth + 2, row_styles)
                walk(requirement.id, depth + 1)

        walk(None, 0)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
