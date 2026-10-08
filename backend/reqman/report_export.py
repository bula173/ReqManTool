"""Export the reports (traceability matrix, verification report, test
execution report) to CSV, XLSX and DOCX.

The input is the report dictionary ``ProjectService`` computes
(``traceability_matrix``, ``verification_report``, ``execution_report``), so
like the other export modules this has no dependency on ``reqman.models`` /
``reqman.service`` (duck typing) and no side effects.

Unlike the requirements CSV (which must round-trip through the CSV import),
report CSV cells that start with ``=``, ``+``, ``-``, ``@``, a tab or a
carriage return get a leading ``'`` so a spreadsheet never runs them as a
formula. XLSX cells holding text are string cells (``xlsx_export``).
"""

from __future__ import annotations

import csv
import io
from datetime import date
from typing import Any, Iterable, Sequence

from docx import Document
from openpyxl import Workbook

from . import richtext
from .xlsx_export import _force_text_cells

STATUS_LABELS = {
    "verified": "Verified",
    "failed": "Failed",
    "not_run": "Not run",
    "not_covered": "Not covered",
    "passed": "Passed",
    "blocked": "Blocked",
    "not_applicable": "Not applicable",
    "not_planned": "Not planned",
}

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def label(status: str | None) -> str:
    return STATUS_LABELS.get(status or "", status or "")


def _csv_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def _sheet(workbook: Workbook, title: str, rows: Iterable[Sequence[Any]], first: bool = False) -> Any:
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title[:31]
    for row in rows:
        values = ["" if v is None else v for v in row]
        sheet.append(values)
        _force_text_cells(sheet[sheet.max_row], values)
    return sheet


def _xlsx_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_bytes(document: Any) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_text(value: Any) -> str:
    return richtext.xml_text("" if value is None else str(value))


def _docx_table(document: Any, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(_docx_text(text)).bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = _docx_text(value)


def _percent(value: float) -> str:
    return f"{value:.1f} %"


# --- traceability matrix ---------------------------------------------------------


def _trace_texts(report: dict[str, Any]) -> dict[str, str]:
    return {item["id"]: item["text"] for module in report["modules"] for item in module["items"]}


def _module_names(report: dict[str, Any]) -> dict[str, str]:
    return {module["id"]: module["name"] for module in report["modules"]}


def _trace_rows(report: dict[str, Any]) -> list[list[Any]]:
    """The trace table: one row per item of the first module, a column per later module."""
    modules = report["modules"]
    texts = _trace_texts(report)
    first = modules[0]["name"]
    header = [f"{first} Id", f"{first} Text"] + [m["name"] for m in modules[1:]] + ["Covered"]
    chain = len(modules) > 2
    if chain:
        header.append(f"Traced to {modules[-1]['name']}")
    rows = [header]
    for row in report["trace"]:
        values = [row["id"], texts.get(row["id"], "")]
        values += ["; ".join(ids) for ids in row["reached"]]
        values.append("yes" if row["reached"] and row["reached"][0] else "no")
        if chain:
            values.append("yes" if row["complete"] else "no")
        rows.append(values)
    return rows


def traceability_to_csv(report: dict[str, Any]) -> str:
    """The trace table as CSV (see ``_trace_rows``)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for row in _trace_rows(report):
        writer.writerow([_csv_cell(v) for v in row])
    return buffer.getvalue()


def traceability_to_xlsx(report: dict[str, Any], project_name: str) -> bytes:
    """Summary, the trace table, one grid per step and the gap lists."""
    modules = report["modules"]
    names = _module_names(report)
    texts = _trace_texts(report)
    summary = report["summary"]
    workbook = Workbook()
    rows: list[list[Any]] = [
        ["Traceability matrix", project_name],
        ["Chain", " → ".join(m["name"] for m in modules)],
        [],
        ["Step", "Covered", "Of", "%", "Uncovered", "Orphans"],
    ]
    for step in report["steps"]:
        rows.append([
            f"{step['config']['name']} ({names[step['from_module_id']]} → {names[step['to_module_id']]})",
            step["covered"], step["total"], step["covered_percent"], len(step["uncovered"]), len(step["orphans"]),
        ])
    rows += [[], [f"Traced to {modules[-1]['name']}", summary["traced"], summary["total"], summary["traced_percent"]]]
    for entry in report["unverified"]:
        rows.append([
            f"Unverified in {names[entry['module_id']]} (no {names[entry['test_module_id']]} test)",
            len(entry["ids"]), entry["total"],
        ])
    _sheet(workbook, "Summary", rows, first=True)
    _sheet(workbook, "Trace", _trace_rows(report))

    for index, step in enumerate(report["steps"], start=1):
        from_items = modules[index - 1]["items"]
        to_items = modules[index]["items"]
        linked = {(row, column) for row, column in step["links"]}
        grid = [["Id", "Text"] + [item["id"] for item in to_items]]
        for item in from_items:
            grid.append(
                [item["id"], item["text"]]
                + ["x" if (item["id"], column["id"]) in linked else "" for column in to_items]
            )
        sheet = _sheet(workbook, f"Matrix {index}", grid)
        sheet.freeze_panes = "C2"

    gaps: list[list[Any]] = [["Kind", "Step", "Module", "Id", "Text"]]
    for index, step in enumerate(report["steps"], start=1):
        for item_id in step["uncovered"]:
            gaps.append(["Uncovered", index, names[step["from_module_id"]], item_id, texts.get(item_id, "")])
        for item_id in step["orphans"]:
            gaps.append(["Orphan", index, names[step["to_module_id"]], item_id, texts.get(item_id, "")])
    for entry in report["unverified"]:
        for item_id in entry["ids"]:
            gaps.append(["Unverified", "", names[entry["module_id"]], item_id, texts.get(item_id, "")])
    _sheet(workbook, "Gaps", gaps)
    return _xlsx_bytes(workbook)


# --- data source (which commit / baseline a report was computed from) -----------------


def data_source_text(report: dict[str, Any]) -> str | None:
    """One line naming the data a report was computed from; ``None`` when the report has no ``data_source``."""
    source = report.get("data_source")
    if not isinstance(source, dict):
        return None
    commit = source.get("commit")
    if not commit:
        text = "Current data (no git history)"
    else:
        text = f"Current data at commit {commit[:12]}"
        if source.get("baselines"):
            text += f" (baseline {', '.join(source['baselines'])})"
        if source.get("pending_changes"):
            text += f" plus {source['pending_changes']} uncommitted change(s)"
        latest = source.get("latest_baseline")
        if latest and latest.get("name") not in (source.get("baselines") or []):
            text += f"; latest baseline: {latest['name']} ({str(latest.get('commit') or '')[:12]})"
    lock = source.get("locked")
    if isinstance(lock, dict):
        by = f" by {lock['by']}" if lock.get("by") else ""
        text += f"; project locked{by}"
    return text


def _data_source_row(report: dict[str, Any]) -> list[list[Any]]:
    text = data_source_text(report)
    return [["Data", text]] if text else []


def _docx_data_source(document: Any, report: dict[str, Any]) -> None:
    text = data_source_text(report)
    if text:
        document.add_paragraph(_docx_text(f"Data: {text}"))


# --- verification report ------------------------------------------------------------


def _test_summary(test: dict[str, Any]) -> str:
    latest = test.get("latest") or {}
    details = [label(test["result"])]
    if latest.get("execution_date"):
        details.append(latest["execution_date"])
    if latest.get("id"):
        details.append(latest["id"])
    return f"{test['id']}: {', '.join(details)}"


def _verification_summary_rows(report: dict[str, Any]) -> list[list[Any]]:
    totals, percent = report["totals"], report["percent"]
    rows = [[label(s), totals[s], _percent(percent[s])] for s in ("verified", "failed", "not_run", "not_covered")]
    rows.append(["Total", totals["total"], ""])
    return rows


def verification_to_xlsx(report: dict[str, Any], project_name: str) -> bytes:
    workbook = Workbook()
    module = report["module"]
    rows: list[list[Any]] = [
        ["Verification report", project_name],
        ["Module", module["name"]],
        ["Test case modules", ", ".join(m["name"] for m in report["test_modules"]) or "(none)"],
        ["Execution modules", ", ".join(m["name"] for m in report["execution_modules"]) or "(none)"],
        ["Generated", date.today().isoformat()],
        *_data_source_row(report),
        [],
        ["Status", "Requirements", "%"],
    ]
    rows += _verification_summary_rows(report)
    rows.append(["Left out (headings, information...)", report["totals"]["excluded"], ""])
    _sheet(workbook, "Summary", rows, first=True)

    detail: list[list[Any]] = [[
        "Requirement", "Text", "Status", "Test case", "Test result", "Execution", "Execution date",
        "Executed by", "Version", "Result from",
    ]]
    for requirement in report["requirements"]:
        head = [requirement["id"], requirement["text"], label(requirement["status"])]
        if not requirement["tests"]:
            detail.append(head + [""] * 7)
        for test in requirement["tests"]:
            latest = test.get("latest") or {}
            detail.append(head + [
                test["id"], label(test["result"]), latest.get("id"), latest.get("execution_date"),
                latest.get("executed_by"), latest.get("version"),
                {"execution": "execution", "test_case": "test case"}.get(test["source"] or "", ""),
            ])
    _sheet(workbook, "Requirements", detail).freeze_panes = "A2"
    return _xlsx_bytes(workbook)


def verification_to_docx(report: dict[str, Any], project_name: str) -> bytes:
    module = report["module"]
    document = Document()
    document.add_heading(_docx_text(f"Verification report: {module['name']}"), level=0)
    document.add_paragraph(_docx_text(f"Project: {project_name}"))
    document.add_paragraph(f"Generated: {date.today().isoformat()}")
    _docx_data_source(document, report)
    document.add_paragraph(_docx_text(
        "Test case modules: " + (", ".join(m["name"] for m in report["test_modules"]) or "(none)")
    ))
    document.add_paragraph(_docx_text(
        "Execution modules: " + (", ".join(m["name"] for m in report["execution_modules"]) or "(none)")
    ))
    document.add_heading("Summary", level=1)
    _docx_table(document, ["Status", "Requirements", "%"], _verification_summary_rows(report))
    document.add_heading("Requirements", level=1)
    _docx_table(
        document,
        ["Requirement", "Text", "Status", "Test cases (latest result)"],
        (
            [r["id"], r["text"], label(r["status"]), "\n".join(_test_summary(t) for t in r["tests"]) or "-"]
            for r in report["requirements"]
        ),
    )
    return _docx_bytes(document)


# --- test execution report --------------------------------------------------------------


def _execution_summary_rows(report: dict[str, Any]) -> list[list[Any]]:
    counts = report["test_counts"]
    date_range = report["date_range"]
    return [
        ["Build/version under test", ", ".join(report["versions"]) or "(not recorded)"],
        ["Executed by", ", ".join(report["executors"]) or "(not recorded)"],
        ["Execution dates", " to ".join(dict.fromkeys(date_range)) if date_range else "(not recorded)"],
        ["Test cases in scope", counts["total"]],
        ["Planned (with a record)", counts["planned"]],
        ["Executed", f"{counts['executed']} of {counts['planned']} ({_percent(counts['executed_percent'])})"],
        ["Passed", counts["passed"]],
        ["Failed", counts["failed"]],
        ["Blocked", counts["blocked"]],
        ["Not run", counts["not_run"]],
        ["Not planned", counts["not_planned"]],
        ["Pass rate (of executed)", _percent(counts["pass_percent"])],
        ["Records", report["record_counts"]["total"]],
        ["Records linked to no test case", len(report["unlinked_records"])],
    ]


def _affected(problem: dict[str, Any]) -> str:
    return ", ".join(r["id"] for r in problem["requirements"]) or "(none)"


def execution_to_xlsx(report: dict[str, Any], project_name: str) -> bytes:
    workbook = Workbook()
    rows: list[list[Any]] = [
        ["Test execution report", project_name],
        ["Module", report["module"]["name"]],
        ["Test case modules", ", ".join(m["name"] for m in report["test_modules"]) or "(none)"],
        ["Generated", date.today().isoformat()],
        *_data_source_row(report),
        [],
    ]
    rows += _execution_summary_rows(report)
    _sheet(workbook, "Summary", rows, first=True)

    tests: list[list[Any]] = [["Test case", "Text", "Status", "Latest record", "Execution date", "Executed by", "Version", "Evidence"]]
    for test in report["test_cases"]:
        latest = test.get("latest") or {}
        tests.append([
            test["id"], test["text"], label(test["status"]), latest.get("id"), latest.get("execution_date"),
            latest.get("executed_by"), latest.get("version"), latest.get("evidence"),
        ])
    _sheet(workbook, "Test cases", tests).freeze_panes = "A2"

    problems: list[list[Any]] = [["Test case", "Text", "Status", "Affected requirements", "Record", "Evidence"]]
    for problem in report["problems"]:
        latest = problem.get("latest") or {}
        problems.append([
            problem["id"], problem["text"], label(problem["status"]), _affected(problem),
            latest.get("id"), latest.get("evidence"),
        ])
    _sheet(workbook, "Failed and blocked", problems)

    records: list[list[Any]] = [["Record", "Text", "Test cases", "Result", "Execution date", "Executed by", "Version", "Evidence"]]
    for record in report["records"]:
        records.append([
            record["id"], record["text"], ", ".join(record["test_case_ids"]),
            record["result_value"] if record["result_value"] is not None else label(record["result"]),
            record["execution_date"], record["executed_by"], record["version"], record["evidence"],
        ])
    _sheet(workbook, "Records", records).freeze_panes = "A2"
    return _xlsx_bytes(workbook)


def execution_to_docx(report: dict[str, Any], project_name: str) -> bytes:
    document = Document()
    document.add_heading(_docx_text(f"Test execution report: {report['module']['name']}"), level=0)
    document.add_paragraph(_docx_text(f"Project: {project_name}"))
    document.add_paragraph(f"Generated: {date.today().isoformat()}")
    _docx_data_source(document, report)
    document.add_paragraph(_docx_text(
        "Test case modules: " + (", ".join(m["name"] for m in report["test_modules"]) or "(none)")
    ))
    document.add_heading("Summary", level=1)
    _docx_table(document, ["Item", "Value"], _execution_summary_rows(report))
    document.add_heading("Failed and blocked tests", level=1)
    if report["problems"]:
        _docx_table(
            document,
            ["Test case", "Text", "Status", "Affected requirements", "Evidence"],
            (
                [p["id"], p["text"], label(p["status"]), _affected(p), (p.get("latest") or {}).get("evidence")]
                for p in report["problems"]
            ),
        )
    else:
        document.add_paragraph("None.")
    document.add_heading("Test cases", level=1)
    _docx_table(
        document,
        ["Test case", "Text", "Status", "Execution date", "Executed by", "Version"],
        (
            [
                t["id"], t["text"], label(t["status"]),
                (t.get("latest") or {}).get("execution_date"),
                (t.get("latest") or {}).get("executed_by"),
                (t.get("latest") or {}).get("version"),
            ]
            for t in report["test_cases"]
        ),
    )
    return _docx_bytes(document)
