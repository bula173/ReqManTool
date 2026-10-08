"""Rich text in the formatted exports: DOCX runs/lists and ReqIF XHTML (rich text, slice 2)."""

import importlib
import io
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fastapi.testclient import TestClient

from reqman.reqif_export import export_project_to_reqif
from reqman.richtext import RICH_TEXT_PREFIX as P
from reqman.service import ProjectService
from reqman.store import ProjectStore

REQIF = "{http://www.omg.org/spec/ReqIF/20110401/reqif.xsd}"
XHTML = "{http://www.w3.org/1999/xhtml}"


def rich(body: str) -> str:
    return P + body + "</div>"


@pytest.fixture
def service(tmp_path):
    return ProjectService(ProjectStore(tmp_path))


@pytest.fixture
def module(service):
    project = service.create_project("P")
    service.add_attribute_definition(project.id, "note", "Note", "text")
    service.add_attribute_definition(project.id, "title", "Title", "string")
    service.add_attribute_definition(project.id, "kind", "Kind", "enum", choices=["Heading", "Req"])
    return project.id, service.create_module(project.id, "M", "M").id


# --- DOCX ------------------------------------------------------------------------------


def _docx(service, project_id):
    return Document(io.BytesIO(service.export_docx(project_id)))


def _rows(document, requirement_id):
    """The attribute paragraphs under the heading of ``requirement_id``."""
    paragraphs = document.paragraphs
    start = next(i for i, p in enumerate(paragraphs) if p.text == requirement_id)
    rows = []
    for p in paragraphs[start + 1:]:
        if p.style.name.startswith("Heading"):
            break
        rows.append(p)
    return rows


def _run(paragraph, text):
    return next(r for r in paragraph.runs if r.text == text)


def test_docx_inline_formatting_becomes_run_properties(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": rich(
        '<p>a <strong>b</strong> <em>i</em> <u>u</u> <s>s</s> '
        '<span style="color: #e03131; font-size: 20px">c<strong>cb</strong></span> '
        '<span style="font-size: 12px"><span style="color: #1971c2">n</span></span></p>'
    )})
    (row,) = _rows(_docx(service, project_id), req.id)
    assert row.text == "ObjectText: a b i u s ccb n"
    assert _run(row, "ObjectText: ").bold is True
    plain = _run(row, "a ")
    assert (plain.bold, plain.italic, plain.underline, plain.font.strike, plain.font.color.rgb, plain.font.size) == (
        None, None, None, None, None, None
    )
    assert _run(row, "b").bold is True
    assert _run(row, "i").italic is True
    assert _run(row, "u").underline is True
    assert _run(row, "s").font.strike is True
    colored = _run(row, "c")
    assert str(colored.font.color.rgb) == "E03131" and colored.font.size == Pt(15)
    nested = _run(row, "cb")
    assert nested.bold is True and str(nested.font.color.rgb) == "E03131" and nested.font.size == Pt(15)
    inner = _run(row, "n")
    assert str(inner.font.color.rgb) == "1971C2" and inner.font.size == Pt(9)


def test_docx_paragraphs_and_line_breaks(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": rich("<p><em>one</em></p><p>l1<br/>l2</p>")})
    first, second = _rows(_docx(service, project_id), req.id)
    assert first.text == "ObjectText: one"
    assert second.style.name == "Normal" and second.text == "l1\nl2"
    assert len(second._p.findall(".//" + qn("w:br"))) == 1


def test_docx_lists_use_list_styles_and_restart_numbering(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": rich(
        "<ul><li><p>b1</p><p>more</p><ul><li><p>b2</p><ul><li><p>b3</p><ul><li><p>b4</p></li></ul></li></ul></li></ul></li></ul>"
        "<ol><li><p>n1</p></li><li><p>n2</p></li></ol><p>between</p><ol><li><p>m1</p></li></ol>"
    )})
    document = _docx(service, project_id)
    rows = _rows(document, req.id)
    # A value that starts with a list keeps its label on a paragraph of its own.
    assert [(p.style.name, p.text) for p in rows] == [
        ("Normal", "ObjectText: "),
        ("List Bullet", "b1"),
        ("List Continue", "more"),
        ("List Bullet 2", "b2"),
        ("List Bullet 3", "b3"),
        ("List Bullet 3", "b4"),
        ("List Number", "n1"),
        ("List Number", "n2"),
        ("Normal", "between"),
        ("List Number", "m1"),
    ]
    # Deeper than the styles go: the level-3 style, indented one level further.
    assert rows[4].paragraph_format.left_indent is None
    assert rows[5].paragraph_format.left_indent == Inches(1)
    assert rows[5].paragraph_format.first_line_indent == -Inches(0.25)

    # Each numbered list has its own numbering instance, restarting at 1.
    def num_id(p):
        return p._p.pPr.numPr.numId.val

    assert num_id(rows[6]) == num_id(rows[7]) != num_id(rows[9])
    numbering = document.part.numbering_part.element
    for p in (rows[6], rows[9]):
        num = numbering.num_having_numId(num_id(p))
        overrides = num.findall(qn("w:lvlOverride"))
        assert len(overrides) == 1
        assert overrides[0].find(qn("w:startOverride")).get(qn("w:val")) == "1"
    # Bullets keep the style's numbering (no direct numPr).
    assert rows[1]._p.pPr.numPr is None


def test_docx_row_style_applies_to_rich_paragraphs_and_span_color_wins(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {
        "object_text": rich('<p>x <span style="color: #1971c2">blue</span></p><ul><li><p>item</p></li></ul>'),
        "kind": "Req",
    })
    service.set_row_styles(project_id, [
        {"attribute_id": "kind", "value": "Req", "font_color": "#e03131", "background_color": "#fff3bf", "bold": True}
    ])
    rows = _rows(_docx(service, project_id), req.id)
    rich_rows = rows[:2]
    assert [p.text for p in rich_rows] == ["ObjectText: x blue", "item"]
    for p in rich_rows:
        assert p._p.pPr.find(qn("w:shd")).get(qn("w:fill")) == "FFF3BF"
        assert all(r.bold for r in p.runs)
    assert str(_run(rows[0], "x ").font.color.rgb) == "E03131"
    assert str(_run(rows[0], "blue").font.color.rgb) == "1971C2"
    assert str(_run(rows[1], "item").font.color.rgb) == "E03131"


def test_docx_plain_values_are_unchanged(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {
        "object_text": "<p>not</p> rich\nsecond line",
        "title": P + "<p><strong>string attribute</strong></p></div>",
    })
    rows = _rows(_docx(service, project_id), req.id)
    assert [[r.text for r in p.runs] for p in rows] == [
        ["ObjectText: ", "<p>not</p> rich\nsecond line"],
        ["Title: ", P + "<p><strong>string attribute</strong></p></div>"],
    ]
    assert [[bool(r.bold) for r in p.runs] for p in rows] == [[True, False], [True, False]]
    assert all(p.style.name == "Normal" for p in rows)


# --- ReqIF ----------------------------------------------------------------------------


def _reqif(xml: str):
    return ET.fromstring(xml.encode())


def _the_value(root, requirement_id, attribute_id):
    """The ``xhtml:div`` of a requirement's XHTML attribute value, or None."""
    for obj in root.iter(REQIF + "SPEC-OBJECT"):
        if obj.get("IDENTIFIER") != requirement_id:
            continue
        for value in obj.iter(REQIF + "ATTRIBUTE-VALUE-XHTML"):
            if value.find(f"{REQIF}DEFINITION/{REQIF}ATTRIBUTE-DEFINITION-XHTML-REF").text == attribute_id:
                (div,) = list(value.find(REQIF + "THE-VALUE"))
                return div
    return None


def _shape(element):
    """``(tag, style, text, children..., tail)`` of an XHTML element, namespace checked."""
    assert element.tag.startswith(XHTML), element.tag
    return (
        element.tag.removeprefix(XHTML), element.get("style"), element.text,
        [_shape(child) for child in element], element.tail,
    )


def test_reqif_text_attributes_are_xhtml(service, module):
    project_id, module_id = module
    service.add_requirement(project_id, module_id, {"object_text": "plain", "title": "t"})
    root = _reqif(service.export_reqif(project_id))
    datatypes = {e.tag.removeprefix(REQIF): e.get("IDENTIFIER") for e in root.find(f".//{REQIF}DATATYPES")}
    assert datatypes["DATATYPE-DEFINITION-XHTML"] == "DATATYPE-XHTML"
    definitions = {
        e.get("IDENTIFIER"): (e.tag.removeprefix(REQIF), e.find(f"{REQIF}TYPE")[0].text)
        for e in root.find(f".//{REQIF}SPEC-ATTRIBUTES")
    }
    assert definitions["object_text"] == ("ATTRIBUTE-DEFINITION-XHTML", "DATATYPE-XHTML")
    assert definitions["note"] == ("ATTRIBUTE-DEFINITION-XHTML", "DATATYPE-XHTML")
    assert definitions["title"] == ("ATTRIBUTE-DEFINITION-STRING", "DATATYPE-STRING")
    assert definitions["kind"] == ("ATTRIBUTE-DEFINITION-STRING", "DATATYPE-STRING")


def test_reqif_formatting_round_trips_into_xhtml_elements(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": rich(
        '<p>a <strong>b</strong><em>i</em><u>u</u><s>s</s><span style="color: #e03131; font-size: 14px">c</span></p>'
        "<p>l1<br/>l2</p><ul><li><p>x</p><ol><li><p>y</p></li></ol></li></ul>"
    )})
    div = _the_value(_reqif(service.export_reqif(project_id)), req.id, "object_text")
    assert _shape(div) == ("div", None, None, [
        ("p", None, "a ", [
            ("strong", None, "b", [], None),
            ("em", None, "i", [], None),
            ("span", "text-decoration: underline", "u", [], None),
            ("span", "text-decoration: line-through", "s", [], None),
            ("span", "color: #e03131; font-size: 14px", "c", [], None),
        ], None),
        ("p", None, "l1", [("br", None, None, [], "l2")], None),
        ("ul", None, None, [
            ("li", None, None, [
                ("p", None, "x", [], None),
                ("ol", None, None, [("li", None, None, [("p", None, "y", [], None)], None)], None),
            ], None),
        ], None),
    ], None)


def test_reqif_escapes_text_safely(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {
        "object_text": rich("<p><strong>a &lt; b &amp;&amp; c &gt; d</strong> &lt;/xhtml:div&gt;</p>"),
        "note": "x < y & z </THE-VALUE> ]]>",
    })
    xml = service.export_reqif(project_id)
    root = _reqif(xml)
    p = _the_value(root, req.id, "object_text")[0]
    assert p[0].text == "a < b && c > d" and p[0].tail == " </xhtml:div>"
    assert _the_value(root, req.id, "note").text == "x < y & z </THE-VALUE> ]]>"
    assert "x &lt; y &amp; z &lt;/THE-VALUE&gt; ]]&gt;" in xml


def test_reqif_plain_value_is_an_xhtml_div(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": "one\r\ntwo\nthree", "note": "<p>literal</p>"})
    root = _reqif(service.export_reqif(project_id))
    assert _shape(_the_value(root, req.id, "object_text")) == (
        "div", None, "one", [("br", None, None, [], "two"), ("br", None, None, [], "three")], None
    )
    assert _shape(_the_value(root, req.id, "note")) == ("div", None, "<p>literal</p>", [], None)


def test_reqif_string_attribute_with_rich_looking_value_stays_a_string(service, module):
    project_id, module_id = module
    value = rich("<p><strong>x</strong></p>")
    service.add_requirement(project_id, module_id, {"object_text": "o", "title": value})
    root = _reqif(service.export_reqif(project_id))
    (string_value,) = list(root.iter(REQIF + "ATTRIBUTE-VALUE-STRING"))
    assert string_value.get("THE-VALUE") == value


def _stub_project(values):
    attrs = [SimpleNamespace(id="object_text", name="ObjectText", type="text")]
    project = SimpleNamespace(id="p", name="P", description="", attributes=attrs)
    module = SimpleNamespace(id="m", name="M", description="")
    reqs = [SimpleNamespace(id=f"R-{i}", attributes={"object_text": v}, parent_id=None) for i, v in enumerate(values)]
    return project, [(module, reqs)]


def test_reqif_empty_and_blank_values_are_empty_divs():
    project, modules = _stub_project(["", rich("<p> </p>"), "a\x01b\x0c"])
    root = _reqif(export_project_to_reqif(project, modules))
    assert _shape(_the_value(root, "R-0", "object_text")) == ("div", None, None, [], None)
    assert _shape(_the_value(root, "R-1", "object_text")) == ("div", None, None, [], None)
    # Characters XML 1.0 cannot hold are dropped, so the document stays well-formed.
    assert _the_value(root, "R-2", "object_text").text == "ab"


def test_reqif_plain_value_keeps_its_spacing_visible():
    project, modules = _stub_project(["  lead  two   three\tx ", "\t"])
    xml = export_project_to_reqif(project, modules)
    nb = " "
    assert "&#160;&#160;lead &#160;two &#160; three&#160;&#160;&#160;&#160;x&#160;</xhtml:div>" in xml
    root = _reqif(xml)
    assert _the_value(root, "R-0", "object_text").text == f"{nb}{nb}lead {nb}two {nb} three{nb * 4}x{nb}"
    assert _the_value(root, "R-1", "object_text").text == nb * 4


def test_reqif_rich_value_keeps_its_spacing_visible():
    # The editor keeps spaces and tabs as typed (preserveWhitespace: 'full'), so a
    # rich value's spacing is made visible like a plain value's, per line of a
    # paragraph, across inline marks.
    project, modules = _stub_project([
        rich("<p>  lead  two<strong>   three</strong>\tx </p><p>a<br/>  b<em> </em></p>"),
        rich("<ul><li><p>x  y</p></li></ul>"),
    ])
    xml = export_project_to_reqif(project, modules)
    nb = " "
    root = _reqif(xml)
    first, second = list(_the_value(root, "R-0", "object_text"))
    assert _shape(first) == (
        "p", None, f"{nb}{nb}lead {nb}two", [("strong", None, f" {nb} three", [], f"{nb * 4}x{nb}")], None
    )
    assert _shape(second) == (
        "p", None, "a", [("br", None, None, [], f"{nb}{nb}b"), ("em", None, nb, [], None)], None
    )
    item = _the_value(root, "R-1", "object_text")[0][0][0]
    assert item.text == f"x {nb}y"


def test_reqif_drops_lone_surrogates_from_text_values():
    # A hand-edited (or git-merged) value may hold a lone surrogate, which no XML can hold.
    project, modules = _stub_project(["a\udc80b", rich("<p>c\ud800d</p>")])
    xml = export_project_to_reqif(project, modules)
    root = _reqif(xml)  # encodes as UTF-8: raises on a surrogate left in
    assert _the_value(root, "R-0", "object_text").text == "ab"
    assert _the_value(root, "R-1", "object_text")[0].text == "cd"


def test_reqif_string_and_enum_values_drop_characters_xml_cannot_hold():
    attrs = [
        SimpleNamespace(id="title", name="Title", type="string"),
        SimpleNamespace(id="kind", name="Kind", type="enum"),
    ]
    project = SimpleNamespace(id="p", name="P", description="", attributes=attrs)
    module = SimpleNamespace(id="m", name="M", description="")
    req = SimpleNamespace(id="R-1", attributes={"title": "x\x01y\udc80", "kind": "H\x0bead"}, parent_id=None)
    root = _reqif(export_project_to_reqif(project, [(module, [req])]))
    values = sorted(v.get("THE-VALUE") for v in root.iter(REQIF + "ATTRIBUTE-VALUE-STRING"))
    assert values == ["Head", "xy"]


def test_reqif_without_text_attributes_has_no_xhtml_datatype():
    project = SimpleNamespace(id="p", name="P", description="", attributes=[
        SimpleNamespace(id="n", name="N", type="integer")
    ])
    xml = export_project_to_reqif(project, [])
    assert "XHTML" not in xml
    _reqif(xml)


def test_reqif_xhtml_parses_with_the_reqif_library(service, module):
    from reqif.parser import ReqIFParser
    from reqif.models.reqif_types import SpecObjectAttributeType

    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": rich("<p><em>x</em> &amp; y</p>")})
    bundle = ReqIFParser.parse_from_string(service.export_reqif(project_id))
    (obj,) = [o for o in bundle.core_content.req_if_content.spec_objects if o.identifier == req.id]
    (attribute,) = obj.attributes
    assert attribute.attribute_type == SpecObjectAttributeType.XHTML
    assert attribute.value_stripped_xhtml == "<div><p><em>x</em> &amp; y</p></div>"


# --- REST -----------------------------------------------------------------------------


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REQMAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(tmp_path / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app)


def test_export_endpoints_with_rich_values(client):
    project_id = client.post("/projects", json={"name": "P"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "M", "prefix": "M"}).json()["id"]
    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": rich("<p><strong>Bold</strong> &lt;x&gt;</p><ul><li>i</li></ul>")}},
    )
    assert resp.status_code in (200, 201), resp.text
    base = f"/projects/{project_id}/export"
    expected = {
        "reqif": "application/xml",
        "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    bodies = {}
    for kind, media_type in expected.items():
        resp = client.get(f"{base}/{kind}")
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(media_type)
        bodies[kind] = resp.content
    assert b"<xhtml:strong>Bold</xhtml:strong> &lt;x&gt;" in bodies["reqif"]
    assert b"Bold <x>\n- i" in bodies["csv"].replace(b"\r\n", b"\n")
    rows = _rows(Document(io.BytesIO(bodies["docx"])), "M-0001")
    assert [p.text for p in rows] == ["ObjectText: Bold <x>", "i"]
    assert _run(rows[0], "Bold").bold is True


# --- gaps: oversized stored values, deep lists with row styles, identifiers, unicode, blocks() ---


def _hand_edit(service, project_id, module_id, requirement_id, **values):
    """Write attribute values straight to the store, bypassing the sanitiser (as a hand edit would)."""
    from dataclasses import replace

    stored = service.store.load_requirement(project_id, module_id, requirement_id)
    service.store.save_requirement(
        project_id, module_id, replace(stored, attributes={**stored.attributes, **values})
    )


def test_oversized_stored_rich_values_export_truncated_like_to_plain(service, module, monkeypatch):
    from reqman import richtext

    # The real length limit; a lower node limit keeps the many-runs DOCX fast.
    monkeypatch.setattr(richtext, "MAX_NODES", 2_000)
    project_id, module_id = module
    too_long = rich("<p>" + "a" * (richtext.MAX_LENGTH + 50) + "</p>")
    too_many_nodes = rich("<p>" + "<em>b</em>" * (richtext.MAX_NODES // 2 + 10) + "</p>")
    for value in (too_long, too_many_nodes):
        with pytest.raises(richtext.RichTextTooLarge):
            richtext.sanitize(value)
    req = service.add_requirement(project_id, module_id, {"object_text": "x"})
    _hand_edit(service, project_id, module_id, req.id, object_text=too_long, note=too_many_nodes)
    expected_long = richtext.to_plain(too_long)
    expected_nodes = richtext.to_plain(too_many_nodes)
    # Truncated, not empty and not the whole value.
    assert set(expected_long) == {"a"} and 0 < len(expected_long) < richtext.MAX_LENGTH + 50
    assert set(expected_nodes) == {"b"} and 0 < len(expected_nodes) < richtext.MAX_NODES // 2 + 10

    rows = _rows(_docx(service, project_id), req.id)
    texts = {p.text.partition(": ")[0]: p.text.partition(": ")[2] for p in rows}
    assert texts["ObjectText"] == expected_long
    assert texts["Note"] == expected_nodes

    root = _reqif(service.export_reqif(project_id))
    long_div = _the_value(root, req.id, "object_text")
    assert "".join(long_div.itertext()) == expected_long
    nodes_div = _the_value(root, req.id, "note")
    assert "".join(nodes_div.itertext()) == expected_nodes
    assert all(e.tag == XHTML + "em" for e in nodes_div[0])


def _identifiers(root):
    return [e.get("IDENTIFIER") for e in root.iter() if e.get("IDENTIFIER") is not None]


def test_reqif_identifiers_are_stable_and_unique_across_modules(service, module):
    project_id, first_module = module
    second_module = service.create_module(project_id, "Second", "S").id
    a = service.add_requirement(project_id, first_module, {"object_text": rich("<p><strong>a</strong></p>")})
    service.add_requirement(project_id, first_module, {"object_text": "child"}, parent_id=a.id)
    b = service.add_requirement(project_id, second_module, {"object_text": rich("<ul><li><p>b</p></li></ul>")})
    service.add_requirement(project_id, second_module, {"object_text": "c", "note": "n"}, parent_id=b.id)

    first_xml = service.export_reqif(project_id)
    second_xml = service.export_reqif(project_id)
    first, second = _identifiers(_reqif(first_xml)), _identifiers(_reqif(second_xml))
    assert first == second
    assert len(first) == len(set(first)), sorted(i for i in first if first.count(i) > 1)
    root = _reqif(first_xml)
    specs = [e.get("IDENTIFIER") for e in root.iter(REQIF + "SPECIFICATION")]
    assert specs == [first_module, second_module]
    objects = {e.get("IDENTIFIER") for e in root.iter(REQIF + "SPEC-OBJECT")}
    assert len(objects) == 4 and a.id in objects and b.id in objects
    # The XHTML datatype is declared exactly once for both modules' text values.
    assert first.count("DATATYPE-XHTML") == 1
    # Editing one value leaves every identifier as it was.
    service.update_requirement(project_id, second_module, b.id, {"object_text": rich("<p><em>b2</em></p>")})
    assert _identifiers(_reqif(service.export_reqif(project_id))) == first


def test_docx_row_style_with_lists_deeper_than_three_levels(service, module):
    project_id, first_module = module
    styled_module = service.create_module(project_id, "Styled", "S").id
    service.add_requirement(project_id, first_module, {"object_text": "unstyled", "kind": "Heading"})
    req = service.add_requirement(project_id, styled_module, {
        "kind": "Req",
        "object_text": rich(
            "<p>intro</p>"
            "<ol><li><p>n1</p>"
            "<ul><li><p>b2</p>"
            "<ol><li><p>n3</p>"
            "<ul><li><p>b4</p>"
            "<ol><li><p>n5</p><p>cont5</p></li><li><p>n5b</p></li></ol>"
            "</li></ul></li></ol></li></ul></li></ol>"
        ),
    })
    service.set_row_styles(project_id, [
        {"attribute_id": "kind", "value": "Req", "font_color": "#e03131", "background_color": "#fff3bf", "bold": True}
    ])
    data = service.export_docx(project_id)
    # The document opens, and survives a save/reopen round trip.
    document = Document(io.BytesIO(data))
    buffer = io.BytesIO()
    document.save(buffer)
    document = Document(io.BytesIO(buffer.getvalue()))

    rows = _rows(document, req.id)
    rich_rows = rows[:8]
    assert [(p.style.name, p.text) for p in rich_rows] == [
        ("Normal", "ObjectText: intro"),
        ("List Number", "n1"),
        ("List Bullet 2", "b2"),
        ("List Number 3", "n3"),
        ("List Bullet 3", "b4"),
        ("List Number 3", "n5"),
        ("List Continue 3", "cont5"),
        ("List Number 3", "n5b"),
    ]
    indents = [(p.paragraph_format.left_indent, p.paragraph_format.first_line_indent) for p in rich_rows]
    assert indents[:4] == [(None, None)] * 4
    assert indents[4] == (Inches(1), -Inches(0.25))
    assert indents[5] == (Inches(1.25), -Inches(0.25))
    assert indents[6] == (Inches(1.25), None)
    assert indents[7] == (Inches(1.25), -Inches(0.25))

    # Each numbered list restarts; the items of one list share its numbering.
    def num_id(p):
        return p._p.pPr.numPr.numId.val

    assert len({num_id(rows[1]), num_id(rows[3]), num_id(rows[5])}) == 3
    assert num_id(rows[5]) == num_id(rows[7])

    # The row style reaches every rich paragraph, list levels included, in schema order.
    for p in rich_rows:
        ppr = p._p.pPr
        assert ppr.find(qn("w:shd")).get(qn("w:fill")) == "FFF3BF", p.text
        assert all(r.bold for r in p.runs), p.text
        assert all(str(r.font.color.rgb) == "E03131" for r in p.runs), p.text
        order = [child.tag for child in ppr]
        if ppr.find(qn("w:ind")) is not None:
            assert order.index(qn("w:shd")) < order.index(qn("w:ind")), p.text
        if ppr.find(qn("w:numPr")) is not None:
            assert order.index(qn("w:numPr")) < order.index(qn("w:shd")), p.text
    # The other module's requirement does not match the rule.
    (unstyled,) = [p for p in document.paragraphs if p.text == "ObjectText: unstyled"]
    assert unstyled._p.pPr is None or unstyled._p.pPr.find(qn("w:shd")) is None


def test_reqif_unicode_emoji_and_span_styles_survive(service, module):
    from reqif.parser import ReqIFParser

    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {
        "object_text": rich(
            '<p>Zażółć 😀 <span style="font-size: 24px">big ✓</span> '
            '<strong><span style="color: rgb(47, 158, 68)">grün 🚀</span></strong></p>'
            '<p><span style="color: #ABC; font-size: 10px">日本語 👩‍💻</span></p>'
        ),
        "note": "Plain 🎉\nzweite Zeile ñ",
    })
    xml = service.export_reqif(project_id)
    assert "😀" in xml and "👩‍💻" in xml and "日本語" in xml
    root = _reqif(xml)
    assert _shape(_the_value(root, req.id, "object_text")) == ("div", None, None, [
        ("p", None, "Zażółć 😀 ", [
            ("span", "font-size: 24px", "big ✓", [], " "),
            ("strong", None, None, [("span", "color: #2f9e44", "grün 🚀", [], None)], None),
        ], None),
        ("p", None, None, [("span", "color: #aabbcc; font-size: 10px", "日本語 👩‍💻", [], None)], None),
    ], None)
    assert _shape(_the_value(root, req.id, "note")) == (
        "div", None, "Plain 🎉", [("br", None, None, [], "zweite Zeile ñ")], None
    )
    # The reqif library reads it back with the characters intact.
    bundle = ReqIFParser.parse_from_string(xml)
    (obj,) = [o for o in bundle.core_content.req_if_content.spec_objects if o.identifier == req.id]
    values = {a.definition_ref: a.value_stripped_xhtml for a in obj.attributes}
    assert "grün 🚀" in values["object_text"] and "👩‍💻" in values["object_text"]
    assert "Plain 🎉" in values["note"]


# --- blocks() on values that are not rich text -------------------------------------------


def test_blocks_of_empty_text_is_empty():
    from reqman import richtext

    assert richtext.blocks("") == []
    assert richtext.blocks(rich("")) == []
    assert richtext.blocks(rich("<p> </p><ul><li> </li></ul>")) == []


def test_blocks_of_a_rich_value_matches_its_plain_projection():
    from reqman import richtext

    value = rich("<p>a <strong>b</strong></p><ol><li><p>x</p></li></ol>")
    (para, lst) = richtext.blocks(value)
    assert isinstance(para, richtext.Paragraph) and isinstance(lst, richtext.ListBlock)
    assert para.children[0] == "a " and isinstance(para.children[1], richtext.Mark)
    assert lst.tag == "ol" and len(lst.items) == 1


@pytest.mark.parametrize("value", [
    "short plain",
    "line one\nline two",
    "A plain value longer than the rich prefix, which mentions <b>markup</b> literally",
    "<p><strong>markup-like</strong> plain</p>",
    " " + P + "<p><em>leading space: not rich</em></p></div>",
    P[:-1],  # the prefix without its closing ">"
])
def test_blocks_of_plain_text_is_empty(value):
    """Plain text is never markup (module docstring): blocks() reads none of it as formatting."""
    from reqman import richtext

    assert not richtext.is_rich(value)
    assert richtext.blocks(value) == []


@pytest.mark.parametrize("value", [42, 0, True, False, None, 1.5, [], ["<p>x</p>"], {"a": 1}, b"<p>x</p>"])
def test_blocks_of_a_non_text_value_never_raises(value):
    """Like to_plain, blocks() is documented to never raise on stored data."""
    from reqman import richtext

    assert richtext.blocks(value) == []


def test_blocks_of_an_over_limit_stored_value_reads_within_the_limits(monkeypatch):
    from reqman import richtext

    monkeypatch.setattr(richtext, "MAX_NODES", 100)
    many = rich("<p>" + "<em>b</em>" * 500 + "</p>")
    (para,) = richtext.blocks(many)  # no RichTextTooLarge
    assert 0 < len(para.children) < 500
    assert all(isinstance(n, richtext.Mark) and n.tag == "em" for n in para.children)
    monkeypatch.setattr(richtext, "MAX_LENGTH", 1_000)
    (para,) = richtext.blocks(rich("<p>" + "a" * 5_000 + "</p>"))
    assert set(para.children[0]) == {"a"} and len(para.children[0]) < 1_000


def test_blocks_returns_the_canonical_tree_of_unsanitised_markup():
    from reqman import richtext

    (para, lst) = richtext.blocks(rich('<p><b onclick="x">B</b><script>s</script><br><i>i</i></p><ul><li>a<ul><li>b</ul></ul>'))
    assert [type(n).__name__ for n in para.children] == ["_Mark", "_Br", "_Mark"]
    assert [(m.tag, m.children) for m in (para.children[0], para.children[2])] == [("strong", ["B"]), ("em", ["i"])]
    assert lst.tag == "ul" and len(lst.items) == 1
    first, nested = lst.items[0]
    assert first.children == ["a"] and nested.tag == "ul" and nested.items[0][0].children == ["b"]


# --- span_style / xml_text ---------------------------------------------------------------


@pytest.mark.parametrize("style, expected", [
    ("color: #e03131; font-size: 14px", ("#e03131", "14px")),
    ("font-size: 20px; color: #1971C2", ("#1971c2", "20px")),
    ("color: #abc", ("#aabbcc", None)),
    ("color: rgb(0, 128, 255)", ("#0080ff", None)),
    ("font-size: 10px", (None, "10px")),
    ("font-size: 13px; color: red", (None, None)),  # not an offered size, not a hex/rgb colour
    ("color: rgb(256, 0, 0); position: fixed", (None, None)),
    ("", (None, None)),
    ("garbage;;: ;color", (None, None)),
])
def test_span_style_returns_canonical_color_and_size(style, expected):
    from reqman import richtext

    assert richtext.span_style(style) == expected


def test_xml_text_drops_only_characters_xml_cannot_hold():
    from reqman import richtext

    kept = "a\tb\nc\rd é 😀 \u00a0 \ud7ff \ue000 \ufffd"
    assert richtext.xml_text(kept) == kept
    assert richtext.xml_text("x\x00\x01\x08\x0b\x0c\x0e\x1fy") == "xy"
    assert richtext.xml_text("s\ud800\udc80\udfffe") == "se"
    assert richtext.xml_text("n\ufffe\uffffm") == "nm"
    assert richtext.xml_text("") == ""


# --- DOCX: plain values with control characters, hand-edited markup -----------------------


def test_docx_export_drops_control_characters_from_plain_values(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {
        "object_text": "a\x01b\x0cc\ttab",
        "title": "t\x0bitle\x1f",
        "kind": "Req",
    })
    rows = _rows(_docx(service, project_id), req.id)
    texts = [p.text for p in rows]
    assert "ObjectText: abc\ttab" in texts
    assert "Title: title" in texts
    assert "Kind: Req" in texts


def test_docx_export_of_a_hand_edited_rich_value_is_sanitised(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": "x"})
    _hand_edit(service, project_id, module_id, req.id, object_text=rich(
        '<p>a\x01<b onclick="x">b</b><script>alert(1)</script></p><h1>head</h1>'
    ))
    (row, second) = _rows(_docx(service, project_id), req.id)
    assert row.text == "ObjectText: ab" and _run(row, "b").bold is True
    assert second.text == "head" and second.style.name == "Normal"


def test_docx_value_starting_with_a_list_leaves_the_label_paragraph_alone(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": rich("<ol><li><p><em>first</em></p></li></ol>")})
    label, item = _rows(_docx(service, project_id), req.id)
    assert [r.text for r in label.runs] == ["ObjectText: "] and label.runs[0].bold is True
    assert item.style.name == "List Number" and _run(item, "first").italic is True


# --- service: what the formatted and table exports receive --------------------------------


def test_export_modules_keeps_rich_values_only_for_formatted_exports(service, module):
    from reqman import richtext

    project_id, module_id = module
    formatted = rich("<p><strong>B</strong></p>")
    blank = rich("<p> </p><ul><li><em> </em></li></ul>")
    req = service.add_requirement(project_id, module_id, {"object_text": formatted})
    _hand_edit(service, project_id, module_id, req.id, note=blank)
    project = service.store.load_project(project_id)

    def values(keep_rich):
        ((_, (exported,)),) = service._export_modules(project, keep_rich=keep_rich)
        return exported.attributes

    kept = values(True)
    assert kept["object_text"] == formatted
    # A rich value with no visible text is exported as empty, not as markup
    # nor as its projection (which would read "- " for the blank list item).
    assert richtext.to_plain(blank) != "" and kept["note"] == ""
    table = values(False)
    assert table["object_text"] == "B" and table["note"] == ""
    # The stored requirement is not touched.
    fresh = ProjectService(ProjectStore(service.store.root))
    stored = fresh.get_requirement(project_id, module_id, req.id).attributes
    assert stored["object_text"] == formatted and stored["note"] == blank


def test_blank_rich_list_value_exports_as_empty_everywhere(service, module):
    # A hand-edited rich value with an empty list item has the plain projection "- ".
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": "o"})
    _hand_edit(service, project_id, module_id, req.id, note=rich("<ul><li> </li></ul>"))
    assert "- " not in service.export_csv(project_id)
    note = _the_value(_reqif(service.export_reqif(project_id)), req.id, "note")
    assert note is None or _shape(note) == ("div", None, None, [], None)
    texts = [p.text for p in _rows(_docx(service, project_id), req.id)]
    assert not any(t.startswith("Note:") and t.strip() != "Note:" for t in texts)


def test_rich_looking_value_of_a_non_text_attribute_stays_literal_in_every_export(service, module):
    from openpyxl import load_workbook

    project_id, module_id = module
    service.add_attribute_definition(project_id, "choice", "Choice", "enum", choices=["A", rich("<p><em>e</em></p>")])
    value = rich("<p><strong>x</strong> &amp; y</p>")
    enum_value = rich("<p><em>e</em></p>")
    req = service.add_requirement(project_id, module_id, {"object_text": "o", "title": value, "choice": enum_value})
    assert service.get_requirement(project_id, module_id, req.id).attributes["title"] == value

    csv_text = service.export_csv(project_id)
    assert '<strong>x</strong> &amp; y' in csv_text and "<em>e</em>" in csv_text
    cells = [c for row in load_workbook(io.BytesIO(service.export_xlsx(project_id))).active.iter_rows(values_only=True) for c in row]
    assert value in cells and enum_value in cells
    texts = [p.text for p in _rows(_docx(service, project_id), req.id)]
    assert f"Title: {value}" in texts and f"Choice: {enum_value}" in texts
    root = _reqif(service.export_reqif(project_id))
    strings = {v.find(f"{REQIF}DEFINITION")[0].text: v.get("THE-VALUE") for v in root.iter(REQIF + "ATTRIBUTE-VALUE-STRING")}
    assert strings == {"title": value, "choice": enum_value}


# --- ReqIF: more value shapes ------------------------------------------------------------


def test_reqif_value_shapes_and_none_values():
    project, modules = _stub_project([
        None,
        "  \n\n",
        "a < b & c > d\r\ne",
        rich("<p>1 &lt; 2</p><ol><li><p>a<br/>b</p></li></ol>"),
    ])
    xml = export_project_to_reqif(project, modules)
    root = _reqif(xml)
    # A None value is omitted, not written as an empty div.
    assert _the_value(root, "R-0", "object_text") is None
    # Blank lines keep their spacing visible and their line breaks.
    assert _shape(_the_value(root, "R-1", "object_text")) == (
        "div", None, "\u00a0\u00a0", [("br", None, None, [], None), ("br", None, None, [], None)], None
    )
    assert _shape(_the_value(root, "R-2", "object_text")) == ("div", None, "a < b & c > d", [("br", None, None, [], "e")], None)
    assert "a &lt; b &amp; c &gt; d<xhtml:br/>e" in xml
    assert _shape(_the_value(root, "R-3", "object_text")) == ("div", None, None, [
        ("p", None, "1 < 2", [], None),
        ("ol", None, None, [("li", None, None, [("p", None, "a", [("br", None, None, [], "b")], None)], None)], None),
    ], None)


def test_reqif_round_trips_through_the_reqif_parser_with_every_value_shape(service, module):
    from reqif.parser import ReqIFParser
    from reqif.unparser import ReqIFUnparser

    project_id, module_id = module
    service.add_attribute_definition(project_id, "n", "N", "integer")
    service.add_attribute_definition(project_id, "flag", "Flag", "boolean")
    req = service.add_requirement(project_id, module_id, {
        "object_text": rich('<p><u>u</u> <s>s</s> <span style="color: #e03131">c</span></p><ul><li><p>i</p></li></ul>'),
        "note": "  two  spaces\tand <tags> & \x01ctl",
        "title": "Ti\x02tle",
        "kind": "Req",
        "n": 7,
        "flag": True,
    })
    xml = service.export_reqif(project_id)
    bundle = ReqIFParser.parse_from_string(xml)
    (obj,) = [o for o in bundle.core_content.req_if_content.spec_objects if o.identifier == req.id]
    values = {a.definition_ref: a for a in obj.attributes}
    assert values["object_text"].value_stripped_xhtml == (
        '<div><p><span style="text-decoration: underline">u</span> '
        '<span style="text-decoration: line-through">s</span> '
        '<span style="color: #e03131">c</span></p><ul><li><p>i</p></li></ul></div>'
    )
    assert "&lt;tags&gt; &amp; ctl" in values["note"].value_stripped_xhtml
    assert values["title"].value == "Title"
    assert values["kind"].value == "Req"
    # The library writes the parsed bundle back out without error.
    assert req.id in ReqIFUnparser.unparse(bundle)
