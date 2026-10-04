import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest

from reqman.service import ProjectService
from reqman.store import ProjectStore


@pytest.fixture
def service(tmp_path):
    return ProjectService(ProjectStore(tmp_path))


def create_project(service, name="Project One"):
    return service.create_project(name)


def create_module(service, project_id, name="Module One", prefix="REQ"):
    return service.create_module(project_id, name, prefix)


def test_create_project_seeds_builtin_attributes(service):
    project = create_project(service)
    ids = {a.id for a in project.attributes}
    assert ids == {"object_text"}


def test_create_project_generates_slug_id(service):
    project = create_project(service, "Project One")
    assert project.id == "project-one"


def test_create_project_duplicate_name_gets_suffixed_id(service):
    project1 = create_project(service, "Project One")
    project2 = create_project(service, "Project One")
    assert project1.id == "project-one"
    assert project2.id == "project-one-2"


def test_add_attribute_definition(service):
    project = create_project(service)
    project = service.add_attribute_definition(project.id, "priority", "Priority", type="enum", choices=["low", "high"])
    assert any(a.id == "priority" for a in project.attributes)


def test_create_module_generates_slug_id_and_keeps_prefix(service):
    project = create_project(service)
    module = create_module(service, project.id, "Module One", "REQ")
    assert module.id == "module-one"
    assert module.prefix == "REQ"


def test_create_module_invalid_prefix_raises(service):
    project = create_project(service)
    with pytest.raises(ValueError):
        service.create_module(project.id, "Module One", "1REQ")


def test_add_requirement_requires_title(service):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(ValueError):
        service.add_requirement(project.id, module.id, {})


def test_add_requirement_allocates_prefixed_id(service):
    project = create_project(service)
    module = create_module(service, project.id, prefix="REQ")
    req1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    req2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "REQ-0002"


def test_add_requirement_shares_counter_across_modules_with_same_prefix(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "REQ")
    req1 = service.add_requirement(project.id, mod1.id, {"object_text": "First"})
    req2 = service.add_requirement(project.id, mod2.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "REQ-0002"


def test_add_requirement_different_prefixes_do_not_share_counter(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    req1 = service.add_requirement(project.id, mod1.id, {"object_text": "First"})
    req2 = service.add_requirement(project.id, mod2.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "TEST-0001"


def test_get_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First requirement"})
    req = service.get_requirement(project.id, module.id, created.id)
    assert req.attributes["object_text"] == "First requirement"


def test_update_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First"})
    updated = service.update_requirement(project.id, module.id, created.id, {"object_text": "Updated"})
    assert updated.attributes["object_text"] == "Updated"


def test_update_requirement_unknown_attribute_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First"})
    with pytest.raises(ValueError):
        service.update_requirement(project.id, module.id, created.id, {"nope": "x"})


def test_remove_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.remove_requirement(project.id, module.id, created.id)
    with pytest.raises(FileNotFoundError):
        service.get_requirement(project.id, module.id, created.id)


def test_create_link_configuration_adds_system_attribute(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    assert config.id == "verifies"

    project = service.get_project(project.id)
    attr = next(a for a in project.attributes if a.id == "verifies")
    assert attr.system is True
    assert attr.type == "text"


def test_link_requirements_across_modules(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})

    result = service.link_requirements(project.id, config.id, source.id, target.id)
    assert result["source"].attributes["verifies"] == f"> {target.id}"
    assert result["target"].attributes["verifies"] == f"< {source.id}"


def test_link_requirements_wrong_module_raises(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    mod3 = create_module(service, project.id, "Other Module", "OTH")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    other = service.add_requirement(project.id, mod3.id, {"object_text": "Other"})

    with pytest.raises(ValueError):
        service.link_requirements(project.id, config.id, source.id, other.id)


def test_link_requirements_wrong_direction_raises(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    a = service.add_requirement(project.id, mod1.id, {"object_text": "A"})
    b = service.add_requirement(project.id, mod2.id, {"object_text": "B"})

    with pytest.raises(ValueError):
        service.link_requirements(project.id, config.id, b.id, a.id)


def test_link_requirements_without_configuration_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    a = service.add_requirement(project.id, module.id, {"object_text": "A"})
    b = service.add_requirement(project.id, module.id, {"object_text": "B"})

    with pytest.raises(FileNotFoundError):
        service.link_requirements(project.id, "does-not-exist", a.id, b.id)


def test_unlink_requirements_clears_both_mirrors(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    result = service.unlink_requirements(project.id, config.id, source.id, target.id)
    assert "verifies" not in result["source"].attributes
    assert "verifies" not in result["target"].attributes


def test_remove_requirement_cleans_up_links_and_other_side_mirror(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    service.remove_requirement(project.id, mod1.id, source.id)
    target = service.get_requirement(project.id, mod2.id, target.id)
    assert "verifies" not in target.attributes
    assert service.store.load_link_pairs(project.id, config.id) == []


def test_remove_link_configuration_cleans_up_everything(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    service.remove_link_configuration(project.id, config.id)

    project_after = service.get_project(project.id)
    assert config.id not in {c.id for c in project_after.link_configurations}
    assert config.id not in {a.id for a in project_after.attributes}
    source = service.get_requirement(project.id, mod1.id, source.id)
    target = service.get_requirement(project.id, mod2.id, target.id)
    assert config.id not in source.attributes
    assert config.id not in target.attributes


def test_setting_system_attribute_directly_raises(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    with pytest.raises(ValueError):
        service.add_requirement(
            project.id, mod1.id, {"object_text": "X", config.id: "> REQ-9999"}
        )


def test_purge_module(service):
    project = create_project(service)
    module = create_module(service, project.id)
    service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.add_requirement(project.id, module.id, {"object_text": "Second"})
    purged = service.purge_module(project.id, module.id)
    assert purged == 2
    assert service.list_requirements(project.id, module.id) == []


def test_purge_project(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    service.add_requirement(project.id, mod1.id, {"object_text": "First"})
    service.add_requirement(project.id, mod2.id, {"object_text": "Second"})
    purged = service.purge_project(project.id)
    assert purged == 2


def test_purge_does_not_reuse_ids(service):
    project = create_project(service)
    module = create_module(service, project.id)
    req1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.purge_module(project.id, module.id)
    req2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "REQ-0002"


def test_add_requirement_with_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Chapter 1"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    assert child.parent_id == parent.id


def test_add_requirement_with_missing_parent_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(FileNotFoundError):
        service.add_requirement(
            project.id, module.id, {"object_text": "Child"}, parent_id="REQ-9999"
        )


def test_move_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    chapter1 = service.add_requirement(project.id, module.id, {"object_text": "Chapter 1"})
    chapter2 = service.add_requirement(project.id, module.id, {"object_text": "Chapter 2"})
    req = service.add_requirement(
        project.id, module.id, {"object_text": "Req"}, parent_id=chapter1.id
    )
    moved = service.move_requirement(project.id, module.id, req.id, chapter2.id)
    assert moved.parent_id == chapter2.id

    moved_to_root = service.move_requirement(project.id, module.id, req.id, None)
    assert moved_to_root.parent_id is None


def test_move_requirement_rejects_self_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    req = service.add_requirement(project.id, module.id, {"object_text": "Req"})
    with pytest.raises(ValueError):
        service.move_requirement(project.id, module.id, req.id, req.id)


def test_move_requirement_rejects_descendant_cycle(service):
    project = create_project(service)
    module = create_module(service, project.id)
    grandparent = service.add_requirement(project.id, module.id, {"object_text": "GP"})
    parent = service.add_requirement(
        project.id, module.id, {"object_text": "Parent"}, parent_id=grandparent.id
    )
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    with pytest.raises(ValueError):
        service.move_requirement(project.id, module.id, grandparent.id, child.id)


def test_move_requirement_missing_parent_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    req = service.add_requirement(project.id, module.id, {"object_text": "Req"})
    with pytest.raises(FileNotFoundError):
        service.move_requirement(project.id, module.id, req.id, "REQ-9999")


def test_remove_requirement_promotes_children(service):
    project = create_project(service)
    module = create_module(service, project.id)
    grandparent = service.add_requirement(project.id, module.id, {"object_text": "GP"})
    parent = service.add_requirement(
        project.id, module.id, {"object_text": "Parent"}, parent_id=grandparent.id
    )
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    service.remove_requirement(project.id, module.id, parent.id)
    child = service.get_requirement(project.id, module.id, child.id)
    assert child.parent_id == grandparent.id


def test_remove_root_requirement_promotes_children_to_root(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    service.remove_requirement(project.id, module.id, parent.id)
    child = service.get_requirement(project.id, module.id, child.id)
    assert child.parent_id is None


def test_export_reqif_nests_hierarchy_by_parent_id(service):
    from reqif.parser import ReqIFParser

    project = create_project(service)
    module = create_module(service, project.id)
    chapter = service.add_requirement(project.id, module.id, {"object_text": "Chapter 1"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=chapter.id
    )

    xml = service.export_reqif(project.id)
    bundle = ReqIFParser.parse_from_string(xml)
    specification = bundle.core_content.req_if_content.specifications[0]

    assert len(specification.children) == 1
    root_node = specification.children[0]
    assert root_node.spec_object == chapter.id
    assert root_node.children is not None
    assert len(root_node.children) == 1
    assert root_node.children[0].spec_object == child.id


def test_export_reqif_includes_link_spec_relations(service):
    from reqif.parser import ReqIFParser

    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    xml = service.export_reqif(project.id)
    bundle = ReqIFParser.parse_from_string(xml)
    relations = bundle.core_content.req_if_content.spec_relations
    assert relations is not None
    assert len(relations) == 1
    assert relations[0].source == source.id
    assert relations[0].target == target.id

    relation_types = [
        st for st in bundle.core_content.req_if_content.spec_types
        if getattr(st, "long_name", None) == "Verifies"
    ]
    assert len(relation_types) == 1


def test_add_requirement_assigns_sequential_order(service):
    project = create_project(service)
    module = create_module(service, project.id)
    r1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    r2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    r3 = service.add_requirement(project.id, module.id, {"object_text": "Third"})
    assert (r1.order, r2.order, r3.order) == (0, 1, 2)


def test_add_requirement_order_is_per_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child1 = service.add_requirement(project.id, module.id, {"object_text": "Child 1"}, parent_id=parent.id)
    child2 = service.add_requirement(project.id, module.id, {"object_text": "Child 2"}, parent_id=parent.id)
    assert parent.order == 0
    assert (child1.order, child2.order) == (0, 1)


def test_reorder_requirement_swaps_adjacent_siblings(service):
    project = create_project(service)
    module = create_module(service, project.id)
    r1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    r2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    r3 = service.add_requirement(project.id, module.id, {"object_text": "Third"})

    service.reorder_requirement(project.id, module.id, r2.id, "up")
    r1 = service.get_requirement(project.id, module.id, r1.id)
    r2 = service.get_requirement(project.id, module.id, r2.id)
    r3 = service.get_requirement(project.id, module.id, r3.id)
    assert (r2.order, r1.order, r3.order) == (0, 1, 2)


def test_reorder_requirement_noop_at_boundary(service):
    project = create_project(service)
    module = create_module(service, project.id)
    r1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    r2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})

    service.reorder_requirement(project.id, module.id, r1.id, "up")
    r1 = service.get_requirement(project.id, module.id, r1.id)
    r2 = service.get_requirement(project.id, module.id, r2.id)
    assert (r1.order, r2.order) == (0, 1)

    service.reorder_requirement(project.id, module.id, r2.id, "down")
    r1 = service.get_requirement(project.id, module.id, r1.id)
    r2 = service.get_requirement(project.id, module.id, r2.id)
    assert (r1.order, r2.order) == (0, 1)


def test_move_requirement_appends_at_end_of_new_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent1 = service.add_requirement(project.id, module.id, {"object_text": "Parent 1"})
    parent2 = service.add_requirement(project.id, module.id, {"object_text": "Parent 2"})
    service.add_requirement(project.id, module.id, {"object_text": "Existing child"}, parent_id=parent2.id)
    moved = service.add_requirement(project.id, module.id, {"object_text": "Moved"}, parent_id=parent1.id)

    moved = service.move_requirement(project.id, module.id, moved.id, parent2.id)
    assert moved.order == 1


def test_copy_requirement_duplicates_subtree_with_new_ids(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child = service.add_requirement(project.id, module.id, {"object_text": "Child"}, parent_id=parent.id)

    copy = service.copy_requirement(project.id, module.id, parent.id, None)
    assert copy.id != parent.id
    assert copy.attributes["object_text"] == "Parent"

    all_reqs = service.list_requirements(project.id, module.id)
    copied_children = [r for r in all_reqs if r.parent_id == copy.id]
    assert len(copied_children) == 1
    assert copied_children[0].id != child.id
    assert copied_children[0].attributes["object_text"] == "Child"


def test_copy_requirement_excludes_system_attributes(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    copy = service.copy_requirement(project.id, mod1.id, source.id, None)
    assert config.id not in copy.attributes


def test_copy_requirement_rejects_copy_under_own_descendant(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child = service.add_requirement(project.id, module.id, {"object_text": "Child"}, parent_id=parent.id)

    with pytest.raises(ValueError):
        service.copy_requirement(project.id, module.id, parent.id, child.id)


def test_export_csv_has_header_and_rows(service):
    import csv as csv_module
    import io

    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    parent = service.add_requirement(project.id, mod1.id, {"object_text": "Parent"})
    child = service.add_requirement(
        project.id, mod1.id, {"object_text": "Child"}, parent_id=parent.id
    )
    other = service.add_requirement(project.id, mod2.id, {"object_text": "Other"})

    csv_text = service.export_csv(project.id)
    rows = list(csv_module.reader(io.StringIO(csv_text)))

    assert rows[0] == ["Id", "Module", "Parent", "ObjectText"]
    assert rows[1] == [parent.id, mod1.id, "", "Parent"]
    assert rows[2] == [child.id, mod1.id, parent.id, "Child"]
    assert rows[3] == [other.id, mod2.id, "", "Other"]


def test_export_csv_includes_custom_attributes(service):
    import csv as csv_module
    import io

    project = create_project(service)
    module = create_module(service, project.id)
    service.add_attribute_definition(project.id, "priority", "Priority", type="string")
    req = service.add_requirement(
        project.id, module.id, {"object_text": "Req", "priority": "High"}
    )

    csv_text = service.export_csv(project.id)
    rows = list(csv_module.reader(io.StringIO(csv_text)))

    assert rows[0] == ["Id", "Module", "Parent", "ObjectText", "Priority"]
    assert rows[1] == [req.id, module.id, "", "Req", "High"]


def test_create_view(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="string")
    view = service.create_view(project.id, "Compact", ["object_text"])
    assert view.id == "compact"
    assert view.attribute_ids == ["object_text"]

    views = service.list_views(project.id)
    assert len(views) == 1
    assert views[0].id == view.id


def test_create_view_rejects_unknown_attribute(service):
    project = create_project(service)
    with pytest.raises(ValueError):
        service.create_view(project.id, "Compact", ["does-not-exist"])


def test_remove_view(service):
    project = create_project(service)
    view = service.create_view(project.id, "Compact", ["object_text"])
    service.remove_view(project.id, view.id)
    assert service.list_views(project.id) == []


def test_remove_view_missing_raises(service):
    project = create_project(service)
    with pytest.raises(FileNotFoundError):
        service.remove_view(project.id, "does-not-exist")


def test_find_requirement_locates_across_modules(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    req = service.add_requirement(project.id, mod2.id, {"object_text": "Somewhere"})

    module_id, found = service.find_requirement(project.id, req.id)
    assert module_id == mod2.id
    assert found.id == req.id


def test_find_requirement_missing_raises(service):
    project = create_project(service)
    with pytest.raises(FileNotFoundError):
        service.find_requirement(project.id, "REQ-9999")


def test_export_xlsx_has_header_and_rows(service):
    import io

    from openpyxl import load_workbook

    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    parent = service.add_requirement(project.id, mod1.id, {"object_text": "Parent"})
    child = service.add_requirement(
        project.id, mod1.id, {"object_text": "Child"}, parent_id=parent.id
    )
    other = service.add_requirement(project.id, mod2.id, {"object_text": "Other"})

    xlsx_bytes = service.export_xlsx(project.id)
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == ("Id", "Module", "Parent", "ObjectText")
    assert rows[1][:2] == (parent.id, mod1.id)
    assert not rows[1][2]  # empty parent, written as ""
    assert rows[1][3] == "Parent"
    assert rows[2] == (child.id, mod1.id, parent.id, "Child")
    assert rows[3][:2] == (other.id, mod2.id)
    assert not rows[3][2]
    assert rows[3][3] == "Other"


def test_export_xlsx_includes_custom_attributes(service):
    import io

    from openpyxl import load_workbook

    project = create_project(service)
    module = create_module(service, project.id)
    service.add_attribute_definition(project.id, "priority", "Priority", type="string")
    req = service.add_requirement(
        project.id, module.id, {"object_text": "Req", "priority": "High"}
    )

    xlsx_bytes = service.export_xlsx(project.id)
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == ("Id", "Module", "Parent", "ObjectText", "Priority")
    assert rows[1][:2] == (req.id, module.id)
    assert not rows[1][2]
    assert rows[1][3:] == ("Req", "High")


def test_export_docx_structure(service):
    import io

    from docx import Document as DocxDocument

    project = create_project(service)
    module = create_module(service, project.id, "Spec Module", "REQ")
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent text"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child text"}, parent_id=parent.id
    )

    docx_bytes = service.export_docx(project.id)
    document = DocxDocument(io.BytesIO(docx_bytes))
    paragraphs = document.paragraphs

    assert paragraphs[0].text == project.name
    assert paragraphs[0].style.name == "Title"

    module_heading = next(p for p in paragraphs if p.text == "Spec Module")
    assert module_heading.style.name == "Heading 1"

    parent_heading = next(p for p in paragraphs if p.text == parent.id)
    assert parent_heading.style.name == "Heading 2"
    child_heading = next(p for p in paragraphs if p.text == child.id)
    assert child_heading.style.name == "Heading 3"

    # The child heading comes after the parent heading (nested under it).
    assert paragraphs.index(child_heading) > paragraphs.index(parent_heading)

    body_texts = [p.text for p in paragraphs]
    assert any("ObjectText: Parent text" in t for t in body_texts)
    assert any("ObjectText: Child text" in t for t in body_texts)


def test_export_docx_skips_empty_attribute_values(service):
    import io

    from docx import Document as DocxDocument

    project = create_project(service)
    service.add_attribute_definition(project.id, "notes", "Notes", type="string")
    module = create_module(service, project.id)
    req = service.add_requirement(project.id, module.id, {"object_text": "Has text, no notes"})

    docx_bytes = service.export_docx(project.id)
    document = DocxDocument(io.BytesIO(docx_bytes))
    body_texts = [p.text for p in document.paragraphs]

    assert any("ObjectText: Has text, no notes" in t for t in body_texts)
    assert not any(t.startswith("Notes:") for t in body_texts)
