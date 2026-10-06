import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import importlib
import json

import pytest
from fastapi.testclient import TestClient


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REQMAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(tmp_path / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app)


def test_create_and_get_project(client):
    resp = client.post("/projects", json={"name": "Project One"})
    assert resp.status_code == 200, resp.text
    project_id = resp.json()["id"]
    assert project_id == "project-one"

    resp = client.get(f"/projects/{project_id}")
    assert resp.status_code == 200
    assert resp.json()["name"] == "Project One"


def test_full_crud_flow(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    mod1_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    mod2_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module Two", "prefix": "REQ"}).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/modules/{mod1_id}/requirements",
        json={"attributes": {"object_text": "First"}},
    )
    assert resp.status_code == 200, resp.text
    req1_id = resp.json()["id"]
    assert req1_id == "REQ-0001"

    resp = client.post(
        f"/projects/{project_id}/modules/{mod2_id}/requirements",
        json={"attributes": {"object_text": "Second"}},
    )
    assert resp.status_code == 200, resp.text
    req2_id = resp.json()["id"]
    assert req2_id == "REQ-0002"

    config_id = client.post(
        f"/projects/{project_id}/link-configurations",
        json={"name": "Derives", "source_module_id": mod2_id, "target_module_id": mod1_id},
    ).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/links",
        json={"config_id": config_id, "source_id": req2_id, "target_id": req1_id},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["source"]["attributes"][config_id] == f"> {req1_id}"
    assert resp.json()["target"]["attributes"][config_id] == f"< {req2_id}"

    resp = client.put(
        f"/projects/{project_id}/modules/{mod1_id}/requirements/{req1_id}",
        json={"attributes": {"object_text": "First updated"}},
    )
    assert resp.status_code == 200
    assert resp.json()["attributes"]["object_text"] == "First updated"

    resp = client.delete(f"/projects/{project_id}/modules/{mod1_id}/requirements/{req1_id}")
    assert resp.status_code == 200

    resp = client.get(f"/projects/{project_id}/modules/{mod2_id}/requirements/{req2_id}")
    assert config_id not in resp.json()["attributes"]


def test_missing_project_returns_404(client):
    resp = client.get("/projects/does-not-exist")
    assert resp.status_code == 404


def test_export_reqif(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "First"}},
    )
    resp = client.get(f"/projects/{project_id}/export/reqif")
    assert resp.status_code == 200, resp.text
    assert b"<?xml" in resp.content or b"REQ-IF" in resp.content or b"REQ-0001" in resp.content


def test_export_csv(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "First"}},
    )
    resp = client.get(f"/projects/{project_id}/export/csv")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    body = resp.content.decode()
    assert "Id,Module,Parent,ObjectText" in body
    assert "REQ-0001" in body


def test_export_xlsx(client):
    import io

    from openpyxl import load_workbook

    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "First"}},
    )
    resp = client.get(f"/projects/{project_id}/export/xlsx")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = load_workbook(io.BytesIO(resp.content))
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0] == ("Id", "Module", "Parent", "ObjectText")
    assert rows[1][0] == "REQ-0001"


def test_export_docx(client):
    import io

    from docx import Document as DocxDocument

    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "First"}},
    )
    resp = client.get(f"/projects/{project_id}/export/docx")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    document = DocxDocument(io.BytesIO(resp.content))
    body_texts = [p.text for p in document.paragraphs]
    assert "Project One" in body_texts
    assert "REQ-0001" in body_texts


def test_websocket_rpc_roundtrip(client):
    with client.websocket_connect("/ws") as ws:
        ws.send_text(json.dumps({"jsonrpc": "2.0", "id": 1, "method": "project.create", "params": {"name": "Project One"}}))
        response = json.loads(ws.receive_text())
        project_id = response["result"]["id"]
        assert project_id == "project-one"

        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 2, "method": "module.create",
            "params": {"project_id": project_id, "name": "Module One", "prefix": "REQ"},
        }))
        response = json.loads(ws.receive_text())
        module_id = response["result"]["id"]
        assert module_id == "module-one"

        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 3, "method": "requirement.add",
            "params": {"project_id": project_id, "module_id": module_id, "attributes": {"object_text": "First"}},
        }))
        response = json.loads(ws.receive_text())
        assert response["result"]["id"] == "REQ-0001"
        assert response["result"]["attributes"]["object_text"] == "First"


def test_broadcast_notifies_other_ws_clients_on_rest_mutation(client):
    with client.websocket_connect("/ws") as ws1, client.websocket_connect("/ws") as ws2:
        resp = client.post("/projects", json={"name": "Project One"})
        project_id = resp.json()["id"]
        event1 = json.loads(ws1.receive_text())
        event2 = json.loads(ws2.receive_text())
        assert event1 == {"event": "changed", "project_id": project_id}
        assert event2 == {"event": "changed", "project_id": project_id}


def test_broadcast_excludes_sender_on_ws_mutation(client):
    with client.websocket_connect("/ws") as ws1, client.websocket_connect("/ws") as ws2:
        ws1.send_text(json.dumps({"jsonrpc": "2.0", "id": 1, "method": "project.create", "params": {"name": "Project One"}}))
        response = json.loads(ws1.receive_text())
        project_id = response["result"]["id"]

        event2 = json.loads(ws2.receive_text())
        assert event2 == {"event": "changed", "project_id": project_id}


def test_nested_requirement_and_move(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]

    chapter_id = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "Chapter 1"}},
    ).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "Child"}, "parent_id": chapter_id},
    )
    assert resp.status_code == 200, resp.text
    child_id = resp.json()["id"]
    assert resp.json()["parent_id"] == chapter_id

    chapter2_id = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "Chapter 2"}},
    ).json()["id"]

    resp = client.put(
        f"/projects/{project_id}/modules/{module_id}/requirements/{child_id}/parent",
        json={"parent_id": chapter2_id},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["parent_id"] == chapter2_id

    resp = client.put(
        f"/projects/{project_id}/modules/{module_id}/requirements/{child_id}/parent",
        json={"parent_id": child_id},
    )
    assert resp.status_code == 400

    resp = client.delete(f"/projects/{project_id}/modules/{module_id}/requirements/{chapter2_id}")
    assert resp.status_code == 200
    resp = client.get(f"/projects/{project_id}/modules/{module_id}/requirements/{child_id}")
    assert resp.json()["parent_id"] is None


def test_get_config_returns_current_data_dir(client, tmp_path):
    resp = client.get("/config")
    assert resp.status_code == 200
    assert resp.json()["data_dir"] == str(tmp_path)


def test_set_config_switches_data_dir_and_persists(client, tmp_path):
    new_dir = tmp_path / "other-root"
    resp = client.post("/config", json={"data_dir": str(new_dir)})
    assert resp.status_code == 200, resp.text
    assert resp.json()["data_dir"] == str(new_dir)

    # New projects are now created under the new root.
    resp = client.post("/projects", json={"name": "Project One"})
    assert resp.status_code == 200, resp.text
    assert (new_dir / "project-one" / "project.json").exists()
    assert not (tmp_path / "project-one" / "project.json").exists()

    # The choice survives a process restart (module reload).
    config_file = tmp_path / "config.json"
    assert json.loads(config_file.read_text())["data_dir"] == str(new_dir)


def test_set_config_broadcasts_root_changed_to_ws_clients(client, tmp_path):
    with client.websocket_connect("/ws") as ws:
        resp = client.post("/config", json={"data_dir": str(tmp_path / "other-root")})
        assert resp.status_code == 200
        event = json.loads(ws.receive_text())
        assert event == {"event": "root_changed"}


def test_link_configuration_crud_and_direction_enforcement(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    src_id = client.post(f"/projects/{project_id}/modules", json={"name": "Source", "prefix": "SRC"}).json()["id"]
    tgt_id = client.post(f"/projects/{project_id}/modules", json={"name": "Target", "prefix": "TGT"}).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/link-configurations",
        json={"name": "Verifies", "source_module_id": src_id, "target_module_id": tgt_id},
    )
    assert resp.status_code == 200, resp.text
    config_id = resp.json()["id"]

    resp = client.get(f"/projects/{project_id}/link-configurations")
    assert resp.status_code == 200
    assert [c["id"] for c in resp.json()] == [config_id]

    source_id = client.post(
        f"/projects/{project_id}/modules/{src_id}/requirements", json={"attributes": {"object_text": "S"}}
    ).json()["id"]
    target_id = client.post(
        f"/projects/{project_id}/modules/{tgt_id}/requirements", json={"attributes": {"object_text": "T"}}
    ).json()["id"]

    # Wrong direction is rejected.
    resp = client.post(
        f"/projects/{project_id}/links",
        json={"config_id": config_id, "source_id": target_id, "target_id": source_id},
    )
    assert resp.status_code == 400

    # Correct direction succeeds and both sides' mirrors are set.
    resp = client.post(
        f"/projects/{project_id}/links",
        json={"config_id": config_id, "source_id": source_id, "target_id": target_id},
    )
    assert resp.status_code == 200, resp.text

    resp = client.get(f"/projects/{project_id}/modules/{src_id}/requirements/{source_id}")
    assert resp.json()["attributes"][config_id] == f"> {target_id}"
    resp = client.get(f"/projects/{project_id}/modules/{tgt_id}/requirements/{target_id}")
    assert resp.json()["attributes"][config_id] == f"< {source_id}"

    # Removing the configuration cleans up the attribute definition and mirrored values.
    resp = client.delete(f"/projects/{project_id}/link-configurations/{config_id}")
    assert resp.status_code == 200

    resp = client.get(f"/projects/{project_id}/modules/{src_id}/requirements/{source_id}")
    assert config_id not in resp.json()["attributes"]
    resp = client.get(f"/projects/{project_id}")
    assert config_id not in {a["id"] for a in resp.json()["attributes"]}


def test_link_requirements_without_configuration_returns_404(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Mod", "prefix": "REQ"}).json()["id"]
    a_id = client.post(f"/projects/{project_id}/modules/{module_id}/requirements", json={"attributes": {"object_text": "A"}}).json()["id"]
    b_id = client.post(f"/projects/{project_id}/modules/{module_id}/requirements", json={"attributes": {"object_text": "B"}}).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/links",
        json={"config_id": "does-not-exist", "source_id": a_id, "target_id": b_id},
    )
    assert resp.status_code == 404


def test_reorder_requirement(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Mod", "prefix": "REQ"}).json()["id"]
    r1_id = client.post(f"/projects/{project_id}/modules/{module_id}/requirements", json={"attributes": {"object_text": "First"}}).json()["id"]
    r2_id = client.post(f"/projects/{project_id}/modules/{module_id}/requirements", json={"attributes": {"object_text": "Second"}}).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements/{r2_id}/reorder",
        json={"direction": "up"},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["order"] == 0

    resp = client.get(f"/projects/{project_id}/modules/{module_id}/requirements/{r1_id}")
    assert resp.json()["order"] == 1


def test_copy_requirement(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Mod", "prefix": "REQ"}).json()["id"]
    parent_id = client.post(f"/projects/{project_id}/modules/{module_id}/requirements", json={"attributes": {"object_text": "Parent"}}).json()["id"]
    child_id = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements",
        json={"attributes": {"object_text": "Child"}, "parent_id": parent_id},
    ).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/requirements/{parent_id}/copy",
        json={"target_parent_id": None},
    )
    assert resp.status_code == 200, resp.text
    copy_id = resp.json()["id"]
    assert copy_id != parent_id

    resp = client.get(f"/projects/{project_id}/modules/{module_id}/requirements")
    all_reqs = resp.json()
    copied_children = [r for r in all_reqs if r["parent_id"] == copy_id]
    assert len(copied_children) == 1
    assert copied_children[0]["id"] != child_id


def test_browse_folders_lists_subdirectories(client, tmp_path):
    (tmp_path / "alpha").mkdir()
    (tmp_path / "beta").mkdir()
    (tmp_path / "not_a_dir.txt").write_text("x")

    resp = client.get("/fs/browse", params={"path": str(tmp_path)})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["path"] == str(tmp_path)
    assert [d["name"] for d in body["directories"]] == ["alpha", "beta"]


def test_browse_folders_missing_path_returns_404(client, tmp_path):
    resp = client.get("/fs/browse", params={"path": str(tmp_path / "does-not-exist")})
    assert resp.status_code == 404


def test_browse_folders_defaults_without_path_param(client):
    resp = client.get("/fs/browse")
    assert resp.status_code == 200
    assert "directories" in resp.json()


def test_view_crud(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]

    resp = client.post(f"/projects/{project_id}/views", json={"name": "Compact", "attribute_ids": ["object_text"]})
    assert resp.status_code == 200, resp.text
    view_id = resp.json()["id"]

    resp = client.get(f"/projects/{project_id}/views")
    assert resp.status_code == 200
    assert [v["id"] for v in resp.json()] == [view_id]

    resp = client.delete(f"/projects/{project_id}/views/{view_id}")
    assert resp.status_code == 200

    resp = client.get(f"/projects/{project_id}/views")
    assert resp.json() == []


def test_create_view_unknown_attribute_returns_400(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    resp = client.post(f"/projects/{project_id}/views", json={"name": "Compact", "attribute_ids": ["nope"]})
    assert resp.status_code == 400


def test_find_requirement_across_modules(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    mod1_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    mod2_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module Two", "prefix": "TEST"}).json()["id"]
    req_id = client.post(
        f"/projects/{project_id}/modules/{mod2_id}/requirements",
        json={"attributes": {"object_text": "Somewhere"}},
    ).json()["id"]

    resp = client.get(f"/projects/{project_id}/requirements/{req_id}")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["module_id"] == mod2_id
    assert body["requirement"]["id"] == req_id


def test_find_requirement_missing_returns_404(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    resp = client.get(f"/projects/{project_id}/requirements/REQ-9999")
    assert resp.status_code == 404


def test_import_csv_rest_and_broadcast(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]

    with client.websocket_connect("/ws") as ws:
        resp = client.post(
            f"/projects/{project_id}/modules/{module_id}/import/csv",
            content="﻿Id,Parent,ObjectText\n1,,Chapter\n2,1,Child\n".encode("utf-8"),
            headers={"Content-Type": "text/csv"},
        )
        assert resp.status_code == 200, resp.text
        assert resp.json() == {"imported": 2, "requirement_ids": ["REQ-0001", "REQ-0002"], "ignored_columns": []}
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": project_id}

    requirements = client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json()
    assert {r["id"]: r["parent_id"] for r in requirements} == {"REQ-0001": None, "REQ-0002": "REQ-0001"}


def test_import_csv_rest_errors(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    url = f"/projects/{project_id}/modules/{module_id}/import/csv"

    resp = client.post(url, content=b"ObjectText,Nope\nx,y\n", headers={"Content-Type": "text/csv"})
    assert resp.status_code == 400
    assert "Unknown column(s) ['Nope']" in resp.json()["detail"]

    resp = client.post(url, content="ObjectText\nZażółć\n".encode("cp1250"), headers={"Content-Type": "text/csv"})
    assert resp.status_code == 400
    assert "UTF-8" in resp.json()["detail"]

    resp = client.post(
        f"/projects/{project_id}/modules/missing/import/csv", content=b"ObjectText\nx\n", headers=CSV_HEADERS
    )
    assert resp.status_code == 404


def test_import_csv_websocket_rpc(client):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]

    with client.websocket_connect("/ws") as ws:
        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 1, "method": "requirement.import_csv",
            "params": {"project_id": project_id, "module_id": module_id, "csv": "ObjectText\nFirst\n"},
        }))
        response = json.loads(ws.receive_text())
        assert response["result"] == {"imported": 1, "requirement_ids": ["REQ-0001"], "ignored_columns": []}

        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 2, "method": "requirement.import_csv",
            "params": {"project_id": project_id, "module_id": module_id, "csv": "ObjectText\n\"\"\n,\n"},
        }))
        response = json.loads(ws.receive_text())
        assert response["error"]["code"] == -32602


CSV_HEADERS = {"Content-Type": "text/csv"}


def _project_and_module(client, name="Project One"):
    project_id = client.post("/projects", json={"name": name}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "Module One", "prefix": "REQ"}).json()["id"]
    return project_id, module_id


def _rpc_import(ws, request_id, project_id, module_id, csv_text):
    ws.send_text(json.dumps({
        "jsonrpc": "2.0", "id": request_id, "method": "requirement.import_csv",
        "params": {"project_id": project_id, "module_id": module_id, "csv": csv_text},
    }))
    return json.loads(ws.receive_text())


def test_import_csv_ws_broadcasts_to_other_clients_only_on_success(client):
    project_id, module_id = _project_and_module(client)

    with client.websocket_connect("/ws") as ws1, client.websocket_connect("/ws") as ws2:
        response = _rpc_import(ws1, 1, project_id, module_id, "ObjectText,Bogus\nx,y\n")
        assert response["error"]["code"] == -32602
        assert "Unknown column(s) ['Bogus']" in response["error"]["message"]

        response = _rpc_import(ws1, 2, project_id, "missing", "ObjectText\nx\n")
        assert response["error"]["code"] == -32001

        # A later, distinguishable mutation is the first event ws2 sees, so the
        # failed imports did not broadcast anything.
        other_project = client.post("/projects", json={"name": "Other"}).json()["id"]
        assert json.loads(ws2.receive_text()) == {"event": "changed", "project_id": other_project}
        assert json.loads(ws1.receive_text()) == {"event": "changed", "project_id": other_project}

        response = _rpc_import(ws1, 3, project_id, module_id, "ObjectText\nx\n")
        assert response["result"]["requirement_ids"] == ["REQ-0001"]
        assert json.loads(ws2.receive_text()) == {"event": "changed", "project_id": project_id}

        # The sender gets no broadcast for its own import: its next message is
        # the event from this REST call.
        client.post("/projects", json={"name": "Third"})
        assert json.loads(ws1.receive_text()) == {"event": "changed", "project_id": "third"}


def test_import_csv_rest_failure_does_not_broadcast(client):
    project_id, module_id = _project_and_module(client)
    url = f"/projects/{project_id}/modules/{module_id}/import/csv"

    with client.websocket_connect("/ws") as ws:
        assert client.post(url, content=b"Id,Parent,ObjectText\nA,A,x\n", headers=CSV_HEADERS).status_code == 400
        assert client.post(url, content="ObjectText\nZażółć\n".encode("cp1250"), headers=CSV_HEADERS).status_code == 400
        assert client.post(
            f"/projects/{project_id}/modules/missing/import/csv", content=b"ObjectText\nx\n", headers=CSV_HEADERS
        ).status_code == 404
        assert client.post(url, content=b"ObjectText\nx\n", headers={"Content-Type": "text/plain"}).status_code == 415

        other_project = client.post("/projects", json={"name": "Other"}).json()["id"]
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": other_project}

    assert client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json() == []
    assert client.get(f"/projects/{project_id}").json()["id_counters"] == {}


def test_import_csv_rest_missing_project_and_empty_body(client):
    project_id, module_id = _project_and_module(client)

    resp = client.post(f"/projects/missing/modules/{module_id}/import/csv", content=b"ObjectText\nx\n", headers=CSV_HEADERS)
    assert resp.status_code == 404

    resp = client.post(f"/projects/{project_id}/modules/{module_id}/import/csv", content=b"", headers=CSV_HEADERS)
    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"]


def test_import_csv_rest_reports_every_row_error_in_detail(client):
    project_id, module_id = _project_and_module(client)

    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/import/csv",
        content=b"Id,Parent,ObjectText\nA,,ok\nB,Z,orphan\nC,,\n",
        headers={"Content-Type": "text/csv"},
    )

    assert resp.status_code == 400
    detail = resp.json()["detail"]
    assert detail.startswith("CSV import failed; nothing was imported.")
    assert "Row 3: Parent 'Z' does not match" in detail
    assert "Row 4: required attribute 'ObjectText' is empty" in detail


def test_import_csv_ws_rejects_missing_or_non_string_csv_param(client):
    project_id, module_id = _project_and_module(client)

    with client.websocket_connect("/ws") as ws:
        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 1, "method": "requirement.import_csv",
            "params": {"project_id": project_id, "module_id": module_id},
        }))
        assert json.loads(ws.receive_text())["error"]["code"] == -32602

        response = _rpc_import(ws, 2, project_id, module_id, ["ObjectText", "x"])
        assert response["error"]["code"] == -32602
        assert "must be a string" in response["error"]["message"]

    assert client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json() == []


def test_import_csv_rest_and_ws_parity(client):
    rest_project, rest_module = _project_and_module(client, "Rest Project")
    ws_project, ws_module = _project_and_module(client, "Ws Project")
    client.post(f"/projects/{rest_project}/attributes", json={"id": "safety", "name": "Safety", "type": "boolean"})
    client.post(f"/projects/{ws_project}/attributes", json={"id": "safety", "name": "Safety", "type": "boolean"})
    csv_text = "Id,Parent,ObjectText,Safety\nC,P,Child,no\nP,,\"Parent, \"\"quoted\"\"\",yes\n"

    rest = client.post(
        f"/projects/{rest_project}/modules/{rest_module}/import/csv",
        content=csv_text.encode("utf-8"),
        headers={"Content-Type": "text/csv"},
    )
    assert rest.status_code == 200, rest.text
    with client.websocket_connect("/ws") as ws:
        ws_result = _rpc_import(ws, 1, ws_project, ws_module, csv_text)["result"]

    assert rest.json() == ws_result == {
        "imported": 2, "requirement_ids": ["REQ-0002", "REQ-0001"], "ignored_columns": [],
    }

    def tree(project_id, module_id):
        reqs = client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json()
        return sorted((r["id"], r["parent_id"], r["order"], r["attributes"]) for r in reqs)

    assert tree(rest_project, rest_module) == tree(ws_project, ws_module) == [
        ("REQ-0001", None, 0, {"object_text": 'Parent, "quoted"', "safety": True}),
        ("REQ-0002", "REQ-0001", 0, {"object_text": "Child", "safety": False}),
    ]


def test_import_csv_rest_roundtrips_rest_export(client):
    project_id, module_id = _project_and_module(client)
    client.post(f"/projects/{project_id}/attributes", json={"id": "priority", "name": "Priority", "type": "integer"})
    base = f"/projects/{project_id}/modules/{module_id}/requirements"
    parent_id = client.post(base, json={"attributes": {"object_text": "Zażółć — \"q\", 🚀", "priority": -2}}).json()["id"]
    client.post(base, json={"attributes": {"object_text": "Line 1\nLine 2"}, "parent_id": parent_id})
    exported = client.get(f"/projects/{project_id}/export/csv")
    assert exported.status_code == 200
    target_id = client.post(f"/projects/{project_id}/modules", json={"name": "Target", "prefix": "TGT"}).json()["id"]

    resp = client.post(
        f"/projects/{project_id}/modules/{target_id}/import/csv",
        content=exported.content,
        headers={"Content-Type": "text/csv"},
    )

    assert resp.status_code == 200, resp.text
    copies = {
        r["attributes"]["object_text"]: r
        for r in client.get(f"/projects/{project_id}/modules/{target_id}/requirements").json()
    }
    parent = copies["Zażółć — \"q\", 🚀"]
    assert parent["attributes"]["priority"] == -2 and parent["parent_id"] is None
    assert copies["Line 1\nLine 2"]["parent_id"] == parent["id"]


@pytest.mark.parametrize(
    "headers",
    [
        {},
        {"Content-Type": "text/plain"},
        {"Content-Type": "application/x-www-form-urlencoded"},
        {"Content-Type": "multipart/form-data; boundary=x"},
        {"Content-Type": "application/json"},
        {"Content-Type": "text/csvx"},
    ],
)
def test_import_csv_rest_requires_text_csv_content_type(client, headers):
    project_id, module_id = _project_and_module(client)

    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/import/csv", content=b"ObjectText\nx\n", headers=headers
    )

    assert resp.status_code == 415
    assert "text/csv" in resp.json()["detail"]
    assert client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json() == []


@pytest.mark.parametrize("content_type", ["text/csv; charset=utf-8", "Text/CSV", " text/csv ;charset=UTF-8"])
def test_import_csv_rest_accepts_text_csv_with_parameters(client, content_type):
    project_id, module_id = _project_and_module(client)

    resp = client.post(
        f"/projects/{project_id}/modules/{module_id}/import/csv",
        content=b"ObjectText\nx\n",
        headers={"Content-Type": content_type},
    )

    assert resp.status_code == 200, resp.text


def test_import_csv_rest_rejects_body_over_the_size_cap(client):
    from reqman.csv_import import MAX_CSV_IMPORT_BYTES

    project_id, module_id = _project_and_module(client)
    url = f"/projects/{project_id}/modules/{module_id}/import/csv"
    header = b"ObjectText\n"
    at_cap = header + b"x" * (MAX_CSV_IMPORT_BYTES - len(header))

    with client.websocket_connect("/ws") as ws:
        resp = client.post(url, content=at_cap + b"x", headers=CSV_HEADERS)
        assert resp.status_code == 413
        assert "too large" in resp.json()["detail"]

        # Without a Content-Length header the body is still capped while reading.
        def chunks():
            yield at_cap
            yield b"x"

        resp = client.post(url, content=chunks(), headers=CSV_HEADERS)
        assert resp.status_code == 413

        other_project = client.post("/projects", json={"name": "Other"}).json()["id"]
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": other_project}

    assert client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json() == []
    assert client.post(url, content=at_cap, headers=CSV_HEADERS).status_code == 200


def test_import_csv_ws_rejects_text_over_the_size_cap(client):
    from reqman.csv_import import MAX_CSV_IMPORT_BYTES

    project_id, module_id = _project_and_module(client)

    with client.websocket_connect("/ws") as ws:
        response = _rpc_import(ws, 1, project_id, module_id, "ObjectText\n" + "x" * MAX_CSV_IMPORT_BYTES)

    assert response["error"]["code"] == -32602
    assert "too large" in response["error"]["message"]
    assert client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json() == []


# --- attribute names and value validation -------------------------------------


@pytest.mark.parametrize("name", ["Id", "module", " Parent ", "ObjectText", "objecttext", "  "])
def test_add_attribute_rejects_reserved_duplicate_and_empty_names(client, name):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    resp = client.post(f"/projects/{project_id}/attributes", json={"id": "custom", "name": name})
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"]

    with client.websocket_connect("/ws") as ws:
        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 1, "method": "attribute.add",
            "params": {"project_id": project_id, "id": "custom", "name": name},
        }))
        response = json.loads(ws.receive_text())
    assert response["error"]["message"] == resp.json()["detail"]


def _typed_project_and_module(client):
    project_id, module_id = _project_and_module(client)
    for body in (
        {"id": "priority", "name": "Priority", "type": "integer"},
        {"id": "safety", "name": "Safety", "type": "boolean"},
        {"id": "status", "name": "Status", "type": "enum", "choices": ["Draft", "Done"]},
    ):
        assert client.post(f"/projects/{project_id}/attributes", json=body).status_code == 200
    return project_id, module_id


@pytest.mark.parametrize(
    "attributes, expected",
    [
        ({"status": "Nope"}, "Status: 'Nope' is not one of ['Draft', 'Done']"),
        ({"priority": "3"}, "Priority: '3' is not an integer"),
        ({"priority": 2.5}, "Priority: '2.5' is not an integer"),
        ({"safety": "yes"}, "Safety: 'yes' is not a boolean"),
        ({"safety": 0}, "Safety: '0' is not a boolean"),
    ],
)
def test_add_and_update_requirement_reject_invalid_values_with_400(client, attributes, expected):
    project_id, module_id = _typed_project_and_module(client)
    url = f"/projects/{project_id}/modules/{module_id}/requirements"

    resp = client.post(url, json={"attributes": {"object_text": "x", **attributes}})
    assert resp.status_code == 400, resp.text
    assert expected in resp.json()["detail"]

    req_id = client.post(url, json={"attributes": {"object_text": "ok"}}).json()["id"]
    resp = client.put(f"{url}/{req_id}", json={"attributes": attributes})
    assert resp.status_code == 400, resp.text
    assert expected in resp.json()["detail"]

    with client.websocket_connect("/ws") as ws:
        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 1, "method": "requirement.update",
            "params": {"project_id": project_id, "module_id": module_id, "id": req_id, "attributes": attributes},
        }))
        response = json.loads(ws.receive_text())
    assert expected in response["error"]["message"]
    assert client.get(f"{url}/{req_id}").json()["attributes"] == {"object_text": "ok"}


def test_add_and_update_requirement_accept_typed_values_and_null(client):
    project_id, module_id = _typed_project_and_module(client)
    url = f"/projects/{project_id}/modules/{module_id}/requirements"

    resp = client.post(url, json={"attributes": {"object_text": "x", "priority": 4, "safety": True, "status": "Done"}})
    assert resp.status_code == 200, resp.text
    req_id = resp.json()["id"]

    resp = client.put(f"{url}/{req_id}", json={"attributes": {"priority": None, "status": None, "safety": False}})
    assert resp.status_code == 200, resp.text
    assert resp.json()["attributes"] == {"object_text": "x", "priority": None, "safety": False, "status": None}


def _rpc(ws, method, params, request_id=1):
    ws.send_text(json.dumps({"jsonrpc": "2.0", "id": request_id, "method": method, "params": params}))
    return json.loads(ws.receive_text())


# --- R1: link configuration names ------------------------------------------------


@pytest.mark.parametrize("name", ["Parent", " id ", "objecttext", "Priority", "  "])
def test_create_link_configuration_rejects_invalid_names_with_same_rest_and_ws_message(client, name):
    project_id, src_id = _project_and_module(client)
    tgt_id = client.post(f"/projects/{project_id}/modules", json={"name": "Target", "prefix": "TGT"}).json()["id"]
    client.post(f"/projects/{project_id}/attributes", json={"id": "priority", "name": "Priority"})
    body = {"name": name, "source_module_id": src_id, "target_module_id": tgt_id}

    resp = client.post(f"/projects/{project_id}/link-configurations", json=body)
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"].startswith("Link name")

    with client.websocket_connect("/ws") as ws:
        response = _rpc(ws, "link_configuration.create", {"project_id": project_id, **body})
    assert response["error"]["code"] == -32602
    assert response["error"]["message"] == resp.json()["detail"]
    assert client.get(f"/projects/{project_id}/link-configurations").json() == []


# --- R2: attribute defaults and choices ------------------------------------------


@pytest.mark.parametrize(
    "body, expected",
    [
        ({"type": "integer", "default": "5"}, "Invalid default value: Priority: '5' is not an integer"),
        ({"type": "boolean", "default": 1}, "Invalid default value: Priority: '1' is not a boolean"),
        (
            {"type": "enum", "choices": ["Low"], "default": "High"},
            "Invalid default value: Priority: 'High' is not one of ['Low']",
        ),
        ({"type": "enum", "choices": []}, "Enum attribute 'Priority' needs a non-empty list of choices"),
        ({"type": "enum"}, "Enum attribute 'Priority' needs a non-empty list of choices"),
        ({"type": "string", "choices": ["a"]}, "Attribute 'Priority' of type 'string' cannot have choices"),
    ],
)
def test_add_attribute_rejects_invalid_default_and_choices_with_same_rest_and_ws_message(client, body, expected):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    body = {"id": "priority", "name": "Priority", **body}

    resp = client.post(f"/projects/{project_id}/attributes", json=body)
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"] == expected

    with client.websocket_connect("/ws") as ws:
        response = _rpc(ws, "attribute.add", {"project_id": project_id, **body})
    assert response["error"]["message"] == expected
    assert [a["id"] for a in client.get(f"/projects/{project_id}").json()["attributes"]] == ["object_text"]


# --- R4: blank values for required attributes ------------------------------------


@pytest.mark.parametrize("value", [None, "", "   "])
def test_required_attribute_rejects_blank_value_on_rest_and_ws_create_and_update(client, value):
    project_id, module_id = _project_and_module(client)
    url = f"/projects/{project_id}/modules/{module_id}/requirements"
    expected = "required attribute 'ObjectText' is empty"

    resp = client.post(url, json={"attributes": {"object_text": value}})
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"] == expected

    req_id = client.post(url, json={"attributes": {"object_text": "keep"}}).json()["id"]
    resp = client.put(f"{url}/{req_id}", json={"attributes": {"object_text": value}})
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"] == expected

    with client.websocket_connect("/ws") as ws:
        created = _rpc(
            ws, "requirement.add",
            {"project_id": project_id, "module_id": module_id, "attributes": {"object_text": value}},
        )
        updated = _rpc(
            ws, "requirement.update",
            {"project_id": project_id, "module_id": module_id, "id": req_id, "attributes": {"object_text": value}},
            request_id=2,
        )
    assert created["error"] == {"code": -32602, "message": expected}
    assert updated["error"] == {"code": -32602, "message": expected}

    reqs = client.get(url).json()
    assert [(r["id"], r["attributes"]) for r in reqs] == [(req_id, {"object_text": "keep"})]


# --- M1/M2: attribute type and blank required default -----------------------------


@pytest.mark.parametrize(
    "body, expected",
    [
        (
            {"type": "Integer", "default": "5"},
            "Unknown attribute type 'Integer'; expected one of string, text, integer, boolean, enum",
        ),
        (
            {"type": "Enum", "choices": ["Low"]},
            "Unknown attribute type 'Enum'; expected one of string, text, integer, boolean, enum",
        ),
        (
            {"type": "string", "required": True, "default": "  "},
            "Attribute 'Priority' is required, so its default must not be empty",
        ),
    ],
)
def test_add_attribute_rejects_unknown_type_and_blank_required_default_on_rest_and_ws(client, body, expected):
    project_id = client.post("/projects", json={"name": "Project One"}).json()["id"]
    body = {"id": "priority", "name": "Priority", **body}

    resp = client.post(f"/projects/{project_id}/attributes", json=body)
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"] == expected

    with client.websocket_connect("/ws") as ws:
        response = _rpc(ws, "attribute.add", {"project_id": project_id, **body})
    assert response["error"] == {"code": -32602, "message": expected}
    assert [a["id"] for a in client.get(f"/projects/{project_id}").json()["attributes"]] == ["object_text"]


# --- requirement move with an index ----------------------------------------------


def _move_fixture(client):
    """A module with a top-level ``target`` holding three children and a top-level ``mover``."""
    project_id, module_id = _project_and_module(client)
    url = f"/projects/{project_id}/modules/{module_id}/requirements"
    target = client.post(url, json={"attributes": {"object_text": "Target"}}).json()["id"]
    kids = [
        client.post(url, json={"attributes": {"object_text": f"C{i}"}, "parent_id": target}).json()["id"]
        for i in range(3)
    ]
    mover = client.post(url, json={"attributes": {"object_text": "Mover"}}).json()["id"]
    return project_id, module_id, target, kids, mover


def _child_ids(client, project_id, module_id, parent_id):
    reqs = client.get(f"/projects/{project_id}/modules/{module_id}/requirements").json()
    children = sorted((r for r in reqs if r["parent_id"] == parent_id), key=lambda r: (r["order"], r["id"]))
    return [(r["id"], r["order"]) for r in children]


def test_move_requirement_rest_with_index_places_and_broadcasts_once(client):
    project_id, module_id, target, kids, mover = _move_fixture(client)
    url = f"/projects/{project_id}/modules/{module_id}/requirements/{mover}/parent"

    with client.websocket_connect("/ws") as ws:
        resp = client.put(url, json={"parent_id": target, "index": 1})
        assert resp.status_code == 200, resp.text
        assert resp.json()["parent_id"] == target
        assert resp.json()["order"] == 1
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": project_id}
        # The next event is a later, distinguishable mutation: the move broadcast once.
        client.post("/projects", json={"name": "Other"})
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": "other"}

    assert _child_ids(client, project_id, module_id, target) == [(kids[0], 0), (mover, 1), (kids[1], 2), (kids[2], 3)]

    # Without an index the move still appends, as before.
    resp = client.put(url, json={"parent_id": None})
    assert resp.status_code == 200, resp.text
    assert resp.json()["parent_id"] is None
    assert resp.json()["order"] == 1


def test_move_requirement_ws_with_index_places_and_broadcasts_once(client):
    project_id, module_id, target, kids, mover = _move_fixture(client)

    with client.websocket_connect("/ws") as ws1, client.websocket_connect("/ws") as ws2:
        response = _rpc(ws1, "requirement.move", {
            "project_id": project_id, "module_id": module_id, "id": mover, "parent_id": target, "index": 0,
        })
        assert response["result"]["parent_id"] == target
        assert response["result"]["order"] == 0
        assert json.loads(ws2.receive_text()) == {"event": "changed", "project_id": project_id}
        client.post("/projects", json={"name": "Other"})
        assert json.loads(ws2.receive_text()) == {"event": "changed", "project_id": "other"}
        # The sender got no broadcast of its own move.
        assert json.loads(ws1.receive_text()) == {"event": "changed", "project_id": "other"}

    assert _child_ids(client, project_id, module_id, target) == [(mover, 0), (kids[0], 1), (kids[1], 2), (kids[2], 3)]


@pytest.mark.parametrize("bad", [-1, 1.5, "1", "abc", True, [0]])
def test_move_requirement_invalid_index_is_rejected_alike_via_rest_and_ws(client, bad):
    project_id, module_id, target, kids, mover = _move_fixture(client)
    body = {"parent_id": target, "index": bad}

    resp = client.put(f"/projects/{project_id}/modules/{module_id}/requirements/{mover}/parent", json=body)
    assert resp.status_code == 400, resp.text
    assert "must be a non-negative integer" in resp.json()["detail"]

    with client.websocket_connect("/ws") as ws:
        response = _rpc(ws, "requirement.move", {"project_id": project_id, "module_id": module_id, "id": mover, **body})
    assert response["error"]["code"] == -32602
    assert response["error"]["message"] == resp.json()["detail"]

    assert _child_ids(client, project_id, module_id, target) == [(kids[0], 0), (kids[1], 1), (kids[2], 2)]
