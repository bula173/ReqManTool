"""Project lifecycle: lock / unlock, remove, and the baseline glue (lock after baseline, report data source).

See docs/ARCHITECTURE.md "Project lifecycle: lock and remove". Tests with git
history work in pytest's ``tmp_path`` (outside any repository) with their own
repositories and local bare remotes; ``conftest.py`` isolates git from the
user's configuration.
"""

from __future__ import annotations

import getpass
import importlib
import io
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient

from reqman.formats import PROJECT_FORMAT
from reqman.git_repo import GitHistory, GitSettings
from reqman.service import ConflictError, ProjectLockedError, ProjectService
from reqman.store import ProjectStore

CHECKOUT = Path(__file__).resolve().parents[2]
STAGING = ".reqman-staging"


class Clock:
    def __init__(self) -> None:
        self.offset = 0.0

    def __call__(self) -> float:
        return time.time() + self.offset


def git(path: Path, *args: str, check: bool = True) -> str:
    assert CHECKOUT not in [path.resolve(), *path.resolve().parents], "never run git in the checkout"
    env = {k: v for k, v in os.environ.items() if k not in ("GIT_DIR", "GIT_WORK_TREE")}
    result = subprocess.run(["git", *args], cwd=path, capture_output=True, text=True, env=env, check=False)
    if check:
        assert result.returncode == 0, result.stderr
    return result.stdout


def tree_bytes(directory: Path) -> dict[str, bytes]:
    return {
        str(p.relative_to(directory)): p.read_bytes()
        for p in sorted(directory.rglob("*"))
        if p.is_file() and ".git" not in p.relative_to(directory).parts
    }


@pytest.fixture
def root(tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    probe = subprocess.run(["git", "rev-parse", "--show-toplevel"], cwd=data, capture_output=True, check=False)
    assert probe.returncode != 0, "tmp_path is inside a git repository"
    return data


@pytest.fixture
def plain(root):
    """History off (plain files)."""
    return ProjectService(ProjectStore(root))


def make_git_service(root: Path, **settings) -> ProjectService:
    settings.setdefault("squash_window_seconds", 0)
    history = GitHistory(GitSettings(enabled=True, **settings), clock=Clock())
    return ProjectService(ProjectStore(root), history=history)


@pytest.fixture
def gitsvc(root):
    return make_git_service(root)


def subjects(folder: Path) -> list[str]:
    return git(folder, "log", "--format=%s").splitlines()


def project_file(root: Path, project_id: str) -> Path:
    return root / project_id / "project.json"


LOGIN = getpass.getuser()
ISO_UTC = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")


# --- lock: data and format --------------------------------------------------------------------


def test_lock_records_who_when_why_and_round_trips(plain, root):
    project = plain.create_project("Alpha")
    locked = plain.lock_project(project.id, "  Release \n 1\t ")
    assert locked.locked is not None
    lock = locked.locked.to_dict()
    assert lock["by"] == LOGIN and lock["email"] == f"{LOGIN}@localhost"
    assert ISO_UTC.match(lock["at"])
    assert lock["reason"] == "Release 1"
    assert "baseline" not in lock

    text = project_file(root, project.id).read_text(encoding="utf-8")
    data = json.loads(text)
    assert text == json.dumps(data, indent=2, ensure_ascii=False) + "\n"
    assert data["format_version"] == PROJECT_FORMAT.current
    keys = list(data)
    assert keys.index("locked") == keys.index("description") + 1
    assert keys.index("attributes") == keys.index("locked") + 1
    assert data["locked"] == {"by": LOGIN, "email": f"{LOGIN}@localhost", "at": lock["at"], "reason": "Release 1"}

    reloaded = ProjectService(ProjectStore(root)).get_project(project.id)
    assert reloaded.locked.to_dict() == lock

    plain.unlock_project(project.id)
    text = project_file(root, project.id).read_text(encoding="utf-8")
    assert "locked" not in json.loads(text)
    assert text.endswith("}\n") and '"locked"' not in text


def test_locked_key_sits_after_references_and_parent(plain, root):
    parent = plain.create_project("Parent")
    child = plain.create_project("Child", parent_project_id=parent.id)
    other = plain.create_project("Other")
    plain.set_project_references(child.id, [other.id])
    plain.lock_project(child.id)
    keys = list(json.loads(project_file(root, child.id).read_text(encoding="utf-8")))
    assert keys[keys.index("parent_project_id") + 1] == "references"
    assert keys[keys.index("references") + 1] == "locked"


def test_lock_without_reason_and_reason_limits(plain):
    project = plain.create_project("Alpha")
    assert plain.lock_project(project.id).locked.reason == ""
    plain.unlock_project(project.id)
    assert plain.lock_project(project.id, "x" * 500).locked.reason == "x" * 500
    plain.unlock_project(project.id)
    with pytest.raises(ValueError, match="too long"):
        plain.lock_project(project.id, "x" * 501)
    with pytest.raises(ValueError, match="must be a string"):
        plain.lock_project(project.id, 5)
    assert plain.get_project(project.id).locked is None
    with pytest.raises(FileNotFoundError):
        plain.lock_project("missing")


def test_locking_a_locked_project_is_refused_and_changes_nothing(plain, root):
    project = plain.create_project("Alpha")
    plain.lock_project(project.id, "first")
    before = project_file(root, project.id).read_bytes()
    with pytest.raises(ConflictError):
        plain.lock_project(project.id, "second")
    assert project_file(root, project.id).read_bytes() == before


def test_unlocking_an_unlocked_project_writes_nothing(plain, root):
    project = plain.create_project("Alpha")
    path = project_file(root, project.id)
    before = (path.read_bytes(), path.stat().st_mtime_ns)
    assert plain.unlock_project(project.id).locked is None
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before


def _rewrite(root: Path, project_id: str, **changes) -> None:
    path = project_file(root, project_id)
    data = json.loads(path.read_text(encoding="utf-8"))
    for key, value in changes.items():
        if value is ...:
            data.pop(key, None)
        else:
            data[key] = value
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


def test_version_7_file_with_a_stray_locked_key_loads_unlocked(plain, root):
    project = plain.create_project("Alpha")
    _rewrite(root, project.id, format_version=7, locked={"by": "Hand", "reason": "stray"})
    fresh = ProjectService(ProjectStore(root))
    assert fresh.get_project(project.id).locked is None
    module = fresh.create_module(project.id, "System", "SYS")  # not refused
    data = json.loads(project_file(root, project.id).read_text(encoding="utf-8"))
    assert data["format_version"] == PROJECT_FORMAT.current and "locked" not in data
    assert module.id in [m.id for m in fresh.list_modules(project.id)]


@pytest.mark.parametrize("value", [True, 1, "yes", ["x"], {}])
def test_truthy_hand_edited_lock_still_means_locked(plain, root, value):
    project = plain.create_project("Alpha")
    _rewrite(root, project.id, locked=value)
    fresh = ProjectService(ProjectStore(root))
    lock = fresh.get_project(project.id).locked
    # Any truthy non-object value (and an empty object) is a lock with unknown details.
    assert lock is not None and lock.to_dict() == {"by": "", "email": "", "at": "", "reason": ""}
    with pytest.raises(ProjectLockedError):
        fresh.create_module(project.id, "System", "SYS")
    fresh.unlock_project(project.id)
    assert "locked" not in json.loads(project_file(root, project.id).read_text(encoding="utf-8"))


@pytest.mark.parametrize("value", [None, False, 0, ""])
def test_falsy_hand_edited_lock_means_unlocked(plain, root, value):
    project = plain.create_project("Alpha")
    _rewrite(root, project.id, locked=value)
    fresh = ProjectService(ProjectStore(root))
    assert fresh.get_project(project.id).locked is None
    fresh.create_module(project.id, "System", "SYS")
    assert "locked" not in json.loads(project_file(root, project.id).read_text(encoding="utf-8"))


def test_hand_edited_lock_with_odd_field_types_keeps_strings_only(plain, root):
    project = plain.create_project("Alpha")
    _rewrite(root, project.id, locked={"by": 5, "email": None, "at": "2026-01-01T00:00:00Z", "reason": ["x"],
                                       "baseline": ""})
    lock = ProjectService(ProjectStore(root)).get_project(project.id).locked
    assert lock.to_dict() == {"by": "", "email": "", "at": "2026-01-01T00:00:00Z", "reason": ""}


# --- lock: who ------------------------------------------------------------------------------


def test_who_with_history_off_is_the_login_name(plain):
    project = plain.create_project("Alpha")
    lock = plain.lock_project(project.id).locked
    assert (lock.by, lock.email) == (LOGIN, f"{LOGIN}@localhost")


def test_who_with_history_off_honours_reqman_git_author_env(root, monkeypatch):
    monkeypatch.setenv("REQMAN_GIT_AUTHOR_NAME", "Rita Reqman")
    monkeypatch.setenv("REQMAN_GIT_AUTHOR_EMAIL", "rita@example.com")
    settings = GitSettings.from_env(dict(os.environ))
    assert settings.enabled is False
    service = ProjectService(ProjectStore(root), history=GitHistory(settings))
    project = service.create_project("Alpha")
    lock = service.lock_project(project.id).locked
    assert (lock.by, lock.email) == ("Rita Reqman", "rita@example.com")


def test_who_with_history_on_but_no_repository_asks_git(plain, root):
    project = plain.create_project("Alpha")
    assert not (root / project.id / ".git").exists()
    service = make_git_service(root)
    lock = service.lock_project(project.id).locked
    assert (lock.by, lock.email) == ("Test Author", "author@example.com")
    assert not (root / project.id / ".git").exists()


def test_who_with_a_repository_is_the_commit_author(gitsvc, root):
    project = gitsvc.create_project("Alpha")
    lock = gitsvc.lock_project(project.id).locked
    author = git(root / project.id, "log", "-1", "--format=%an|%ae").strip()
    assert f"{lock.by}|{lock.email}" == author == "Test Author|author@example.com"


def test_who_with_reqman_git_author_env_and_history_on(root):
    service = make_git_service(root, author_name="Rita Reqman", author_email="rita@example.com")
    project = service.create_project("Alpha")
    lock = service.lock_project(project.id).locked
    assert (lock.by, lock.email) == ("Rita Reqman", "rita@example.com")
    assert git(root / project.id, "log", "-1", "--format=%an <%ae>").strip() == "Rita Reqman <rita@example.com>"


# --- lock: history messages -------------------------------------------------------------------


def test_lock_and_unlock_commit_messages(gitsvc, root):
    project = gitsvc.create_project("Alpha")
    folder = root / project.id
    gitsvc.lock_project(project.id, "Release 1")
    assert subjects(folder)[0] == "Project locked: Release 1"
    gitsvc.unlock_project(project.id)
    assert subjects(folder)[0] == "Project unlocked"
    gitsvc.lock_project(project.id)
    assert subjects(folder)[0] == "Project locked"
    gitsvc.unlock_project(project.id)
    count = len(subjects(folder))
    gitsvc.unlock_project(project.id)  # nothing written, nothing committed
    assert len(subjects(folder)) == count
    assert git(folder, "status", "--porcelain") == ""


# --- lock: the guard (service) ------------------------------------------------------------------


class Fixture:
    """A project with two modules, a link configuration, a link, a view and a module folder."""

    def __init__(self, service: ProjectService) -> None:
        self.service = service
        self.pid = service.create_project("Locked").id
        self.other = service.create_project("Other").id
        service.add_attribute_definition(self.pid, "status", "Status", "enum", choices=["Draft", "Reviewed"])
        self.sys = service.create_module(self.pid, "System", "SYS").id
        self.tst = service.create_module(self.pid, "Tests", "TST").id
        self.r1 = service.add_requirement(self.pid, self.sys, {"object_text": "one"}).id
        self.r2 = service.add_requirement(self.pid, self.sys, {"object_text": "two"}).id
        self.t1 = service.add_requirement(self.pid, self.tst, {"object_text": "test"}).id
        self.t2 = service.add_requirement(self.pid, self.tst, {"object_text": "test 2"}).id
        self.cfg = service.create_link_configuration(self.pid, "Verifies", self.tst, self.sys).id
        service.link_requirements(self.pid, self.cfg, self.t1, self.r1)
        self.view = service.create_view(self.pid, "Review", ["status"]).id
        self.folder = service.create_module_folder(self.pid, "Folder")["id"]


SERVICE_MUTATIONS = {
    "add_requirement": lambda s, f: s.add_requirement(f.pid, f.sys, {"object_text": "x"}),
    "update_requirement": lambda s, f: s.update_requirement(f.pid, f.sys, f.r1, {"object_text": "y"}),
    "remove_requirement": lambda s, f: s.remove_requirement(f.pid, f.sys, f.r2),
    "move_requirement": lambda s, f: s.move_requirement(f.pid, f.sys, f.r2, f.r1),
    "reorder_requirement": lambda s, f: s.reorder_requirement(f.pid, f.sys, f.r1, "down"),
    "copy_requirement": lambda s, f: s.copy_requirement(f.pid, f.sys, f.r1, None),
    "create_module": lambda s, f: s.create_module(f.pid, "New", "NEW"),
    "reorder_module": lambda s, f: s.reorder_module(f.pid, f.tst, "up"),
    "move_module": lambda s, f: s.move_module(f.pid, f.sys, f.folder),
    "purge_module": lambda s, f: s.purge_module(f.pid, f.tst),
    "purge_project": lambda s, f: s.purge_project(f.pid),
    "add_attribute": lambda s, f: s.add_attribute_definition(f.pid, "x", "X"),
    "set_attribute_modules": lambda s, f: s.set_attribute_modules(f.pid, "status", [f.sys]),
    "set_row_styles": lambda s, f: s.set_row_styles(f.pid, []),
    "set_references": lambda s, f: s.set_project_references(f.pid, [f.other]),
    "set_parent": lambda s, f: s.set_project_parent(f.pid, f.other),
    "create_link_configuration": lambda s, f: s.create_link_configuration(f.pid, "Other", f.sys, f.tst),
    "remove_link_configuration": lambda s, f: s.remove_link_configuration(f.pid, f.cfg),
    "link": lambda s, f: s.link_requirements(f.pid, f.cfg, f.t2, f.r2),
    "unlink": lambda s, f: s.unlink_requirements(f.pid, f.cfg, f.t1, f.r1),
    "create_view": lambda s, f: s.create_view(f.pid, "Another", []),
    "update_view": lambda s, f: s.update_view(f.pid, f.view, {"name": "Renamed"}),
    "set_view_column_widths": lambda s, f: s.set_view_column_widths(f.pid, f.view, {}),
    "remove_view": lambda s, f: s.remove_view(f.pid, f.view),
    "import_csv": lambda s, f: s.import_csv(f.pid, f.sys, "ObjectText\nimported\n"),
    "create_module_folder": lambda s, f: s.create_module_folder(f.pid, "Another"),
    "rename_module_folder": lambda s, f: s.rename_module_folder(f.pid, f.folder, "Renamed"),
    "delete_module_folder": lambda s, f: s.delete_module_folder(f.pid, f.folder),
    "move_module_folder": lambda s, f: s.move_module_folder(f.pid, f.folder, None, 0),
    "reorder_module_folder": lambda s, f: s.reorder_module_folder(f.pid, f.folder, "down"),
}


@pytest.mark.parametrize("operation", sorted(SERVICE_MUTATIONS))
def test_every_service_mutation_is_refused_while_locked(plain, root, operation):
    f = Fixture(plain)
    plain.lock_project(f.pid, "Release")
    before = tree_bytes(root)
    with pytest.raises(ProjectLockedError, match="is locked .*Release.* Unlock it to make changes; nothing was changed"):
        SERVICE_MUTATIONS[operation](plain, f)
    assert tree_bytes(root) == before
    # Unlocked, the same operation works (the fixture is valid for it).
    plain.unlock_project(f.pid)
    SERVICE_MUTATIONS[operation](plain, f)


def test_allowed_while_locked(plain, root):
    f = Fixture(plain)
    plain.lock_project(f.pid, "Release")
    before = tree_bytes(root / f.pid)
    assert plain.get_project(f.pid).locked is not None
    assert [m.id for m in plain.list_modules(f.pid)] == [f.sys, f.tst]
    assert [r.id for r in plain.list_requirements(f.pid, f.sys)] == [f.r1, f.r2]
    assert plain.export_csv(f.pid)
    assert plain.export_xlsx(f.pid)
    assert plain.export_docx(f.pid)
    assert plain.export_reqif(f.pid)
    assert plain.traceability_matrix(f.pid, [f.cfg])["modules"]
    assert plain.verification_report(f.pid, f.sys)["data_source"]["locked"]["reason"] == "Release"
    template = plain.save_project_as_template(f.pid, "From locked")
    assert template
    info = plain.project_removal_info(f.pid)
    assert info["locked"]["reason"] == "Release" and info["can_remove"] is False
    # A subproject under a locked project changes only the new project.
    child = plain.create_project("Child", parent_project_id=f.pid)
    assert child.parent_project_id == f.pid
    assert tree_bytes(root / f.pid) == before


def test_linking_from_another_project_into_a_locked_one_is_allowed(plain, root):
    f = Fixture(plain)
    plain.create_module(f.other, "Src", "SRC")
    src = plain.add_requirement(f.other, "src", {"object_text": "s"}).id
    plain.set_project_references(f.other, [f.pid])
    config = plain.create_link_configuration(f.other, "Traces", "src", f.sys, f.pid).id
    plain.lock_project(f.pid)
    before = tree_bytes(root / f.pid)
    plain.link_requirements(f.other, config, src, f.r1)
    assert tree_bytes(root / f.pid) == before
    # But removing the linked requirement would clean a link out of... the other way round:
    # the locked project's own requirement can't be removed at all.
    with pytest.raises(ProjectLockedError):
        plain.remove_requirement(f.pid, f.sys, f.r1)


def test_cascade_into_a_locked_referencing_project_is_refused(plain, root):
    f = Fixture(plain)
    plain.create_module(f.other, "Target", "TGT")
    target = plain.add_requirement(f.other, "target", {"object_text": "t"}).id
    plain.set_project_references(f.pid, [f.other])
    config = plain.create_link_configuration(f.pid, "Traces", f.sys, "target", f.other).id
    plain.link_requirements(f.pid, config, f.r1, target)
    plain.lock_project(f.pid)
    before = tree_bytes(root)
    with pytest.raises(ProjectLockedError):
        plain.remove_requirement(f.other, "target", target)
    with pytest.raises(ProjectLockedError):
        plain.purge_module(f.other, "target")
    with pytest.raises(ProjectLockedError):
        plain.purge_project(f.other)
    assert tree_bytes(root) == before


# --- lock: REST / WebSocket --------------------------------------------------------------------


@pytest.fixture
def client(root, monkeypatch):
    monkeypatch.setenv("REQMAN_DATA_DIR", str(root))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(root.parent / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app)


def rpc(ws, method: str, params: dict, request_id: int = 1) -> dict:
    ws.send_json({"jsonrpc": "2.0", "id": request_id, "method": method, "params": params})
    return ws.receive_json()


def _locked_fixture(client, root) -> Fixture:
    f = Fixture(ProjectService(ProjectStore(root)))
    response = client.put(f"/projects/{f.pid}/lock", json={"reason": "Release"})
    assert response.status_code == 200, response.text
    assert response.json()["locked"]["reason"] == "Release"
    return f


REST_MUTATIONS = {
    "requirement add": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.sys}/requirements",
                                  {"json": {"attributes": {"object_text": "x"}}}),
    "requirement update": lambda f: ("PUT", f"/projects/{f.pid}/modules/{f.sys}/requirements/{f.r1}",
                                     {"json": {"attributes": {"object_text": "y"}}}),
    "requirement remove": lambda f: ("DELETE", f"/projects/{f.pid}/modules/{f.sys}/requirements/{f.r2}", {}),
    "requirement move": lambda f: ("PUT", f"/projects/{f.pid}/modules/{f.sys}/requirements/{f.r2}/parent",
                                   {"json": {"parent_id": f.r1}}),
    "requirement reorder": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.sys}/requirements/{f.r1}/reorder",
                                      {"json": {"direction": "down"}}),
    "requirement copy": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.sys}/requirements/{f.r1}/copy",
                                   {"json": {}}),
    "module create": lambda f: ("POST", f"/projects/{f.pid}/modules", {"json": {"name": "N", "prefix": "NN"}}),
    "module reorder": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.tst}/reorder", {"json": {"direction": "up"}}),
    "module move": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.sys}/move", {"json": {"folder_id": f.folder}}),
    "module purge": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.tst}/purge", {}),
    "project purge": lambda f: ("POST", f"/projects/{f.pid}/purge", {}),
    "attribute add": lambda f: ("POST", f"/projects/{f.pid}/attributes", {"json": {"id": "x", "name": "X"}}),
    "attribute modules": lambda f: ("PUT", f"/projects/{f.pid}/attributes/status/modules",
                                    {"json": {"modules": [f.sys]}}),
    "row styles": lambda f: ("PUT", f"/projects/{f.pid}/row-styles", {"json": {"row_styles": []}}),
    "references": lambda f: ("PUT", f"/projects/{f.pid}/references", {"json": {"references": [f.other]}}),
    "parent": lambda f: ("PUT", f"/projects/{f.pid}/parent", {"json": {"parent_project_id": f.other}}),
    "link configuration create": lambda f: ("POST", f"/projects/{f.pid}/link-configurations",
                                            {"json": {"name": "O", "source_module_id": f.sys,
                                                      "target_module_id": f.tst}}),
    "link configuration remove": lambda f: ("DELETE", f"/projects/{f.pid}/link-configurations/{f.cfg}", {}),
    "link": lambda f: ("POST", f"/projects/{f.pid}/links",
                       {"json": {"config_id": f.cfg, "source_id": f.t2, "target_id": f.r2}}),
    "unlink": lambda f: ("DELETE", f"/projects/{f.pid}/links",
                         {"json": {"config_id": f.cfg, "source_id": f.t1, "target_id": f.r1}}),
    "view create": lambda f: ("POST", f"/projects/{f.pid}/views", {"json": {"name": "V2"}}),
    "view update": lambda f: ("PUT", f"/projects/{f.pid}/views/{f.view}", {"json": {"name": "V3"}}),
    "view widths": lambda f: ("PUT", f"/projects/{f.pid}/views/{f.view}/column-widths",
                              {"json": {"column_widths": {}}}),
    "view remove": lambda f: ("DELETE", f"/projects/{f.pid}/views/{f.view}", {}),
    "import csv": lambda f: ("POST", f"/projects/{f.pid}/modules/{f.sys}/import/csv",
                             {"content": b"ObjectText\nimported\n", "headers": {"Content-Type": "text/csv"}}),
    "folder create": lambda f: ("POST", f"/projects/{f.pid}/module-folders", {"json": {"name": "F2"}}),
    "folder rename": lambda f: ("PUT", f"/projects/{f.pid}/module-folders/{f.folder}/name", {"json": {"name": "G"}}),
    "folder delete": lambda f: ("DELETE", f"/projects/{f.pid}/module-folders/{f.folder}", {}),
    "folder move": lambda f: ("POST", f"/projects/{f.pid}/module-folders/{f.folder}/move",
                              {"json": {"folder_id": None, "index": 0}}),
    "folder reorder": lambda f: ("POST", f"/projects/{f.pid}/module-folders/{f.folder}/reorder",
                                 {"json": {"direction": "down"}}),
    "lock again": lambda f: ("PUT", f"/projects/{f.pid}/lock", {"json": {"reason": "again"}}),
}


@pytest.mark.parametrize("operation", sorted(REST_MUTATIONS))
def test_rest_mutations_are_423_while_locked_and_do_not_broadcast(client, root, operation):
    f = _locked_fixture(client, root)
    method, url, kwargs = REST_MUTATIONS[operation](f)
    before = tree_bytes(root)
    with client.websocket_connect("/ws") as watcher:
        response = client.request(method, url, **kwargs)
        assert response.status_code == 423, response.text
        assert "is locked" in response.json()["detail"]
        assert tree_bytes(root) == before
        # Nothing was broadcast: the next event is the unlock's.
        assert client.delete(f"/projects/{f.pid}/lock").status_code == 200
        assert watcher.receive_json() == {"event": "changed", "project_id": f.pid}
    # Unlocked, the same request is accepted.
    method, url, kwargs = REST_MUTATIONS[operation](f)
    assert client.request(method, url, **kwargs).status_code == 200


def test_rest_module_from_template_is_423_while_locked(client, root):
    f = _locked_fixture(client, root)
    templates = client.get("/module-templates").json()
    if not templates:
        pytest.skip("no module template available")
    response = client.post(f"/projects/{f.pid}/modules/from-template", json={"template_id": templates[0]["id"]})
    assert response.status_code == 423


WS_MUTATIONS = {
    "requirement.add": lambda f: {"project_id": f.pid, "module_id": f.sys, "attributes": {"object_text": "x"}},
    "requirement.update": lambda f: {"project_id": f.pid, "module_id": f.sys, "id": f.r1,
                                     "attributes": {"object_text": "y"}},
    "requirement.remove": lambda f: {"project_id": f.pid, "module_id": f.sys, "id": f.r2},
    "requirement.move": lambda f: {"project_id": f.pid, "module_id": f.sys, "id": f.r2, "parent_id": f.r1},
    "requirement.reorder": lambda f: {"project_id": f.pid, "module_id": f.sys, "id": f.r1, "direction": "down"},
    "requirement.copy": lambda f: {"project_id": f.pid, "module_id": f.sys, "id": f.r1},
    "requirement.import_csv": lambda f: {"project_id": f.pid, "module_id": f.sys, "csv": "ObjectText\nx\n"},
    "requirement.link": lambda f: {"project_id": f.pid, "config_id": f.cfg, "source_id": f.t2, "target_id": f.r2},
    "requirement.unlink": lambda f: {"project_id": f.pid, "config_id": f.cfg, "source_id": f.t1, "target_id": f.r1},
    "module.create": lambda f: {"project_id": f.pid, "name": "N", "prefix": "NN"},
    "module.reorder": lambda f: {"project_id": f.pid, "id": f.tst, "direction": "up"},
    "module.move": lambda f: {"project_id": f.pid, "id": f.sys, "folder_id": f.folder},
    "module.purge": lambda f: {"project_id": f.pid, "id": f.tst},
    "project.purge": lambda f: {"id": f.pid},
    "project.set_row_styles": lambda f: {"id": f.pid, "row_styles": []},
    "project.set_references": lambda f: {"id": f.pid, "references": [f.other]},
    "project.set_parent": lambda f: {"id": f.pid, "parent_project_id": f.other},
    "project.lock": lambda f: {"id": f.pid, "reason": "again"},
    "attribute.add": lambda f: {"project_id": f.pid, "id": "x", "name": "X"},
    "attribute.set_modules": lambda f: {"project_id": f.pid, "id": "status", "modules": [f.sys]},
    "link_configuration.create": lambda f: {"project_id": f.pid, "name": "O", "source_module_id": f.sys,
                                            "target_module_id": f.tst},
    "link_configuration.remove": lambda f: {"project_id": f.pid, "id": f.cfg},
    "view.create": lambda f: {"project_id": f.pid, "name": "V2"},
    "view.update": lambda f: {"project_id": f.pid, "id": f.view, "name": "V3"},
    "view.set_column_widths": lambda f: {"project_id": f.pid, "id": f.view, "column_widths": {}},
    "view.remove": lambda f: {"project_id": f.pid, "id": f.view},
    "module_folder.create": lambda f: {"project_id": f.pid, "name": "F2"},
    "module_folder.rename": lambda f: {"project_id": f.pid, "id": f.folder, "name": "G"},
    "module_folder.delete": lambda f: {"project_id": f.pid, "id": f.folder},
    "module_folder.move": lambda f: {"project_id": f.pid, "id": f.folder, "folder_id": None, "index": 0},
    "module_folder.reorder": lambda f: {"project_id": f.pid, "id": f.folder, "direction": "down"},
}


@pytest.mark.parametrize("method", sorted(WS_MUTATIONS))
def test_ws_mutations_are_refused_with_32004_while_locked_and_do_not_broadcast(client, root, method):
    f = _locked_fixture(client, root)
    before = tree_bytes(root)
    with client.websocket_connect("/ws") as ws, client.websocket_connect("/ws") as watcher:
        response = rpc(ws, method, WS_MUTATIONS[method](f))
        assert response["error"]["code"] == -32004, response
        assert "is locked" in response["error"]["message"]
        assert tree_bytes(root) == before
        unlocked = rpc(ws, "project.unlock", {"id": f.pid}, 2)
        assert unlocked["result"]["id"] == f.pid and "locked" not in unlocked["result"]
        assert watcher.receive_json() == {"event": "changed", "project_id": f.pid}
        response = rpc(ws, method, WS_MUTATIONS[method](f), 3)
        assert "result" in response, response


def test_ws_lock_and_unlock_broadcast_and_match_rest(client, root):
    pid = client.post("/projects", json={"name": "Alpha"}).json()["id"]
    with client.websocket_connect("/ws") as ws, client.websocket_connect("/ws") as watcher:
        locked = rpc(ws, "project.lock", {"id": pid, "reason": " Freeze  now "})
        assert locked["result"]["locked"]["reason"] == "Freeze now"
        assert watcher.receive_json() == {"event": "changed", "project_id": pid}
        assert client.get(f"/projects/{pid}").json()["locked"] == locked["result"]["locked"]
        assert rpc(ws, "project.lock", {"id": pid, "reason": 5}, 2)["error"]["code"] == -32004  # locked already
        unlocked = rpc(ws, "project.unlock", {"id": pid}, 3)
        assert "locked" not in unlocked["result"]
        assert watcher.receive_json() == {"event": "changed", "project_id": pid}
        assert rpc(ws, "project.lock", {"id": pid, "reason": "x" * 501}, 4)["error"]["code"] == -32602
        assert rpc(ws, "project.lock", {"id": "missing"}, 5)["error"]["code"] == -32001
    assert client.put(f"/projects/{pid}/lock", json={"reason": "x" * 501}).status_code == 400
    assert client.put("/projects/missing/lock", json={}).status_code == 404
    assert client.delete("/projects/missing/lock").status_code == 404
    # Unlocking an unlocked project is fine (and writes nothing).
    assert client.delete(f"/projects/{pid}/lock").status_code == 200


def test_rest_allowed_while_locked(client, root):
    f = _locked_fixture(client, root)
    before = tree_bytes(root / f.pid)
    for url in (
        f"/projects/{f.pid}",
        f"/projects/{f.pid}/modules",
        f"/projects/{f.pid}/modules/{f.sys}/requirements",
        f"/projects/{f.pid}/module-tree",
        f"/projects/{f.pid}/views",
        f"/projects/{f.pid}/removal",
        f"/projects/{f.pid}/export/reqif",
        f"/projects/{f.pid}/export/csv",
        f"/projects/{f.pid}/export/xlsx",
        f"/projects/{f.pid}/export/docx",
        f"/projects/{f.pid}/reports/traceability?config={f.cfg}",
        f"/projects/{f.pid}/modules/{f.sys}/reports/verification",
        f"/projects/{f.pid}/modules/{f.sys}/reports/verification/export/xlsx",
    ):
        response = client.get(url)
        assert response.status_code == 200, (url, response.text)
    response = client.post(
        f"/projects/{f.pid}/modules/{f.sys}/import/csv/preview",
        content=b"ObjectText\nx\n",
        headers={"Content-Type": "text/csv"},
    )
    assert response.status_code == 200
    with client.websocket_connect("/ws") as watcher:
        response = client.post(f"/projects/{f.pid}/save-as-template", json={"name": "From locked"})
        assert response.status_code == 200
        assert watcher.receive_json()["event"] == "templates_changed"
        child = client.post("/projects", json={"name": "Child", "parent_project_id": f.pid})
        assert child.status_code == 200
    assert tree_bytes(root / f.pid) == before
    assert client.get(f"/projects/{f.pid}").json()["locked"]["reason"] == "Release"


def test_ws_cascade_into_a_locked_project_is_32004(client, root):
    service = ProjectService(ProjectStore(root))
    a = service.create_project("A").id
    b = service.create_project("B").id
    service.create_module(a, "Srs", "SRS")
    service.create_module(b, "Sys", "SYS")
    src = service.add_requirement(a, "srs", {"object_text": "s"}).id
    target = service.add_requirement(b, "sys", {"object_text": "t"}).id
    service.set_project_references(a, [b])
    config = service.create_link_configuration(a, "Traces", "srs", "sys", b).id
    service.link_requirements(a, config, src, target)
    assert client.put(f"/projects/{a}/lock", json={"reason": "Release"}).status_code == 200
    before = tree_bytes(root)
    with client.websocket_connect("/ws") as ws:
        response = rpc(ws, "requirement.remove", {"project_id": b, "module_id": "sys", "id": target})
        assert response["error"]["code"] == -32004
        assert rpc(ws, "module.purge", {"project_id": b, "id": "sys"}, 2)["error"]["code"] == -32004
        assert rpc(ws, "project.purge", {"id": b}, 3)["error"]["code"] == -32004
    assert tree_bytes(root) == before


# --- remove: rules --------------------------------------------------------------------------------


def test_remove_deletes_the_folder_and_nothing_else(plain, root):
    keep = plain.create_project("Keep")
    plain.create_module(keep.id, "Mod", "M")
    gone = plain.create_project("Gone")
    plain.create_module(gone.id, "Mod", "M")
    plain.add_requirement(gone.id, "mod", {"object_text": "x"})
    (root / "notes.txt").write_text("outside\n", encoding="utf-8")
    (root / gone.id / "extra.bin").write_bytes(b"\x00\x01")
    before = {k: v for k, v in tree_bytes(root).items() if not k.startswith(f"{gone.id}/")}

    result = plain.remove_project(gone.id, "Gone")
    assert result == {"removed": gone.id, "name": "Gone", "leftover": None}
    assert not (root / gone.id).exists()
    assert not (root / STAGING).exists()
    assert tree_bytes(root) == before
    assert [p.id for p in plain.list_projects()] == [keep.id]
    assert [p.id for p in ProjectService(ProjectStore(root)).list_projects()] == [keep.id]
    with pytest.raises(FileNotFoundError):
        plain.get_project(gone.id)
    with pytest.raises(FileNotFoundError):
        plain.remove_project(gone.id, "Gone")
    # The id is free again.
    assert plain.create_project("Gone").id == gone.id


@pytest.mark.parametrize("confirm", ["gone", "Gon", "", None, 5, ["Gone"], "Gone!"])
def test_remove_needs_the_exact_name(plain, root, confirm):
    project = plain.create_project("Gone")
    with pytest.raises(ValueError, match="Type the project name 'Gone'"):
        plain.remove_project(project.id, confirm)
    assert (root / project.id / "project.json").is_file()


def test_remove_ignores_surrounding_whitespace_of_the_name(plain, root):
    project = plain.create_project("Gone")
    plain.remove_project(project.id, "  Gone \n")
    assert not (root / project.id).exists()


@pytest.mark.parametrize("data_dir", ["/somewhere/else", "", 5, ["x"], {"path": "x"}, "data-folder:0000000000000000"])
def test_remove_refuses_another_data_dir(plain, root, data_dir):
    project = plain.create_project("Gone")
    with pytest.raises(ConflictError, match="data folder changed"):
        plain.remove_project(project.id, "Gone", data_dir)
    assert (root / project.id / "project.json").is_file()


def test_remove_data_dir_check_runs_before_the_name_check(plain, root):
    project = plain.create_project("Gone")
    with pytest.raises(ConflictError):
        plain.remove_project(project.id, "wrong", "/elsewhere")


@pytest.mark.parametrize("form", ["path", "trailing-slash", "symlink", "token"])
def test_remove_accepts_the_current_data_dir(plain, root, tmp_path, form):
    project = plain.create_project("Gone")
    if form == "path":
        data_dir = str(root)
    elif form == "trailing-slash":
        data_dir = str(root) + os.sep
    elif form == "symlink":
        link = tmp_path / "data-link"
        link.symlink_to(root, target_is_directory=True)
        data_dir = str(link)
    else:
        data_dir = plain.data_dir_token()
    plain.remove_project(project.id, "Gone", data_dir)
    assert not (root / project.id).exists()


def test_remove_without_data_dir_is_allowed(plain, root):
    """``data_dir`` is checked only when given (docs: "when given"); null means no check."""
    project = plain.create_project("Gone")
    plain.remove_project(project.id, "Gone", None)
    assert not (root / project.id).exists()


def test_remove_is_refused_while_locked(plain, root):
    project = plain.create_project("Gone")
    plain.lock_project(project.id, "Keep it")
    with pytest.raises(ProjectLockedError, match="Unlock it to remove it; nothing was removed"):
        plain.remove_project(project.id, "Gone")
    assert (root / project.id / "project.json").is_file()
    plain.unlock_project(project.id)
    plain.remove_project(project.id, "Gone")


def test_subprojects_block_the_removal(plain, root):
    parent = plain.create_project("Parent")
    child = plain.create_project("Child", parent_project_id=parent.id)
    info = plain.project_removal_info(parent.id)
    assert info["subprojects"] == [{"id": child.id, "name": "Child"}]
    assert info["referenced_by"] == [] and info["can_remove"] is False
    with pytest.raises(ConflictError, match=r"it has subprojects 'Child' \(child\)") as excinfo:
        plain.remove_project(parent.id, "Parent")
    assert not isinstance(excinfo.value, ProjectLockedError)
    assert "Move or remove the subprojects first; nothing was removed" in str(excinfo.value)
    assert (root / parent.id).is_dir()
    # The child itself can go, then the parent.
    plain.remove_project(child.id, "Child")
    plain.remove_project(parent.id, "Parent")


def test_referencing_projects_block_the_removal(plain, root):
    target = plain.create_project("Target")
    user = plain.create_project("User")
    plain.set_project_references(user.id, [target.id])
    info = plain.project_removal_info(target.id)
    assert info["referenced_by"] == [{"id": user.id, "name": "User"}] and info["can_remove"] is False
    with pytest.raises(ConflictError, match=r"it is referenced by 'User' \(user\)") as excinfo:
        plain.remove_project(target.id, "Target")
    assert "Remove those references first" in str(excinfo.value)
    plain.set_project_references(user.id, [])
    assert plain.project_removal_info(target.id)["can_remove"] is True
    plain.remove_project(target.id, "Target")


def test_subprojects_and_references_are_both_listed(plain):
    target = plain.create_project("Target")
    child = plain.create_project("Child", parent_project_id=target.id)
    user = plain.create_project("User")
    plain.set_project_references(user.id, [target.id])
    with pytest.raises(ConflictError) as excinfo:
        plain.remove_project(target.id, "Target")
    message = str(excinfo.value)
    assert f"it has subprojects 'Child' ({child.id}) and it is referenced by 'User' ({user.id})" in message
    assert "Move or remove the subprojects and remove those references first" in message


@pytest.mark.parametrize("breakage", ["newer-format", "bad-attributes"])
def test_broken_projects_still_block_the_removal(plain, root, breakage):
    target = plain.create_project("Target")
    child = plain.create_project("Child", parent_project_id=target.id)
    user = plain.create_project("User")
    plain.set_project_references(user.id, [target.id])
    for pid in (child.id, user.id):
        if breakage == "newer-format":
            _rewrite(root, pid, format_version=PROJECT_FORMAT.current + 1)
        else:
            _rewrite(root, pid, attributes="not a list", id_counters=[1])
    fresh = ProjectService(ProjectStore(root))
    info = fresh.project_removal_info(target.id)
    assert [p["id"] for p in info["subprojects"]] == [child.id]
    assert [p["id"] for p in info["referenced_by"]] == [user.id]
    assert info["can_remove"] is False
    with pytest.raises(ConflictError):
        fresh.remove_project(target.id, "Target")
    assert (root / target.id).is_dir()


def test_unparseable_projects_do_not_block_the_removal(plain, root):
    target = plain.create_project("Target")
    junk = plain.create_project("Junk")
    project_file(root, junk.id).write_text("{ not json", encoding="utf-8")
    info = ProjectService(ProjectStore(root)).project_removal_info(target.id)
    assert info["subprojects"] == [] and info["referenced_by"] == [] and info["can_remove"] is True


def test_symlinked_project_folder_is_refused(plain, root, tmp_path):
    outside = tmp_path / "outside"
    outside.mkdir()
    real = plain.create_project("Real")
    os.rename(root / real.id, outside / "linked")
    (root / "linked").symlink_to(outside / "linked", target_is_directory=True)
    _rewrite(root, "linked", id="linked", name="Linked")
    fresh = ProjectService(ProjectStore(root))
    assert fresh.get_project("linked").name == "Linked"
    before = tree_bytes(outside)
    with pytest.raises(ValueError, match="not a plain folder"):
        fresh.remove_project("linked", "Linked")
    assert tree_bytes(outside) == before
    assert (root / "linked").is_symlink()
    assert not (root / STAGING).exists()


def test_existing_staging_contents_are_left_alone(plain, root):
    project = plain.create_project("Gone")
    other = root / STAGING / "in-progress"
    other.mkdir(parents=True)
    (other / "file.json").write_text("{}\n", encoding="utf-8")
    assert plain.remove_project(project.id, "Gone")["leftover"] is None
    assert not (root / project.id).exists()
    assert (other / "file.json").read_text(encoding="utf-8") == "{}\n"
    assert [p.name for p in (root / STAGING).iterdir()] == ["in-progress"]
    # The staging folder never shows up as a project.
    assert [p.id for p in plain.list_projects()] == []


def test_leftover_is_reported_when_the_folder_cannot_be_deleted(plain, root, monkeypatch):
    import reqman.store as store_module

    project = plain.create_project("Gone")
    real_rmtree = store_module.shutil.rmtree

    def failing_rmtree(path, *args, **kwargs):
        if STAGING in str(path):
            raise OSError("busy")
        return real_rmtree(path, *args, **kwargs)

    monkeypatch.setattr(store_module.shutil, "rmtree", failing_rmtree)
    result = plain.remove_project(project.id, "Gone")
    monkeypatch.undo()
    leftover = Path(result["leftover"])
    assert leftover.parent == root / STAGING and leftover.name.startswith("removed-")
    assert (leftover / "project.json").is_file()
    assert not (root / project.id).exists()
    assert [p.id for p in plain.list_projects()] == []


# --- remove: REST / WebSocket ---------------------------------------------------------------------


def test_rest_remove_broadcasts_removed(client, root):
    pid = client.post("/projects", json={"name": "Gone"}).json()["id"]
    other = client.post("/projects", json={"name": "Keep"}).json()["id"]
    with client.websocket_connect("/ws") as watcher:
        assert client.request("DELETE", f"/projects/{pid}", json={"confirm_name": "nope"}).status_code == 400
        assert client.request(
            "DELETE", f"/projects/{pid}", json={"confirm_name": "Gone", "data_dir": "/elsewhere"}
        ).status_code == 409
        assert client.request("DELETE", f"/projects/{pid}", json={"confirm_name": "Gone", "data_dir": 5}).status_code == 409
        assert client.request("DELETE", "/projects/missing", json={"confirm_name": "x"}).status_code == 404
        response = client.request("DELETE", f"/projects/{pid}", json={"confirm_name": "Gone", "data_dir": str(root)})
        assert response.status_code == 200, response.text
        assert response.json() == {"removed": pid, "name": "Gone", "leftover": None}
        # The failed attempts broadcast nothing; the first event is the removal.
        assert watcher.receive_json() == {"event": "changed", "project_id": pid, "removed": True}
    assert [p["id"] for p in client.get("/projects").json()] == [other]
    assert client.get(f"/projects/{pid}").status_code == 404


def test_rest_remove_status_codes_for_blockers(client, root):
    parent = client.post("/projects", json={"name": "Parent"}).json()["id"]
    client.post("/projects", json={"name": "Child", "parent_project_id": parent})
    response = client.request("DELETE", f"/projects/{parent}", json={"confirm_name": "Parent"})
    assert response.status_code == 409 and "Child" in response.json()["detail"]
    info = client.get(f"/projects/{parent}/removal").json()
    assert info["can_remove"] is False and info["history"] is None
    assert info["name"] == "Parent" and info["project_id"] == parent and info["locked"] is None

    locked = client.post("/projects", json={"name": "Locked"}).json()["id"]
    client.put(f"/projects/{locked}/lock", json={"reason": "Keep"})
    response = client.request("DELETE", f"/projects/{locked}", json={"confirm_name": "Locked"})
    assert response.status_code == 423
    assert client.get(f"/projects/{locked}/removal").json()["can_remove"] is False
    assert client.get("/projects/missing/removal").status_code == 404


def test_ws_remove_broadcasts_removed_to_others_and_matches_rest(client, root):
    pid = client.post("/projects", json={"name": "Gone"}).json()["id"]
    with client.websocket_connect("/ws") as ws, client.websocket_connect("/ws") as watcher:
        info = rpc(ws, "project.removal_info", {"id": pid})["result"]
        assert info == client.get(f"/projects/{pid}/removal").json()
        assert info["can_remove"] is True and info["subprojects"] == [] and info["referenced_by"] == []
        assert rpc(ws, "project.remove", {"id": pid, "confirm_name": "x"}, 2)["error"]["code"] == -32602
        response = rpc(ws, "project.remove", {"id": pid, "confirm_name": "Gone", "data_dir": "/x"}, 3)
        assert response["error"]["code"] == -32602
        response = rpc(ws, "project.remove", {"id": pid, "confirm_name": "Gone", "data_dir": str(root)}, 4)
        assert response["result"] == {"removed": pid, "name": "Gone", "leftover": None}
        assert watcher.receive_json() == {"event": "changed", "project_id": pid, "removed": True}
        assert rpc(ws, "project.remove", {"id": pid, "confirm_name": "Gone"}, 5)["error"]["code"] == -32001
    assert not (root / pid).exists()


def test_ws_remove_locked_is_32004(client, root):
    pid = client.post("/projects", json={"name": "Locked"}).json()["id"]
    client.put(f"/projects/{pid}/lock", json={})
    with client.websocket_connect("/ws") as ws:
        assert rpc(ws, "project.remove", {"id": pid, "confirm_name": "Locked"})["error"]["code"] == -32004
    assert (root / pid).is_dir()


# --- remove: git state and recovery hints -----------------------------------------------------------


@pytest.fixture
def bare(tmp_path):
    remote = tmp_path / "remote.git"
    subprocess.run(["git", "init", "--bare", "-q", str(remote)], check=True, capture_output=True)
    return remote


def test_recovery_hint_with_history_off(plain):
    project = plain.create_project("Alpha")
    info = plain.project_removal_info(project.id)
    assert info["history"] is None
    assert "no git history of its own" in info["recovery"]


def test_recovery_hint_without_remote(gitsvc):
    project = gitsvc.create_project("Alpha")
    info = gitsvc.project_removal_info(project.id)
    assert info["history"]["state"] == "ready"
    assert "It has no remote" in info["recovery"] and "including its baselines" in info["recovery"]


def test_recovery_hint_with_a_pushed_remote(gitsvc, root, bare):
    project = gitsvc.create_project("Alpha")
    gitsvc.set_git_remote(project.id, str(bare))
    gitsvc.git_push(project.id)
    info = gitsvc.project_removal_info(project.id)
    assert f"A copy remains at its remote {bare}" in info["recovery"]
    assert info["recovery"].endswith("Everything is pushed.")

    # Unpushed commits and uncommitted changes are named.
    gitsvc.create_module(project.id, "Mod", "M")
    path = project_file(root, project.id)
    data = json.loads(path.read_text(encoding="utf-8"))
    data["description"] = "edited by hand"
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")
    info = gitsvc.project_removal_info(project.id)
    assert "Not on the remote, and lost: uncommitted changes, 1 unpushed commit(s) — push first." in info["recovery"]


def test_recovery_hint_with_a_never_fetched_remote(gitsvc, bare):
    project = gitsvc.create_project("Alpha")
    gitsvc.set_git_remote(project.id, str(bare))
    info = gitsvc.project_removal_info(project.id)
    assert "possibly unpushed commits (the remote was never fetched)" in info["recovery"]


def test_recovery_hint_inside_an_enclosing_repository(tmp_path):
    outer = tmp_path / "outer"
    outer.mkdir()
    git(outer, "init", "-q", "-b", "main")
    service = make_git_service(outer / "data")
    project = service.create_project("Inner")
    info = service.project_removal_info(project.id)
    assert info["history"]["state"] == "nested"
    assert f"inside the git repository at '{outer.resolve()}'" in info["recovery"]
    # Removing it touches nothing of the enclosing repository but the folder.
    service.remove_project(project.id, "Inner")
    assert not (outer / "data" / project.id).exists()
    assert (outer / ".git").is_dir()


def test_remove_with_history_deletes_its_repository_but_not_the_remote(gitsvc, root, bare):
    project = gitsvc.create_project("Alpha")
    gitsvc.set_git_remote(project.id, str(bare))
    gitsvc.create_baseline(project.id, "v1")
    gitsvc.git_push(project.id)
    remote_refs = git(bare, "for-each-ref", "--format=%(refname)")
    gitsvc.remove_project(project.id, "Alpha")
    assert not (root / project.id).exists()
    assert git(bare, "for-each-ref", "--format=%(refname)") == remote_refs
    assert "refs/tags/baseline/v1" in remote_refs
    # A new project with the same id gets a fresh repository.
    again = gitsvc.create_project("Alpha")
    assert subjects(root / again.id) == ["Create project 'Alpha'"]


# --- baseline glue ----------------------------------------------------------------------------------


def _lab(service: ProjectService) -> tuple[str, str]:
    pid = service.create_project("Lab").id
    mid = service.create_module(pid, "Reqs", "REQ").id
    service.add_requirement(pid, mid, {"object_text": "r"})
    return pid, mid


def test_baseline_with_lock_locks_first_then_tags(gitsvc, root):
    pid, _mid = _lab(gitsvc)
    folder = root / pid
    result = gitsvc.create_baseline(pid, "v1", "First", lock=True)
    assert result["name"] == "v1"
    assert result["locked"]["reason"] == "Baseline v1" and result["locked"]["baseline"] == "v1"
    assert result["locked"]["by"] == "Test Author"
    lock = gitsvc.get_project(pid).locked
    assert lock.to_dict() == result["locked"]
    tagged = git(folder, "rev-parse", "baseline/v1^{commit}").strip()
    assert tagged == git(folder, "rev-parse", "HEAD").strip()
    assert subjects(folder)[0] == "Project locked: Baseline v1"
    # The baseline is exactly the locked state.
    snapshot = gitsvc.get_baseline(pid, "v1")
    assert snapshot["project"]["locked"] == result["locked"]
    assert git(folder, "status", "--porcelain") == ""


def test_baseline_lock_with_custom_reason(gitsvc):
    pid, _mid = _lab(gitsvc)
    result = gitsvc.create_baseline(pid, "rel-2", lock=True, lock_reason="  Frozen   for audit ")
    assert result["locked"]["reason"] == "Frozen for audit" and result["locked"]["baseline"] == "rel-2"


def test_baseline_without_lock_reports_locked_null(gitsvc):
    pid, _mid = _lab(gitsvc)
    assert gitsvc.create_baseline(pid, "v1")["locked"] is None
    assert gitsvc.get_project(pid).locked is None


def test_failed_baseline_undoes_the_lock_duplicate_name(gitsvc, root):
    pid, _mid = _lab(gitsvc)
    gitsvc.create_baseline(pid, "v1")
    with pytest.raises(ConflictError):
        gitsvc.create_baseline(pid, "v1", lock=True)
    assert gitsvc.get_project(pid).locked is None
    assert "locked" not in json.loads(project_file(root, pid).read_text(encoding="utf-8"))
    # Still editable.
    gitsvc.create_module(pid, "More", "MORE")


def test_failed_baseline_undoes_the_lock_history_off(plain, root):
    pid, _mid = _lab(plain)
    before = project_file(root, pid).read_bytes()
    with pytest.raises(ValueError):
        plain.create_baseline(pid, "v1", lock=True)
    assert plain.get_project(pid).locked is None
    assert json.loads(project_file(root, pid).read_bytes()) == json.loads(before)


@pytest.mark.parametrize("name", ["", "a b", "current", "-x", "x" * 101])
def test_invalid_baseline_name_with_lock_never_locks(gitsvc, root, name):
    pid, _mid = _lab(gitsvc)
    before = project_file(root, pid).read_bytes()
    with pytest.raises(ValueError):
        gitsvc.create_baseline(pid, name, lock=True)
    assert project_file(root, pid).read_bytes() == before


def test_invalid_lock_reason_never_locks(gitsvc, root):
    pid, _mid = _lab(gitsvc)
    with pytest.raises(ValueError, match="too long"):
        gitsvc.create_baseline(pid, "v1", lock=True, lock_reason="x" * 501)
    assert gitsvc.get_project(pid).locked is None
    assert gitsvc.list_baselines(pid)["baselines"] == []


def test_baseline_of_an_already_locked_project_keeps_its_lock(gitsvc, root):
    pid, _mid = _lab(gitsvc)
    gitsvc.lock_project(pid, "Manual")
    lock = gitsvc.get_project(pid).locked.to_dict()
    result = gitsvc.create_baseline(pid, "v1", lock=True)
    assert result["locked"] == lock and "baseline" not in result["locked"]
    result = gitsvc.create_baseline(pid, "v2")
    assert result["locked"] == lock
    assert sorted(b["name"] for b in gitsvc.list_baselines(pid)["baselines"]) == ["v1", "v2"]
    # A failing baseline of an already locked project leaves the lock alone.
    with pytest.raises(ConflictError):
        gitsvc.create_baseline(pid, "v1", lock=True)
    assert gitsvc.get_project(pid).locked.to_dict() == lock


def test_rest_and_ws_baseline_with_lock(root, monkeypatch):
    monkeypatch.setenv("REQMAN_GIT", "on")
    monkeypatch.setenv("REQMAN_GIT_SQUASH_MINUTES", "0")
    monkeypatch.setenv("REQMAN_DATA_DIR", str(root))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(root.parent / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    client = TestClient(app_module.app)
    a = client.post("/projects", json={"name": "A"}).json()["id"]
    b = client.post("/projects", json={"name": "B"}).json()["id"]
    with client.websocket_connect("/ws") as watcher:
        response = client.post(f"/projects/{a}/baselines", json={"name": "v1", "lock": True, "lock_reason": "R"})
        assert response.status_code == 200, response.text
        assert response.json()["locked"]["reason"] == "R" and response.json()["locked"]["baseline"] == "v1"
        assert watcher.receive_json() == {"event": "changed", "project_id": a}
        assert client.post(f"/projects/{a}/modules", json={"name": "M", "prefix": "M"}).status_code == 423
        # Duplicate: 409, and an already locked project stays locked.
        assert client.post(f"/projects/{a}/baselines", json={"name": "v1", "lock": True}).status_code == 409
        assert client.get(f"/projects/{a}").json()["locked"]["reason"] == "R"

    with client.websocket_connect("/ws") as ws, client.websocket_connect("/ws") as watcher:
        result = rpc(ws, "baseline.create", {"project_id": b, "name": "w1", "lock": True})["result"]
        assert result["locked"]["reason"] == "Baseline w1"
        assert watcher.receive_json() == {"event": "changed", "project_id": b}
        response = rpc(ws, "baseline.create", {"project_id": b, "name": "w2"}, 2)
        assert response["result"]["locked"]["reason"] == "Baseline w1"
        assert watcher.receive_json() == {"event": "changed", "project_id": b}
    # Removal info of a locked project.
    info = client.get(f"/projects/{a}/removal").json()
    assert info["can_remove"] is False and info["history"]["state"] == "ready"


# --- report data source -------------------------------------------------------------------------------


def _xlsx_values(content: bytes) -> list[list]:
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    rows = []
    for sheet in workbook.worksheets:
        rows += [[c for c in row if c is not None] for row in sheet.iter_rows(values_only=True)]
    return rows


def _docx_text(content: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(content)).paragraphs]


def _test_lab(service: ProjectService) -> tuple[str, str, str]:
    pid = service.create_project("Lab").id
    service.add_attribute_definition(
        pid, "verdict", "Verdict", type="enum", choices=["Pass", "Fail"], module_types=["test_execution"]
    )
    reqs = service.create_module(pid, "Reqs", "REQ").id
    execs = service.create_module(pid, "Exec", "EX", type="test_execution").id
    service.add_requirement(pid, reqs, {"object_text": "r"})
    return pid, reqs, execs


def test_report_data_source_after_a_locked_baseline(gitsvc, root):
    pid, reqs, execs = _test_lab(gitsvc)
    folder = root / pid
    gitsvc.create_baseline(pid, "v1", lock=True)
    commit = git(folder, "rev-parse", "HEAD").strip()
    source = gitsvc.verification_report(pid, reqs)["data_source"]
    lock = gitsvc.get_project(pid).locked.to_dict()
    assert source == {
        "history": "ready",
        "commit": commit,
        "pending_changes": 0,
        "baselines": ["v1"],
        "latest_baseline": {"name": "v1", "commit": commit},
        "locked": lock,
    }
    assert gitsvc.execution_report(pid, execs)["data_source"] == source

    expected = f"Current data at commit {commit[:12]} (baseline v1); project locked by Test Author"
    assert ["Data", expected] in _xlsx_values(gitsvc.export_verification_report(pid, reqs, "xlsx"))
    assert f"Data: {expected}" in _docx_text(gitsvc.export_verification_report(pid, reqs, "docx"))
    assert ["Data", expected] in _xlsx_values(gitsvc.export_execution_report(pid, execs, "xlsx"))
    assert f"Data: {expected}" in _docx_text(gitsvc.export_execution_report(pid, execs, "docx"))


def test_report_data_source_after_later_commits_and_pending_changes(gitsvc, root):
    pid, reqs, _execs = _test_lab(gitsvc)
    folder = root / pid
    gitsvc.create_baseline(pid, "v1")
    tagged = git(folder, "rev-parse", "HEAD").strip()
    gitsvc.add_requirement(pid, reqs, {"object_text": "later"})
    head = git(folder, "rev-parse", "HEAD").strip()
    path = project_file(root, pid)
    data = json.loads(path.read_text(encoding="utf-8"))
    data["description"] = "pending"
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")

    source = gitsvc.verification_report(pid, reqs)["data_source"]
    assert source["commit"] == head and source["baselines"] == []
    assert source["pending_changes"] == 1
    assert source["latest_baseline"] == {"name": "v1", "commit": tagged}
    assert source["locked"] is None
    expected = (
        f"Current data at commit {head[:12]} plus 1 uncommitted change(s); latest baseline: v1 ({tagged[:12]})"
    )
    assert ["Data", expected] in _xlsx_values(gitsvc.export_verification_report(pid, reqs, "xlsx"))
    assert f"Data: {expected}" in _docx_text(gitsvc.export_verification_report(pid, reqs, "docx"))


def test_report_data_source_text_without_history(plain):
    pid, reqs, _execs = _test_lab(plain)
    plain.lock_project(pid, "Audit")
    expected = f"Current data (no git history); project locked by {LOGIN}"
    assert ["Data", expected] in _xlsx_values(plain.export_verification_report(pid, reqs, "xlsx"))
    assert f"Data: {expected}" in _docx_text(plain.export_verification_report(pid, reqs, "docx"))


# --- remove: review fixes -------------------------------------------------------------------------

posix_non_root = pytest.mark.skipif(
    sys.platform == "win32" or (hasattr(os, "geteuid") and os.geteuid() == 0),
    reason="needs POSIX permissions that apply (not root)",
)


@posix_non_root
def test_remove_with_an_unreadable_subfolder_still_succeeds_and_reports_the_leftover(plain, root):
    """rmtree fails on os.open of a mode 000 folder: not retried (no TypeError), the removal still
    succeeds (the folder is already out of the data directory) and reports what was left behind."""
    project = plain.create_project("Gone")
    locked_dir = root / project.id / "sealed"
    locked_dir.mkdir()
    (locked_dir / "inner.txt").write_text("x\n", encoding="utf-8")
    locked_dir.chmod(0o000)
    leftover = None
    try:
        result = plain.remove_project(project.id, "Gone")
        leftover = result["leftover"]
        assert result["removed"] == project.id and result["name"] == "Gone"
        assert leftover is not None and Path(leftover).parent == root / STAGING
        assert not (root / project.id).exists()
        assert [p.id for p in plain.list_projects()] == []
    finally:
        for folder in [locked_dir, *([Path(leftover) / "sealed"] if leftover else [])]:
            if folder.exists() or folder.is_symlink():
                folder.chmod(0o700)


@posix_non_root
def test_remove_never_changes_the_mode_of_a_symlink_target_outside_the_project(plain, root, tmp_path):
    project = plain.create_project("Gone")
    outside = tmp_path / "outside.txt"
    outside.write_text("keep\n", encoding="utf-8")
    outside.chmod(0o400)
    read_only = root / project.id / "ro"
    read_only.mkdir()
    (read_only / "link").symlink_to(outside)
    read_only.chmod(0o500)  # its entries cannot be unlinked until the folder is made writable
    try:
        result = plain.remove_project(project.id, "Gone")
        assert result["leftover"] is None
        assert not (root / project.id).exists()
        assert outside.read_text(encoding="utf-8") == "keep\n"
        assert outside.stat().st_mode & 0o777 == 0o400
    finally:
        outside.chmod(0o600)
        if read_only.exists():
            read_only.chmod(0o700)


def test_a_differently_cased_id_cannot_bypass_the_removal_blockers(plain, root):
    parent = plain.create_project("Alpha")
    plain.create_project("Child", parent_project_id=parent.id)
    user = plain.create_project("User")
    plain.set_project_references(user.id, [parent.id])
    other_case = parent.id.upper()
    if not (root / other_case).exists():
        pytest.skip("the temporary filesystem is case-sensitive")
    with pytest.raises(FileNotFoundError):
        plain.project_removal_info(other_case)
    with pytest.raises(FileNotFoundError):
        plain.remove_project(other_case, "Alpha")
    assert (root / parent.id / "project.json").is_file()


def test_lock_of_a_locked_project_is_423(plain):
    """lock_project docstring: already locked is refused as locked (423), not 409."""
    project = plain.create_project("Locked")
    plain.lock_project(project.id, "Release")
    with pytest.raises(ProjectLockedError):
        plain.lock_project(project.id, "Again")
