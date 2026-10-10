"""Rich text values of ``text`` attributes: sanitiser, projection, service write paths, project format."""

import importlib
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from fastapi.testclient import TestClient

from reqman.formats import PROJECT_FORMAT, upgrade
from reqman.history import format_value
from reqman.richtext import RICH_TEXT_PREFIX as P, from_plain, is_rich, sanitize, to_plain
from reqman.service import ProjectService
from reqman.store import ProjectStore

END = "</div>"


def rich(body: str) -> str:
    return P + body + END


# --- sanitiser -------------------------------------------------------------------


@pytest.mark.parametrize(
    "value",
    ["", "plain", "a < b && c > d", "<p>not rich</p>", "<script>alert(1)</script>", " " + P + "<p>x</p></div>"],
)
def test_plain_values_pass_through_unchanged(value):
    assert not is_rich(value)
    assert sanitize(value) == value
    assert to_plain(value) == value


@pytest.mark.parametrize(
    "dirty, clean",
    [
        ("<p>a<script>alert(1)</script>b</p>", "<p>a</p>"),  # demoted below; see the next test
        ('<p><strong onclick="x()">x</strong><img src=x onerror="alert(1)"></p>', "<p><strong>x</strong></p>"),
        ('<p><a href="javascript:alert(1)">link</a> <b>t</b></p>', "<p>link <strong>t</strong></p>"),
        ('<p><span style="color: red; background: url(javascript:x)">c</span><u>u</u></p>', "<p>c<u>u</u></p>"),
        ('<p><span style="color:expression(alert(1))">c</span><em>e</em></p>', "<p>c<em>e</em></p>"),
        ("<style>p{}</style><iframe src=x>i</iframe><p><s>s</s></p><!-- c -->", "<p><s>s</s></p>"),
        ('<svg onload="alert(1)"><text>t</text></svg><p><i>i</i></p>', "<p><em>i</em></p>"),
        ("<p><del>d</del><strike>s</strike></p>", "<p><s>d</s><s>s</s></p>"),
    ],
)
def test_xss_vectors_are_stripped(dirty, clean):
    result = sanitize(rich(dirty))
    if "<strong>" in clean or "<u>" in clean or "<em>" in clean or "<s>" in clean:
        assert result == rich(clean)
    for bad in ("script", "onclick", "onerror", "onload", "javascript", "href", "iframe", "svg", "style>"):
        assert bad not in result


def test_unformatted_rich_text_is_stored_as_plain_text():
    assert sanitize(rich("<p>a<script>alert(1)</script>b</p>")) == "ab"
    assert sanitize(rich("<p>one</p><p></p><p>two</p>")) == "one\n\ntwo"
    assert sanitize(rich("<p> </p>")) == ""
    assert sanitize(rich("<ul><li></li></ul>")) == ""


def test_styles_are_canonical():
    value = rich(
        '<p><span style="FONT-SIZE: 14px; color: RGB(255, 0, 16); position: fixed">a</span>'
        '<span style="color:#ABC;font-size:13px">b</span><span style="font-size: 99px">c</span></p>'
    )
    assert sanitize(value) == rich(
        '<p><span style="color: #ff0010; font-size: 14px">a</span><span style="color: #aabbcc">b</span>c</p>'
    )


def test_nesting_lists_and_entities():
    value = rich(
        "\n<ul>\n <li>one &amp; <b>two</b><ol><li>x &lt;y&gt;</li></ol></li>loose</ul>"
        "text<br>more<p>a<p>b</p></p>&nbsp;<p>&nbsp;</p>"
    )
    result = sanitize(value)
    # As in HTML (the browser mirror): <p> closes an open <p>, so "a" and "b"
    # are separate paragraphs, and the stray </p> after them is an empty one.
    assert result == rich(
        "<ul><li><p>one &amp; <strong>two</strong></p><ol><li><p>x &lt;y&gt;</p></li></ol></li>"
        "<li><p>loose</p></li></ul><p>text<br/>more</p><p>a</p><p>b</p><p></p><p>\u00a0</p>"
    )
    assert to_plain(result) == "- one & two\n  1. x <y>\n- loose\ntext\nmore\na\nb\n\n\u00a0"


@pytest.mark.parametrize(
    "body, expected",
    [
        ("<ul><li>a<li>b</ul>", rich("<ul><li><p>a</p></li><li><p>b</p></li></ul>")),
        # A formatting element closed implicitly is reopened, as by the browser.
        (
            "<ol><li><strong>a<li>b</strong></ol>",
            rich("<ol><li><p><strong>a</strong></p></li><li><p><strong>b</strong></p></li></ol>"),
        ),
        # An <li> inside a nested list does not close the outer item.
        (
            "<ul><li>a<ul><li>b<li>c</ul><li>d</ul>",
            rich("<ul><li><p>a</p><ul><li><p>b</p></li><li><p>c</p></li></ul></li><li><p>d</p></li></ul>"),
        ),
        ("<p>a<p>b", "a\nb"),
        ("<p><em>a</em><p>b", rich("<p><em>a</em></p><p>b</p>")),
        ("<p><em>a<ul><li>b</ul>", rich("<p><em>a</em></p><ul><li><p><em>b</em></p></li></ul>")),
        ("<p><em>a</p><p>b</p>", rich("<p><em>a</em></p><p><em>b</em></p>")),
        ("<p><del>a<p>b", rich("<p><s>a</s></p><p>b</p>")),  # del is not an HTML formatting element
        ("<p><em>a</em><div>b</div>", rich("<p><em>a</em></p><p>b</p>")),
        ("<ul><li><p><em>a</em><li>b</ul>", rich("<ul><li><p><em>a</em></p></li><li><p>b</p></li></ul>")),
        ("a</p>b", "a\n\nb"),
        # Formatting elements equal but for (dropped) attributes count as equal entries.
        ("<p><b a=1><b a=2><b a=3><b a=4>x<p>y", rich("<p><strong><strong><strong><strong>x</strong></strong></strong></strong></p><p><strong><strong><strong>y</strong></strong></strong></p>")),
        # The self-closing flag of a non-void element is ignored, as by the browser.
        ("<p><b/>x</p>", rich("<p><strong>x</strong></p>")),
        ("<p/>a<p/>b", "a\nb"),
        ("<ul><li/>a<li/>b</ul>", rich("<ul><li><p>a</p></li><li><p>b</p></li></ul>")),
    ],
)
def test_implied_end_tags_match_html(body, expected):
    result = sanitize(rich(body))
    assert result == expected
    assert sanitize(result) == result


@pytest.mark.parametrize(
    "value",
    [
        rich("<p><b><i>x</i></b> <span style='color:#000000'>y</span></p><ol><li>a<ul><li>b</li></ul></li></ol>"),
        rich("<span style='color:#123456'>" * 500 + "deep"),
        rich("<ul><li>" * 300 + "deep"),
        rich("<p>a\tb\r\nc</p>\x01<p>&lt;div&gt;</p>"),
        rich("<p>" + P + "</p>"),
    ],
)
def test_sanitising_is_idempotent(value):
    once = sanitize(value)
    assert sanitize(once) == once


def test_text_that_starts_with_the_prefix_stays_literal():
    literal = P + "<p>x</p>"
    stored = from_plain(literal)
    assert is_rich(stored) and sanitize(stored) == stored
    assert to_plain(stored) == literal
    assert sanitize(rich("<p>" + P.replace("<", "&lt;") + "</p>")) == rich("<p>" + P.replace("<", "&lt;").replace(">", "&gt;") + "</p>")
    assert from_plain("a <b>") == "a <b>"


def test_plain_projection_and_commit_messages():
    value = sanitize(rich("<p><strong>Shall</strong> work</p><ul><li>a</li><li>b<br>c</li></ul>"))
    assert to_plain(value) == "Shall work\n- a\n- b\n  c"
    assert format_value(value) == "Shall work - a - b c"


# --- service write paths -------------------------------------------------------------


@pytest.fixture
def service(tmp_path):
    return ProjectService(ProjectStore(tmp_path))


@pytest.fixture
def module(service):
    project = service.create_project("P")
    service.add_attribute_definition(project.id, "note", "Note", "text")
    service.add_attribute_definition(project.id, "title", "Title", "string")
    return project.id, service.create_module(project.id, "M", "M").id


DIRTY = rich('<p><b onclick="x">Bold</b><script>alert(1)</script></p>')
CLEAN = rich("<p><strong>Bold</strong></p>")


def test_add_update_and_copy_sanitise_text_values_only(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": DIRTY, "note": DIRTY, "title": DIRTY})
    assert req.attributes["object_text"] == CLEAN
    assert req.attributes["note"] == CLEAN
    assert req.attributes["title"] == DIRTY  # string attributes are not rich text
    saved = service.get_requirement(project_id, module_id, req.id)
    assert saved.attributes["object_text"] == CLEAN

    updated = service.update_requirement(project_id, module_id, req.id, {"note": rich("<p><u>u</u></p><x-y>")})
    assert updated.attributes["note"] == rich("<p><u>u</u></p>")

    # A hand-edited file holding unsanitised markup is cleaned when copied.
    path = service.store.root / project_id / "modules" / module_id / "requirements" / f"{req.id}.json"
    data = json.loads(path.read_text())
    data["attributes"]["object_text"] = DIRTY
    path.write_text(json.dumps(data, indent=2) + "\n")
    copy = service.copy_requirement(project_id, module_id, req.id, None)
    assert copy.attributes["object_text"] == CLEAN


def test_required_text_with_no_visible_text_is_blank(service, module):
    project_id, module_id = module
    with pytest.raises(ValueError, match="required attribute 'ObjectText' is empty"):
        service.add_requirement(project_id, module_id, {"object_text": rich("<ul><li><b> </b></li></ul>")})


def test_text_default_is_sanitised(service):
    project = service.create_project("D")
    attr = service.add_attribute_definition(project.id, "note", "Note", "text", default=DIRTY).attributes[-1]
    assert attr.default == CLEAN


def test_csv_import_keeps_markup_like_cells_literal(service, module):
    project_id, module_id = module
    literal = P + "<p>x</p>" + END
    csv_text = 'ObjectText,Note\n"' + literal.replace('"', '""') + '","<b>plain</b>"\n'
    result = service.import_csv(project_id, module_id, csv_text)
    req = service.get_requirement(project_id, module_id, result["requirement_ids"][0])
    assert req.attributes["note"] == "<b>plain</b>"
    assert is_rich(req.attributes["object_text"])
    assert to_plain(req.attributes["object_text"]) == literal


def test_table_exports_write_the_plain_projection(service, module):
    """CSV and XLSX stay plain; DOCX and ReqIF keep formatting (``test_rich_text_exports.py``)."""
    from docx import Document
    from openpyxl import load_workbook

    project_id, module_id = module
    service.add_requirement(project_id, module_id, {"object_text": rich("<p><em>Shall</em> run</p><ol><li>a</li></ol>")})
    text = service.export_csv(project_id)
    assert "Shall run" in text and "1. a" in text and "<em>" not in text and "&lt;em" not in text
    cells = [c for row in load_workbook(io.BytesIO(service.export_xlsx(project_id))).active.iter_rows(values_only=True) for c in row]
    assert "Shall run\n1. a" in cells
    paragraphs = "\n".join(p.text for p in Document(io.BytesIO(service.export_docx(project_id))).paragraphs)
    assert "Shall run" in paragraphs and "<em>" not in paragraphs
    # ReqIF writes the formatting as XHTML elements, not escaped markup.
    reqif = service.export_reqif(project_id)
    assert "<xhtml:p><xhtml:em>Shall</xhtml:em> run</xhtml:p><xhtml:ol><xhtml:li><xhtml:p>a</xhtml:p></xhtml:li></xhtml:ol>" in reqif
    assert "&lt;em" not in reqif


# --- project format: rich text needs format 4 or newer -----------------------------------


def test_v3_to_v4_migration_keeps_every_value():
    # Plain values, markup-like text included, mean the same in format 4: nothing is rewritten.
    data = {"id": "p", "attributes": [{"id": "object_text", "type": "text", "default": "a < b <p>x</p>"}]}
    assert PROJECT_FORMAT.migrations[3](json.loads(json.dumps(data))) == data
    assert PROJECT_FORMAT.current >= 4


def test_v3_project_loads_unchanged_and_saves_at_the_current_format(service, module):
    project_id, module_id = module
    req = service.add_requirement(project_id, module_id, {"object_text": "a < b <p>x</p>"})
    path = service.store.root / project_id / "project.json"
    current = json.loads(path.read_text())
    # The same project as a genuine version 3 file holds it (module_order as id strings, format 7).
    v3 = {k: v for k, v in current.items() if k not in ("locked", "references")}
    v3["format_version"] = 3
    v3["module_order"] = [e["module"] for e in current["module_order"]]
    path.write_text(json.dumps(v3, indent=2) + "\n")

    upgraded = upgrade(v3, PROJECT_FORMAT, "Project 'p'")
    assert upgraded == {k: v for k, v in current.items() if k != "format_version"}
    assert service.get_requirement(project_id, module_id, req.id).attributes["object_text"] == "a < b <p>x</p>"

    service.update_requirement(project_id, module_id, req.id, {"note": rich("<p><b>x</b></p>")})
    text = path.read_text()
    assert json.loads(text)["format_version"] == PROJECT_FORMAT.current and text.endswith("}\n")
    assert json.loads(text) == current
    # A fresh service reads the rich value back.
    fresh = ProjectService(ProjectStore(service.store.root))
    assert fresh.get_requirement(project_id, module_id, req.id).attributes["note"] == rich("<p><strong>x</strong></p>")


# --- REST / WebSocket parity -----------------------------------------------------------


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REQMAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(tmp_path / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app)


def test_rest_and_websocket_strip_the_same_way(client):
    project_id = client.post("/projects", json={"name": "P"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "M", "prefix": "M"}).json()["id"]
    base = f"/projects/{project_id}/modules/{module_id}/requirements"
    resp = client.post(base, json={"attributes": {"object_text": DIRTY}})
    assert resp.status_code == 200, resp.text
    rest_id = resp.json()["id"]
    assert resp.json()["attributes"]["object_text"] == CLEAN

    with client.websocket_connect("/ws") as ws:
        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 1, "method": "requirement.add",
            "params": {"project_id": project_id, "module_id": module_id, "attributes": {"object_text": DIRTY}},
        }))
        added = json.loads(ws.receive_text())
        while "id" not in added:
            added = json.loads(ws.receive_text())
        assert added["result"]["attributes"]["object_text"] == CLEAN
        ws.send_text(json.dumps({
            "jsonrpc": "2.0", "id": 2, "method": "requirement.update",
            "params": {"project_id": project_id, "module_id": module_id, "id": rest_id,
                       "attributes": {"object_text": rich("<p><i>i</i><a href='javascript:x'>l</a></p>")}},
        }))
        updated = json.loads(ws.receive_text())
        while "id" not in updated:
            updated = json.loads(ws.receive_text())
        assert updated["result"]["attributes"]["object_text"] == rich("<p><em>i</em>l</p>")

    resp = client.put(f"{base}/{rest_id}", json={"attributes": {"object_text": rich("<p><i>i</i><a href='javascript:x'>l</a></p>")}})
    assert resp.json()["attributes"]["object_text"] == rich("<p><em>i</em>l</p>")


@pytest.mark.parametrize("method", ["feed", "close"])
def test_a_failing_html_parser_keeps_the_markup_as_literal_text(monkeypatch, method):
    # html.parser raises on some markup (e.g. "<![foo[x]]>") on some Python
    # versions; that must never fail a write.
    import reqman.richtext as richtext

    def boom(self, *args):
        raise AssertionError("unexpected call to parse_marked_section()")

    monkeypatch.setattr(richtext._TreeBuilder, method, boom)
    value = rich("<p><strong>a</strong><![foo[x]]></p>")
    result = sanitize(value)
    assert result == "<p><strong>a</strong><![foo[x]]></p>"  # literal text, not markup
    assert sanitize(result) == result
    assert to_plain(value) == result


# --- size limits ---------------------------------------------------------------------


def test_oversize_rich_text_is_rejected_by_service_rest_and_websocket(client):
    from reqman.richtext import MAX_LENGTH, MAX_NODES

    too_long = rich("<p>" + "x" * MAX_LENGTH + "</p>")
    too_many_nodes = rich("<p>" + "x<br/>" * (MAX_NODES // 2 + 1) + "</p>")
    assert len(too_many_nodes) < MAX_LENGTH
    for value in (too_long, too_many_nodes):
        with pytest.raises(ValueError, match="Rich text is too large"):
            sanitize(value)
    project_id = client.post("/projects", json={"name": "P"}).json()["id"]
    module_id = client.post(f"/projects/{project_id}/modules", json={"name": "M", "prefix": "M"}).json()["id"]
    base = f"/projects/{project_id}/modules/{module_id}/requirements"
    ok = client.post(base, json={"attributes": {"object_text": "ok"}}).json()["id"]
    for value in (too_long, too_many_nodes):
        resp = client.post(base, json={"attributes": {"object_text": value}})
        assert resp.status_code == 400 and "Rich text is too large" in resp.text, resp.text[:300]
        resp = client.put(f"{base}/{ok}", json={"attributes": {"object_text": value}})
        assert resp.status_code == 400 and "Rich text is too large" in resp.text, resp.text[:300]
    resp = client.post(
        f"/projects/{project_id}/attributes", json={"id": "note", "name": "Note", "type": "text", "default": too_long}
    )
    assert resp.status_code == 400 and "Rich text is too large" in resp.text, resp.text[:300]
    with client.websocket_connect("/ws") as ws:
        for request_id, (method, extra) in enumerate(
            [("requirement.add", {}), ("requirement.update", {"id": ok})], start=1
        ):
            ws.send_text(json.dumps({
                "jsonrpc": "2.0", "id": request_id, "method": method,
                "params": {"project_id": project_id, "module_id": module_id, **extra,
                           "attributes": {"object_text": too_many_nodes}},
            }))
            reply = json.loads(ws.receive_text())
            while "id" not in reply:
                reply = json.loads(ws.receive_text())
            assert reply["error"]["code"] == -32602
            assert "Rich text is too large" in reply["error"]["message"]
    assert client.get(f"{base}/{ok}").json()["attributes"]["object_text"] == "ok"


def test_stored_over_limit_value_reads_bounded(service, module):
    import time

    from reqman.richtext import MAX_LENGTH

    project_id, module_id = module
    # Hand-edited or pulled from git: over both limits, with hostile nesting.
    stored = rich("<p>" + "".join(f"<b a={i}>" for i in range(200)) + "<p>x" * (MAX_LENGTH // 4))
    start = time.perf_counter()
    plain = to_plain(stored)
    assert time.perf_counter() - start < 15  # shared CI runners are slow; linearity is what matters here
    assert 0 < len(plain) <= len(stored)
    assert set(plain.split("\n")) == {"", "x"}
    assert format_value(stored)  # commit messages / history use to_plain too
    req = service.add_requirement(project_id, module_id, {"object_text": "ok"})
    path = service.store.root / project_id / "modules" / module_id / "requirements" / f"{req.id}.json"
    data = json.loads(path.read_text())
    data["attributes"]["object_text"] = stored
    path.write_text(json.dumps(data, indent=2) + "\n")
    assert service.get_requirement(project_id, module_id, req.id).attributes["object_text"] == stored
    assert "x\nx" in service.export_csv(project_id)
