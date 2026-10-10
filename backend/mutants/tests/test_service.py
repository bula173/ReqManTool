import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import json

import pytest

from reqman.service import ProjectService
from reqman.store import ProjectStore


@pytest.fixture
def service(tmp_path):
    return ProjectService(ProjectStore(tmp_path))


def create_project(service, name="Project One"):
    return service.create_project(name)


def create_module(service, project_id, name="Module One", prefix="REQ"):
    return service.create_module(project_id, name, prefix)


def test_create_project_seeds_builtin_attributes(service):
    project = create_project(service)
    ids = {a.id for a in project.attributes}
    assert ids == {"object_text"}


def test_create_project_generates_slug_id(service):
    project = create_project(service, "Project One")
    assert project.id == "project-one"


def test_create_project_duplicate_name_gets_suffixed_id(service):
    project1 = create_project(service, "Project One")
    project2 = create_project(service, "Project One")
    assert project1.id == "project-one"
    assert project2.id == "project-one-2"


def test_add_attribute_definition(service):
    project = create_project(service)
    project = service.add_attribute_definition(project.id, "priority", "Priority", type="enum", choices=["low", "high"])
    assert any(a.id == "priority" for a in project.attributes)


def test_create_module_generates_slug_id_and_keeps_prefix(service):
    project = create_project(service)
    module = create_module(service, project.id, "Module One", "REQ")
    assert module.id == "module-one"
    assert module.prefix == "REQ"


def test_create_module_invalid_prefix_raises(service):
    project = create_project(service)
    with pytest.raises(ValueError):
        service.create_module(project.id, "Module One", "1REQ")


def test_add_requirement_requires_title(service):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(ValueError):
        service.add_requirement(project.id, module.id, {})


def test_add_requirement_allocates_prefixed_id(service):
    project = create_project(service)
    module = create_module(service, project.id, prefix="REQ")
    req1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    req2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "REQ-0002"


def test_add_requirement_shares_counter_across_modules_with_same_prefix(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "REQ")
    req1 = service.add_requirement(project.id, mod1.id, {"object_text": "First"})
    req2 = service.add_requirement(project.id, mod2.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "REQ-0002"


def test_add_requirement_different_prefixes_do_not_share_counter(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    req1 = service.add_requirement(project.id, mod1.id, {"object_text": "First"})
    req2 = service.add_requirement(project.id, mod2.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "TEST-0001"


def test_get_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First requirement"})
    req = service.get_requirement(project.id, module.id, created.id)
    assert req.attributes["object_text"] == "First requirement"


def test_update_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First"})
    updated = service.update_requirement(project.id, module.id, created.id, {"object_text": "Updated"})
    assert updated.attributes["object_text"] == "Updated"


def test_update_requirement_unknown_attribute_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First"})
    with pytest.raises(ValueError):
        service.update_requirement(project.id, module.id, created.id, {"nope": "x"})


def test_remove_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    created = service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.remove_requirement(project.id, module.id, created.id)
    with pytest.raises(FileNotFoundError):
        service.get_requirement(project.id, module.id, created.id)


def test_create_link_configuration_adds_system_attribute(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    assert config.id == "verifies"

    project = service.get_project(project.id)
    attr = next(a for a in project.attributes if a.id == "verifies")
    assert attr.system is True
    assert attr.type == "text"


def test_link_requirements_across_modules(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})

    result = service.link_requirements(project.id, config.id, source.id, target.id)
    assert result["source"].attributes["verifies"] == f"> {target.id}"
    assert result["target"].attributes["verifies"] == f"< {source.id}"


def test_link_requirements_wrong_module_raises(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    mod3 = create_module(service, project.id, "Other Module", "OTH")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    other = service.add_requirement(project.id, mod3.id, {"object_text": "Other"})

    with pytest.raises(ValueError):
        service.link_requirements(project.id, config.id, source.id, other.id)


def test_link_requirements_wrong_direction_raises(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    a = service.add_requirement(project.id, mod1.id, {"object_text": "A"})
    b = service.add_requirement(project.id, mod2.id, {"object_text": "B"})

    with pytest.raises(ValueError):
        service.link_requirements(project.id, config.id, b.id, a.id)


def test_link_requirements_without_configuration_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    a = service.add_requirement(project.id, module.id, {"object_text": "A"})
    b = service.add_requirement(project.id, module.id, {"object_text": "B"})

    with pytest.raises(FileNotFoundError):
        service.link_requirements(project.id, "does-not-exist", a.id, b.id)


def test_unlink_requirements_clears_both_mirrors(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    result = service.unlink_requirements(project.id, config.id, source.id, target.id)
    assert "verifies" not in result["source"].attributes
    assert "verifies" not in result["target"].attributes


def test_remove_requirement_cleans_up_links_and_other_side_mirror(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    service.remove_requirement(project.id, mod1.id, source.id)
    target = service.get_requirement(project.id, mod2.id, target.id)
    assert "verifies" not in target.attributes
    assert service.store.load_link_pairs(project.id, config.id) == []


def test_remove_link_configuration_cleans_up_everything(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    service.remove_link_configuration(project.id, config.id)

    project_after = service.get_project(project.id)
    assert config.id not in {c.id for c in project_after.link_configurations}
    assert config.id not in {a.id for a in project_after.attributes}
    source = service.get_requirement(project.id, mod1.id, source.id)
    target = service.get_requirement(project.id, mod2.id, target.id)
    assert config.id not in source.attributes
    assert config.id not in target.attributes


def test_setting_system_attribute_directly_raises(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    with pytest.raises(ValueError):
        service.add_requirement(
            project.id, mod1.id, {"object_text": "X", config.id: "> REQ-9999"}
        )


def test_purge_module(service):
    project = create_project(service)
    module = create_module(service, project.id)
    service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.add_requirement(project.id, module.id, {"object_text": "Second"})
    purged = service.purge_module(project.id, module.id)
    assert purged == 2
    assert service.list_requirements(project.id, module.id) == []


def test_purge_project(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    service.add_requirement(project.id, mod1.id, {"object_text": "First"})
    service.add_requirement(project.id, mod2.id, {"object_text": "Second"})
    purged = service.purge_project(project.id)
    assert purged == 2


def test_purge_does_not_reuse_ids(service):
    project = create_project(service)
    module = create_module(service, project.id)
    req1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.purge_module(project.id, module.id)
    req2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    assert req1.id == "REQ-0001"
    assert req2.id == "REQ-0002"


def test_add_requirement_with_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Chapter 1"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    assert child.parent_id == parent.id


def test_add_requirement_with_missing_parent_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(FileNotFoundError):
        service.add_requirement(
            project.id, module.id, {"object_text": "Child"}, parent_id="REQ-9999"
        )


def test_move_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    chapter1 = service.add_requirement(project.id, module.id, {"object_text": "Chapter 1"})
    chapter2 = service.add_requirement(project.id, module.id, {"object_text": "Chapter 2"})
    req = service.add_requirement(
        project.id, module.id, {"object_text": "Req"}, parent_id=chapter1.id
    )
    moved = service.move_requirement(project.id, module.id, req.id, chapter2.id)
    assert moved.parent_id == chapter2.id

    moved_to_root = service.move_requirement(project.id, module.id, req.id, None)
    assert moved_to_root.parent_id is None


def test_move_requirement_rejects_self_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    req = service.add_requirement(project.id, module.id, {"object_text": "Req"})
    with pytest.raises(ValueError):
        service.move_requirement(project.id, module.id, req.id, req.id)


def test_move_requirement_rejects_descendant_cycle(service):
    project = create_project(service)
    module = create_module(service, project.id)
    grandparent = service.add_requirement(project.id, module.id, {"object_text": "GP"})
    parent = service.add_requirement(
        project.id, module.id, {"object_text": "Parent"}, parent_id=grandparent.id
    )
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    with pytest.raises(ValueError):
        service.move_requirement(project.id, module.id, grandparent.id, child.id)


def test_move_requirement_missing_parent_raises(service):
    project = create_project(service)
    module = create_module(service, project.id)
    req = service.add_requirement(project.id, module.id, {"object_text": "Req"})
    with pytest.raises(FileNotFoundError):
        service.move_requirement(project.id, module.id, req.id, "REQ-9999")


def test_remove_requirement_promotes_children(service):
    project = create_project(service)
    module = create_module(service, project.id)
    grandparent = service.add_requirement(project.id, module.id, {"object_text": "GP"})
    parent = service.add_requirement(
        project.id, module.id, {"object_text": "Parent"}, parent_id=grandparent.id
    )
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    service.remove_requirement(project.id, module.id, parent.id)
    child = service.get_requirement(project.id, module.id, child.id)
    assert child.parent_id == grandparent.id


def test_remove_root_requirement_promotes_children_to_root(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=parent.id
    )
    service.remove_requirement(project.id, module.id, parent.id)
    child = service.get_requirement(project.id, module.id, child.id)
    assert child.parent_id is None


def test_export_reqif_nests_hierarchy_by_parent_id(service):
    from reqif.parser import ReqIFParser

    project = create_project(service)
    module = create_module(service, project.id)
    chapter = service.add_requirement(project.id, module.id, {"object_text": "Chapter 1"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child"}, parent_id=chapter.id
    )

    xml = service.export_reqif(project.id)
    bundle = ReqIFParser.parse_from_string(xml)
    specification = bundle.core_content.req_if_content.specifications[0]

    assert len(specification.children) == 1
    root_node = specification.children[0]
    assert root_node.spec_object == chapter.id
    assert root_node.children is not None
    assert len(root_node.children) == 1
    assert root_node.children[0].spec_object == child.id


def test_export_reqif_includes_link_spec_relations(service):
    from reqif.parser import ReqIFParser

    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    xml = service.export_reqif(project.id)
    bundle = ReqIFParser.parse_from_string(xml)
    relations = bundle.core_content.req_if_content.spec_relations
    assert relations is not None
    assert len(relations) == 1
    assert relations[0].source == source.id
    assert relations[0].target == target.id

    relation_types = [
        st for st in bundle.core_content.req_if_content.spec_types
        if getattr(st, "long_name", None) == "Verifies"
    ]
    assert len(relation_types) == 1


def test_add_requirement_assigns_sequential_order(service):
    project = create_project(service)
    module = create_module(service, project.id)
    r1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    r2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    r3 = service.add_requirement(project.id, module.id, {"object_text": "Third"})
    assert (r1.order, r2.order, r3.order) == (0, 1, 2)


def test_add_requirement_order_is_per_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child1 = service.add_requirement(project.id, module.id, {"object_text": "Child 1"}, parent_id=parent.id)
    child2 = service.add_requirement(project.id, module.id, {"object_text": "Child 2"}, parent_id=parent.id)
    assert parent.order == 0
    assert (child1.order, child2.order) == (0, 1)


def test_reorder_requirement_swaps_adjacent_siblings(service):
    project = create_project(service)
    module = create_module(service, project.id)
    r1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    r2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    r3 = service.add_requirement(project.id, module.id, {"object_text": "Third"})

    service.reorder_requirement(project.id, module.id, r2.id, "up")
    r1 = service.get_requirement(project.id, module.id, r1.id)
    r2 = service.get_requirement(project.id, module.id, r2.id)
    r3 = service.get_requirement(project.id, module.id, r3.id)
    assert (r2.order, r1.order, r3.order) == (0, 1, 2)


def test_reorder_requirement_noop_at_boundary(service):
    project = create_project(service)
    module = create_module(service, project.id)
    r1 = service.add_requirement(project.id, module.id, {"object_text": "First"})
    r2 = service.add_requirement(project.id, module.id, {"object_text": "Second"})

    service.reorder_requirement(project.id, module.id, r1.id, "up")
    r1 = service.get_requirement(project.id, module.id, r1.id)
    r2 = service.get_requirement(project.id, module.id, r2.id)
    assert (r1.order, r2.order) == (0, 1)

    service.reorder_requirement(project.id, module.id, r2.id, "down")
    r1 = service.get_requirement(project.id, module.id, r1.id)
    r2 = service.get_requirement(project.id, module.id, r2.id)
    assert (r1.order, r2.order) == (0, 1)


def test_move_requirement_appends_at_end_of_new_parent(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent1 = service.add_requirement(project.id, module.id, {"object_text": "Parent 1"})
    parent2 = service.add_requirement(project.id, module.id, {"object_text": "Parent 2"})
    service.add_requirement(project.id, module.id, {"object_text": "Existing child"}, parent_id=parent2.id)
    moved = service.add_requirement(project.id, module.id, {"object_text": "Moved"}, parent_id=parent1.id)

    moved = service.move_requirement(project.id, module.id, moved.id, parent2.id)
    assert moved.order == 1


def test_add_requirement_after_removal_does_not_reuse_an_existing_order(service):
    project = create_project(service)
    module = create_module(service, project.id)
    first = service.add_requirement(project.id, module.id, {"object_text": "First"})
    second = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    third = service.add_requirement(project.id, module.id, {"object_text": "Third"})
    service.remove_requirement(project.id, module.id, first.id)

    added = service.add_requirement(project.id, module.id, {"object_text": "Added"})

    # Existing orders are not renumbered; the new one goes after the last.
    assert service.get_requirement(project.id, module.id, second.id).order == 1
    assert service.get_requirement(project.id, module.id, third.id).order == 2
    assert added.order == 3


def test_move_requirement_after_removal_does_not_reuse_an_existing_order(service):
    project = create_project(service)
    module = create_module(service, project.id)
    target = service.add_requirement(project.id, module.id, {"object_text": "Target"})
    gone = service.add_requirement(project.id, module.id, {"object_text": "Gone"}, parent_id=target.id)
    kept = service.add_requirement(project.id, module.id, {"object_text": "Kept"}, parent_id=target.id)
    mover = service.add_requirement(project.id, module.id, {"object_text": "Mover"})
    service.remove_requirement(project.id, module.id, gone.id)

    moved = service.move_requirement(project.id, module.id, mover.id, target.id)

    assert service.get_requirement(project.id, module.id, kept.id).order == 1
    assert moved.order == 2


# --- move_requirement with an index ------------------------------------------


def _children(service, project_id, module_id, parent_id):
    """``[(id, order), ...]`` of ``parent_id``'s children, sorted by (order, id)."""
    reqs = service.list_requirements(project_id, module_id)
    return sorted(((r.id, r.order) for r in reqs if r.parent_id == parent_id), key=lambda t: (t[1], t[0]))


def _tree(service, count=3):
    """A project/module with a top-level ``target`` holding ``count`` children and a top-level ``mover``."""
    project = create_project(service)
    module = create_module(service, project.id)
    target = service.add_requirement(project.id, module.id, {"object_text": "Target"})
    children = [
        service.add_requirement(project.id, module.id, {"object_text": f"C{i}"}, parent_id=target.id)
        for i in range(count)
    ]
    mover = service.add_requirement(project.id, module.id, {"object_text": "Mover"})
    return project, module, target, [c.id for c in children], mover


def test_move_requirement_without_index_still_appends_and_leaves_siblings_untouched(service):
    project, module, target, kids, mover = _tree(service)
    service.remove_requirement(project.id, module.id, kids[0])  # leaves a gap at order 0

    moved = service.move_requirement(project.id, module.id, mover.id, target.id)

    assert moved.order == 3
    assert _children(service, project.id, module.id, target.id) == [(kids[1], 1), (kids[2], 2), (mover.id, 3)]


@pytest.mark.parametrize(
    "index, position",
    [(0, 0), (1, 1), (2, 2), (3, 3)],
)
def test_move_requirement_with_index_places_and_renumbers(service, index, position):
    project, module, target, kids, mover = _tree(service)

    moved = service.move_requirement(project.id, module.id, mover.id, target.id, index)

    expected = list(kids)
    expected.insert(position, mover.id)
    assert moved.parent_id == target.id
    assert moved.order == position
    assert _children(service, project.id, module.id, target.id) == [(rid, i) for i, rid in enumerate(expected)]
    # The returned requirement matches what was persisted.
    assert service.get_requirement(project.id, module.id, mover.id).order == position


def test_move_requirement_index_past_the_end_is_clamped_to_append(service):
    project, module, target, kids, mover = _tree(service)

    moved = service.move_requirement(project.id, module.id, mover.id, target.id, 99)

    assert moved.order == 3
    assert _children(service, project.id, module.id, target.id) == [(kids[0], 0), (kids[1], 1), (kids[2], 2), (mover.id, 3)]


@pytest.mark.parametrize("bad", [-1, 1.5, "1", True, False, [0], {"i": 0}])
def test_move_requirement_rejects_an_invalid_index_without_writing(service, bad):
    project, module, target, kids, mover = _tree(service)
    before = _children(service, project.id, module.id, None)

    with pytest.raises(ValueError, match="must be a non-negative integer"):
        service.move_requirement(project.id, module.id, mover.id, target.id, bad)

    assert service.get_requirement(project.id, module.id, mover.id).parent_id is None
    assert _children(service, project.id, module.id, None) == before


def test_move_requirement_same_parent_with_index_repositions(service):
    project, module, target, kids, mover = _tree(service, count=4)

    moved = service.move_requirement(project.id, module.id, kids[3], target.id, 0)
    assert moved.order == 0
    assert [rid for rid, _ in _children(service, project.id, module.id, target.id)] == [kids[3], kids[0], kids[1], kids[2]]

    service.move_requirement(project.id, module.id, kids[3], target.id, 2)
    assert _children(service, project.id, module.id, target.id) == [
        (kids[0], 0), (kids[1], 1), (kids[3], 2), (kids[2], 3)
    ]

    # Moving to its current position changes nothing.
    service.move_requirement(project.id, module.id, kids[3], target.id, 2)
    assert _children(service, project.id, module.id, target.id) == [
        (kids[0], 0), (kids[1], 1), (kids[3], 2), (kids[2], 3)
    ]


def test_move_requirement_to_root_with_index(service):
    project, module, target, kids, mover = _tree(service)

    moved = service.move_requirement(project.id, module.id, kids[1], None, 1)

    assert moved.parent_id is None
    assert _children(service, project.id, module.id, None) == [(target.id, 0), (kids[1], 1), (mover.id, 2)]
    # The old parent keeps its remaining children in order (the gap is not renumbered).
    assert _children(service, project.id, module.id, target.id) == [(kids[0], 0), (kids[2], 2)]


def test_move_requirement_with_index_renumbers_duplicate_legacy_orders(service, tmp_path):
    project, module, target, kids, mover = _tree(service)
    # Simulate hand-edited/legacy data: every child of target has order 0.
    for rid in kids:
        req = service.get_requirement(project.id, module.id, rid)
        req.order = 0
        service.store.save_requirement(project.id, module.id, req)

    service.move_requirement(project.id, module.id, mover.id, target.id, 1)

    # Ties are broken by id, deterministically; the result is contiguous.
    ordered = sorted(kids)
    assert _children(service, project.id, module.id, target.id) == [
        (ordered[0], 0), (mover.id, 1), (ordered[1], 2), (ordered[2], 3)
    ]


def test_move_requirement_with_index_writes_one_history_record(service, monkeypatch):
    project, module, target, kids, mover = _tree(service)
    calls = []
    monkeypatch.setattr(service, "_auto_commit", lambda project_id: calls.append(project_id))

    service.move_requirement(project.id, module.id, mover.id, target.id, 0)

    assert calls == [project.id]


def test_move_requirement_with_index_keeps_parent_constraints(service):
    project, module, target, kids, mover = _tree(service)
    other = create_module(service, project.id, "Other", "OTH")
    foreign = service.add_requirement(project.id, other.id, {"object_text": "Foreign"})
    before = _children(service, project.id, module.id, target.id)

    with pytest.raises(FileNotFoundError):
        service.move_requirement(project.id, module.id, mover.id, foreign.id, 0)
    with pytest.raises(ValueError):
        service.move_requirement(project.id, module.id, target.id, kids[0], 0)
    with pytest.raises(ValueError):
        service.move_requirement(project.id, module.id, target.id, target.id, 0)

    assert service.get_requirement(project.id, module.id, mover.id).parent_id is None
    assert service.get_requirement(project.id, module.id, target.id).parent_id is None
    assert _children(service, project.id, module.id, target.id) == before


def test_move_requirement_with_index_keeps_json_format(service, tmp_path):
    project, module, target, kids, mover = _tree(service)

    service.move_requirement(project.id, module.id, mover.id, target.id, 0)

    files = list(tmp_path.rglob(f"{mover.id}.json"))
    assert len(files) == 1
    text = files[0].read_text()
    data = json.loads(text)
    assert text == json.dumps(data, indent=2, ensure_ascii=False) + "\n"
    assert list(data) == list(service.get_requirement(project.id, module.id, kids[0]).to_dict())


def test_copy_requirement_after_removal_does_not_reuse_an_existing_order(service):
    project = create_project(service)
    module = create_module(service, project.id)
    first = service.add_requirement(project.id, module.id, {"object_text": "First"})
    second = service.add_requirement(project.id, module.id, {"object_text": "Second"})
    service.remove_requirement(project.id, module.id, first.id)

    copy = service.copy_requirement(project.id, module.id, second.id, None)

    assert service.get_requirement(project.id, module.id, second.id).order == 1
    assert copy.order == 2


def test_copy_requirement_duplicates_subtree_with_new_ids(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child = service.add_requirement(project.id, module.id, {"object_text": "Child"}, parent_id=parent.id)

    copy = service.copy_requirement(project.id, module.id, parent.id, None)
    assert copy.id != parent.id
    assert copy.attributes["object_text"] == "Parent"

    all_reqs = service.list_requirements(project.id, module.id)
    copied_children = [r for r in all_reqs if r.parent_id == copy.id]
    assert len(copied_children) == 1
    assert copied_children[0].id != child.id
    assert copied_children[0].attributes["object_text"] == "Child"


def test_copy_requirement_excludes_system_attributes(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Source Module", "SRC")
    mod2 = create_module(service, project.id, "Target Module", "TGT")
    config = service.create_link_configuration(project.id, "Verifies", mod1.id, mod2.id)
    source = service.add_requirement(project.id, mod1.id, {"object_text": "Source"})
    target = service.add_requirement(project.id, mod2.id, {"object_text": "Target"})
    service.link_requirements(project.id, config.id, source.id, target.id)

    copy = service.copy_requirement(project.id, mod1.id, source.id, None)
    assert config.id not in copy.attributes


def test_copy_requirement_rejects_copy_under_own_descendant(service):
    project = create_project(service)
    module = create_module(service, project.id)
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent"})
    child = service.add_requirement(project.id, module.id, {"object_text": "Child"}, parent_id=parent.id)

    with pytest.raises(ValueError):
        service.copy_requirement(project.id, module.id, parent.id, child.id)


def test_export_csv_has_header_and_rows(service):
    import csv as csv_module
    import io

    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    parent = service.add_requirement(project.id, mod1.id, {"object_text": "Parent"})
    child = service.add_requirement(
        project.id, mod1.id, {"object_text": "Child"}, parent_id=parent.id
    )
    other = service.add_requirement(project.id, mod2.id, {"object_text": "Other"})

    csv_text = service.export_csv(project.id)
    rows = list(csv_module.reader(io.StringIO(csv_text)))

    assert rows[0] == ["Id", "Module", "Parent", "ObjectText"]
    assert rows[1] == [parent.id, mod1.id, "", "Parent"]
    assert rows[2] == [child.id, mod1.id, parent.id, "Child"]
    assert rows[3] == [other.id, mod2.id, "", "Other"]


def test_export_csv_includes_custom_attributes(service):
    import csv as csv_module
    import io

    project = create_project(service)
    module = create_module(service, project.id)
    service.add_attribute_definition(project.id, "priority", "Priority", type="string")
    req = service.add_requirement(
        project.id, module.id, {"object_text": "Req", "priority": "High"}
    )

    csv_text = service.export_csv(project.id)
    rows = list(csv_module.reader(io.StringIO(csv_text)))

    assert rows[0] == ["Id", "Module", "Parent", "ObjectText", "Priority"]
    assert rows[1] == [req.id, module.id, "", "Req", "High"]


def test_export_csv_writes_rows_depth_first_with_siblings_by_order(service):
    import csv as csv_module
    import io

    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    a = service.add_requirement(project.id, mod1.id, {"object_text": "A"})
    b = service.add_requirement(project.id, mod1.id, {"object_text": "B"})
    a1 = service.add_requirement(project.id, mod1.id, {"object_text": "A1"}, parent_id=a.id)
    a2 = service.add_requirement(project.id, mod1.id, {"object_text": "A2"}, parent_id=a.id)
    other = service.add_requirement(project.id, mod2.id, {"object_text": "Other"})
    service.reorder_requirement(project.id, mod1.id, b.id, "up")
    service.reorder_requirement(project.id, mod1.id, a2.id, "up")

    rows = list(csv_module.reader(io.StringIO(service.export_csv(project.id))))

    assert [row[0] for row in rows[1:]] == [b.id, a.id, a2.id, a1.id, other.id]
    assert rows[0] == ["Id", "Module", "Parent", "ObjectText"]
    assert rows[3] == [a2.id, mod1.id, a.id, "A2"]


def test_create_view(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="string")
    view = service.create_view(project.id, "Compact", ["object_text"])
    assert view.id == "compact"
    assert view.attribute_ids == ["object_text"]

    views = service.list_views(project.id)
    assert len(views) == 1
    assert views[0].id == view.id


def test_create_view_rejects_unknown_attribute(service):
    project = create_project(service)
    with pytest.raises(ValueError):
        service.create_view(project.id, "Compact", ["does-not-exist"])


def test_remove_view(service):
    project = create_project(service)
    view = service.create_view(project.id, "Compact", ["object_text"])
    service.remove_view(project.id, view.id)
    assert service.list_views(project.id) == []


def test_remove_view_missing_raises(service):
    project = create_project(service)
    with pytest.raises(FileNotFoundError):
        service.remove_view(project.id, "does-not-exist")


def test_find_requirement_locates_across_modules(service):
    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    req = service.add_requirement(project.id, mod2.id, {"object_text": "Somewhere"})

    module_id, found = service.find_requirement(project.id, req.id)
    assert module_id == mod2.id
    assert found.id == req.id


def test_find_requirement_missing_raises(service):
    project = create_project(service)
    with pytest.raises(FileNotFoundError):
        service.find_requirement(project.id, "REQ-9999")


def test_export_xlsx_has_header_and_rows(service):
    import io

    from openpyxl import load_workbook

    project = create_project(service)
    mod1 = create_module(service, project.id, "Module One", "REQ")
    mod2 = create_module(service, project.id, "Module Two", "TEST")
    parent = service.add_requirement(project.id, mod1.id, {"object_text": "Parent"})
    child = service.add_requirement(
        project.id, mod1.id, {"object_text": "Child"}, parent_id=parent.id
    )
    other = service.add_requirement(project.id, mod2.id, {"object_text": "Other"})

    xlsx_bytes = service.export_xlsx(project.id)
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == ("Id", "Module", "Parent", "ObjectText")
    assert rows[1][:2] == (parent.id, mod1.id)
    assert not rows[1][2]  # empty parent, written as ""
    assert rows[1][3] == "Parent"
    assert rows[2] == (child.id, mod1.id, parent.id, "Child")
    assert rows[3][:2] == (other.id, mod2.id)
    assert not rows[3][2]
    assert rows[3][3] == "Other"


def test_export_xlsx_includes_custom_attributes(service):
    import io

    from openpyxl import load_workbook

    project = create_project(service)
    module = create_module(service, project.id)
    service.add_attribute_definition(project.id, "priority", "Priority", type="string")
    req = service.add_requirement(
        project.id, module.id, {"object_text": "Req", "priority": "High"}
    )

    xlsx_bytes = service.export_xlsx(project.id)
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == ("Id", "Module", "Parent", "ObjectText", "Priority")
    assert rows[1][:2] == (req.id, module.id)
    assert not rows[1][2]
    assert rows[1][3:] == ("Req", "High")


def test_export_docx_structure(service):
    import io

    from docx import Document as DocxDocument

    project = create_project(service)
    module = create_module(service, project.id, "Spec Module", "REQ")
    parent = service.add_requirement(project.id, module.id, {"object_text": "Parent text"})
    child = service.add_requirement(
        project.id, module.id, {"object_text": "Child text"}, parent_id=parent.id
    )

    docx_bytes = service.export_docx(project.id)
    document = DocxDocument(io.BytesIO(docx_bytes))
    paragraphs = document.paragraphs

    assert paragraphs[0].text == project.name
    assert paragraphs[0].style.name == "Title"

    module_heading = next(p for p in paragraphs if p.text == "Spec Module")
    assert module_heading.style.name == "Heading 1"

    parent_heading = next(p for p in paragraphs if p.text == parent.id)
    assert parent_heading.style.name == "Heading 2"
    child_heading = next(p for p in paragraphs if p.text == child.id)
    assert child_heading.style.name == "Heading 3"

    # The child heading comes after the parent heading (nested under it).
    assert paragraphs.index(child_heading) > paragraphs.index(parent_heading)

    body_texts = [p.text for p in paragraphs]
    assert any("ObjectText: Parent text" in t for t in body_texts)
    assert any("ObjectText: Child text" in t for t in body_texts)


def test_export_docx_skips_empty_attribute_values(service):
    import io

    from docx import Document as DocxDocument

    project = create_project(service)
    service.add_attribute_definition(project.id, "notes", "Notes", type="string")
    module = create_module(service, project.id)
    req = service.add_requirement(project.id, module.id, {"object_text": "Has text, no notes"})

    docx_bytes = service.export_docx(project.id)
    document = DocxDocument(io.BytesIO(docx_bytes))
    body_texts = [p.text for p in document.paragraphs]

    assert any("ObjectText: Has text, no notes" in t for t in body_texts)
    assert not any(t.startswith("Notes:") for t in body_texts)


# --- CSV import ----------------------------------------------------------------


def test_import_csv_creates_nested_requirements_with_server_ids(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    service.add_attribute_definition(
        project.id, "status", "Status", type="enum", choices=["Draft", "Approved"]
    )
    module = create_module(service, project.id)
    existing = service.add_requirement(project.id, module.id, {"object_text": "Existing"})

    csv_text = (
        "Id,Module,Parent,ObjectText,Priority,Safety,Status\n"
        "X-1,other,X-2,Child,3,yes,Draft\n"
        "X-2,other,,Chapter,,FALSE,\n"
    )
    result = service.import_csv(project.id, module.id, csv_text)

    assert result == {
        "imported": 2,
        "requirement_ids": ["REQ-0003", "REQ-0002"],
        "ignored_columns": [],
    }
    chapter = service.get_requirement(project.id, module.id, "REQ-0002")
    child = service.get_requirement(project.id, module.id, "REQ-0003")
    assert chapter.parent_id is None and chapter.order == 1
    assert chapter.attributes == {"object_text": "Chapter", "safety": False}
    assert child.parent_id == "REQ-0002" and child.order == 0
    assert child.attributes == {"object_text": "Child", "priority": 3, "safety": True, "status": "Draft"}
    assert service.get_requirement(project.id, module.id, existing.id).order == 0


def test_import_csv_roundtrips_export_and_ignores_link_mirror_columns(service):
    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "A")
    mod_b = create_module(service, project.id, "Module B", "B")
    parent = service.add_requirement(project.id, mod_a.id, {"object_text": "Parent"})
    service.add_requirement(project.id, mod_a.id, {"object_text": "Child"}, parent_id=parent.id)
    target = service.add_requirement(project.id, mod_b.id, {"object_text": "Target"})
    config = service.create_link_configuration(project.id, "Satisfies", mod_a.id, mod_b.id)
    service.link_requirements(project.id, config.id, parent.id, target.id)

    exported = service.export_csv(project.id)
    module_a_only = "\n".join(
        line for line in exported.splitlines() if not line.split(",")[1] == mod_b.id
    )
    result = service.import_csv(project.id, mod_b.id, module_a_only)

    assert result["imported"] == 2
    assert result["ignored_columns"] == ["Satisfies"]
    imported = {r.id: r for r in service.list_requirements(project.id, mod_b.id) if r.id != target.id}
    by_text = {r.attributes["object_text"]: r for r in imported.values()}
    assert by_text["Child"].parent_id == by_text["Parent"].id
    assert all(config.id not in r.attributes for r in imported.values())


def test_import_csv_is_all_or_nothing_with_row_numbered_errors(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    csv_text = "Id,Parent,ObjectText,Priority\nA,,Fine,1\nB,,,2\nC,Z,Orphan,high\n"
    with pytest.raises(ValueError) as excinfo:
        service.import_csv(project.id, module.id, csv_text)

    message = str(excinfo.value)
    assert "Row 3: required attribute 'ObjectText' is empty" in message
    assert "Row 4: Priority: 'high' is not an integer" in message
    assert "Row 4: Parent 'Z' does not match" in message
    assert service.list_requirements(project.id, module.id) == []
    assert service.get_project(project.id).id_counters == {}


@pytest.mark.parametrize(
    "csv_text, expected",
    [
        ("Id,ObjectText,Bogus\n1,x,y\n", "Unknown column(s) ['Bogus']"),
        ("Id,Parent,ObjectText\nA,B,x\nB,A,y\n", "form a cycle"),
        ("Id,Parent,ObjectText\nA,A,x\n", "cannot be its own parent"),
        ("Id,ObjectText\nA,x\nA,y\n", "Id 'A' is already used by row 2"),
        ("Id,Module,ObjectText\n1,m1,x\n2,m2,y\n", "contains several modules"),
        ("Id,ObjectText\n", "no requirement rows"),
        ("", "empty"),
    ],
)
def test_import_csv_rejects_invalid_files(service, csv_text, expected):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(ValueError) as excinfo:
        service.import_csv(project.id, module.id, csv_text)
    assert expected in str(excinfo.value)
    assert service.list_requirements(project.id, module.id) == []


def test_import_csv_missing_module_raises_not_found(service):
    project = create_project(service)
    with pytest.raises(FileNotFoundError):
        service.import_csv(project.id, "missing", "ObjectText\nx\n")


def _import_error(service, project_id, module_id, csv_text):
    with pytest.raises(ValueError) as excinfo:
        service.import_csv(project_id, module_id, csv_text)
    assert service.list_requirements(project_id, module_id) == []
    return str(excinfo.value)


def _by_text(service, project_id, module_id):
    return {r.attributes.get("object_text"): r for r in service.list_requirements(project_id, module_id)}


def test_import_csv_keeps_quoted_multiline_commas_and_quotes_verbatim(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "notes", "Notes")
    module = create_module(service, project.id)

    csv_text = (
        'Id,ObjectText,Notes\r\n'
        '1,"Line one\nLine two, with comma","He said ""hi"", then left"\r\n'
        '2,"CRLF\r\ninside","  padded  "\r\n'
        '3,"""quoted"" start",\r\n'
    )
    result = service.import_csv(project.id, module.id, csv_text)

    assert result["imported"] == 3
    reqs = {r.id: r.attributes for r in service.list_requirements(project.id, module.id)}
    assert reqs["REQ-0001"] == {"object_text": "Line one\nLine two, with comma", "notes": 'He said "hi", then left'}
    assert reqs["REQ-0002"] == {"object_text": "CRLF\r\ninside", "notes": "  padded  "}
    assert reqs["REQ-0003"] == {"object_text": '"quoted" start'}


def test_import_csv_accepts_bare_cr_line_endings(service):
    project = create_project(service)
    module = create_module(service, project.id)

    # Excel for Mac's "CSV (Macintosh)" format ends lines with a bare CR.
    csv_text = 'Id,Parent,ObjectText\r1,,hello\r2,1,"multi\r\nline"\r'
    result = service.import_csv(project.id, module.id, csv_text)

    assert result["imported"] == 2
    reqs = _by_text(service, project.id, module.id)
    assert reqs["multi\r\nline"].parent_id == reqs["hello"].id


def test_import_csv_row_numbers_count_records_not_physical_lines(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, 'ObjectText,Priority\n"a\nb\nc",1\nd,x\n')
    assert "Row 3: Priority: 'x' is not an integer" in message


def test_import_csv_allows_empty_header_columns_while_their_cells_are_empty(service):
    project = create_project(service)
    module = create_module(service, project.id)

    csv_text = "Id,,ObjectText,,\n1,,First,,\n2, ,Second\n3,,Third,,,\n"
    result = service.import_csv(project.id, module.id, csv_text)

    assert result["imported"] == 3
    assert set(_by_text(service, project.id, module.id)) == {"First", "Second", "Third"}


@pytest.mark.parametrize(
    "csv_text, expected",
    [
        ("Id,ObjectText,,\n1,x,,oops\n", "Row 2: column(s) [4] have a value but no header"),
        ("Id,,ObjectText\n1,oops,x\n", "Row 2: column(s) [2] have a value but no header"),
        ("Id,ObjectText\n1,x,extra\n", "Row 2: has 3 cells but the header has 2 columns"),
    ],
)
def test_import_csv_rejects_values_without_a_header(service, csv_text, expected):
    project = create_project(service)
    module = create_module(service, project.id)
    assert expected in _import_error(service, project.id, module.id, csv_text)


def test_import_csv_multi_level_parents_declared_after_children(service):
    project = create_project(service)
    module = create_module(service, project.id)
    existing = service.add_requirement(project.id, module.id, {"object_text": "Existing"})

    csv_text = (
        "Id,Parent,ObjectText\n"
        "C,B,Grandchild\n"
        "B2,A,Second child\n"
        ",A,Keyless child\n"
        "B,A,Third child\n"
        "A,,Root\n"
    )
    result = service.import_csv(project.id, module.id, csv_text)

    # Parents are created first; ids are returned in file order.
    assert result["requirement_ids"] == ["REQ-0006", "REQ-0003", "REQ-0004", "REQ-0005", "REQ-0002"]
    reqs = _by_text(service, project.id, module.id)
    root, grandchild = reqs["Root"], reqs["Grandchild"]
    assert (root.id, root.parent_id, root.order) == ("REQ-0002", None, 1)
    assert [(reqs[t].parent_id, reqs[t].order) for t in ("Second child", "Keyless child", "Third child")] == [
        (root.id, 0),
        (root.id, 1),
        (root.id, 2),
    ]
    assert (grandchild.parent_id, grandchild.order) == (reqs["Third child"].id, 0)
    assert reqs["Existing"].id == existing.id and reqs["Existing"].order == 0


def test_import_csv_rejects_longer_cycles_and_rows_hanging_off_them(service):
    project = create_project(service)
    module = create_module(service, project.id)

    csv_text = "Id,Parent,ObjectText\nA,C,a\nB,A,b\nC,B,c\nD,A,d\nE,,fine\n"
    message = _import_error(service, project.id, module.id, csv_text)
    for row in (2, 3, 4, 5):
        assert f"Row {row}: the Parent references form a cycle" in message
    assert "Row 6" not in message
    assert service.get_project(project.id).id_counters == {}


def test_import_csv_rejects_self_parent_inside_otherwise_valid_tree(service):
    project = create_project(service)
    module = create_module(service, project.id)

    message = _import_error(
        service, project.id, module.id, "Id,Parent,ObjectText\nA,,root\nB,A,child\nC,C,self\n"
    )
    assert "Row 4: a requirement cannot be its own parent" in message
    assert "Row 2" not in message and "Row 3" not in message


def test_import_csv_parent_must_be_in_file_not_an_existing_requirement(service):
    project = create_project(service)
    module = create_module(service, project.id)
    existing = service.add_requirement(project.id, module.id, {"object_text": "Existing"})

    with pytest.raises(ValueError) as excinfo:
        service.import_csv(project.id, module.id, f"Parent,ObjectText\n{existing.id},child\n")
    assert f"Row 2: Parent '{existing.id}' does not match the Id of any row in this file" in str(excinfo.value)
    assert [r.id for r in service.list_requirements(project.id, module.id)] == [existing.id]


def test_import_csv_duplicate_ids_are_trimmed_and_case_sensitive(service):
    project = create_project(service)
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, "Id,ObjectText\nA,x\n A ,y\nb,z\nA,w\n")
    assert "Row 3: Id 'A' is already used by row 2" in message
    assert "Row 5: Id 'A' is already used by row 2" in message

    result = service.import_csv(project.id, module.id, "Id,ObjectText\na,x\nA,y\n")
    assert result["imported"] == 2


def test_import_csv_module_column_single_value_allowed_multiple_rejected(service):
    project = create_project(service)
    module = create_module(service, project.id)

    result = service.import_csv(
        project.id, module.id, "Module,ObjectText\nsome-other-module,x\n,y\n some-other-module ,z\n"
    )
    assert result["imported"] == 3

    other = create_module(service, project.id, "Module Two", "REQ")
    message = _import_error(service, project.id, other.id, "Module,ObjectText\nm1,x\nM1,y\n")
    assert "contains several modules ['M1', 'm1']" in message


def test_import_csv_matches_columns_by_attribute_id_and_name(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    result = service.import_csv(project.id, module.id, "object_text, priority \nBy id,7\n")
    assert result["imported"] == 1
    assert service.get_requirement(project.id, module.id, "REQ-0001").attributes == {
        "object_text": "By id",
        "priority": 7,
    }


def test_import_csv_name_match_takes_precedence_over_id_match(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "x", "y")
    service.add_attribute_definition(project.id, "y", "z")
    module = create_module(service, project.id)

    service.import_csv(project.id, module.id, "ObjectText,y\nrow,value\n")

    assert service.get_requirement(project.id, module.id, "REQ-0001").attributes == {
        "object_text": "row",
        "x": "value",
    }


def test_import_csv_rejects_name_and_id_columns_for_same_attribute(service):
    project = create_project(service)
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, "ObjectText,object_text\nx,y\n")
    assert "Column 'object_text' maps to attribute 'object_text', which another column already uses" in message

    message = _import_error(service, project.id, module.id, "ObjectText,ObjectText\nx,y\n")
    assert "Duplicate column 'ObjectText'" in message


@pytest.mark.parametrize(
    "cell, expected",
    [
        ("true", True), ("TRUE", True), ("True", True), ("yes", True), ("Yes", True),
        ("y", True), ("Y", True), ("1", True), (" yes ", True),
        ("false", False), ("FALSE", False), ("no", False), ("NO", False),
        ("n", False), ("N", False), ("0", False), (" 0 ", False),
    ],
)
def test_import_csv_boolean_variants(service, cell, expected):
    project = create_project(service)
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    module = create_module(service, project.id)

    service.import_csv(project.id, module.id, f"ObjectText,Safety\nx,{cell}\n")

    assert service.get_requirement(project.id, module.id, "REQ-0001").attributes["safety"] is expected


@pytest.mark.parametrize("cell", ["on", "off", "2", "t", "nope", "-1"])
def test_import_csv_rejects_invalid_booleans(service, cell):
    project = create_project(service)
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, f"ObjectText,Safety\nx,{cell}\n")
    assert f"Row 2: Safety: '{cell}' is not a boolean" in message


@pytest.mark.parametrize(
    "cell, expected",
    [("5", 5), ("+5", 5), ("-3", -3), (" 5", 5), ("5 ", 5), ("0", 0), ("007", 7), ("-0", 0)],
)
def test_import_csv_integer_variants(service, cell, expected):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    service.import_csv(project.id, module.id, f'ObjectText,Priority\nx,"{cell}"\n')

    value = service.get_requirement(project.id, module.id, "REQ-0001").attributes["priority"]
    assert value == expected and type(value) is int


@pytest.mark.parametrize("cell", ["1.0", "1e3", "1_000", "- 3", "++5", "0x10", "five", "٣", "１２"])
def test_import_csv_rejects_invalid_integers(service, cell):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, f'ObjectText,Priority\nx,"{cell}"\n')
    assert f"Row 2: Priority: '{cell}' is not an integer" in message


def test_import_csv_enum_requires_exact_match(service):
    project = create_project(service)
    service.add_attribute_definition(
        project.id, "status", "Status", type="enum", choices=["Draft", "In review"]
    )
    module = create_module(service, project.id)

    message = _import_error(
        service, project.id, module.id, 'ObjectText,Status\na," Draft"\nb,"Draft "\nc,draft\nd,In  review\n'
    )
    for row, value in ((2, " Draft"), (3, "Draft "), (4, "draft"), (5, "In  review")):
        assert f"Row {row}: Status: '{value}' is not one of ['Draft', 'In review']" in message

    service.import_csv(project.id, module.id, "ObjectText,Status\na,In review\nb,   \n")
    reqs = _by_text(service, project.id, module.id)
    assert reqs["a"].attributes["status"] == "In review"
    assert "status" not in reqs["b"].attributes


def test_import_csv_required_attribute_without_default_must_be_filled(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "owner", "Owner", required=True)
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, 'ObjectText,Owner\na,alice\nb,\nc,"   "\n')
    assert "Row 3: required attribute 'Owner' is empty" in message
    assert "Row 4: required attribute 'Owner' is empty" in message
    assert "Row 2" not in message

    message = _import_error(service, project.id, module.id, "ObjectText\na\nb\n")
    assert "Row 2: required attribute 'Owner' is empty" in message
    assert "Row 3: required attribute 'Owner' is empty" in message


def test_import_csv_applies_attribute_defaults_to_empty_cells(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer", default=5)
    service.add_attribute_definition(project.id, "owner", "Owner", required=True, default="team")
    service.add_attribute_definition(
        project.id, "status", "Status", type="enum", choices=["Draft", "Done"], default="Draft"
    )
    module = create_module(service, project.id)

    service.import_csv(project.id, module.id, "ObjectText,Priority,Owner\nempty,,\nfilled,9,alice\n")
    service.import_csv(project.id, module.id, "ObjectText,Status\nno columns,\nexplicit,Done\n")

    reqs = _by_text(service, project.id, module.id)
    assert reqs["empty"].attributes == {"object_text": "empty", "priority": 5, "owner": "team", "status": "Draft"}
    assert reqs["filled"].attributes == {"object_text": "filled", "priority": 9, "owner": "alice", "status": "Draft"}
    assert reqs["no columns"].attributes == {
        "object_text": "no columns", "priority": 5, "owner": "team", "status": "Draft",
    }
    assert reqs["explicit"].attributes["status"] == "Done"


def test_import_csv_skips_rows_where_every_cell_is_empty(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    result = service.import_csv(project.id, module.id, "ObjectText,Priority\nfirst,1\n,\n  ,  \n\n\"\",\nsecond,2\n\n")
    assert result == {"imported": 2, "requirement_ids": ["REQ-0001", "REQ-0002"], "ignored_columns": []}
    assert [(r.attributes["object_text"], r.order) for r in service.list_requirements(project.id, module.id)] == [
        ("first", 0),
        ("second", 1),
    ]

    other = create_module(service, project.id, "Module Two", "OTH")
    message = _import_error(service, project.id, other.id, "ObjectText,Priority\nfirst,1\n,\n\nsecond,x\n")
    assert "Row 5: Priority: 'x' is not an integer" in message


def test_import_csv_whitespace_only_cell_means_not_set(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "notes", "Notes")
    module = create_module(service, project.id)

    service.import_csv(project.id, module.id, 'ObjectText,Notes\nx,"   "\n')

    assert service.get_requirement(project.id, module.id, "REQ-0001").attributes == {"object_text": "x"}


def test_import_csv_appends_roots_after_existing_roots(service):
    project = create_project(service)
    module = create_module(service, project.id)
    first = service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.add_requirement(project.id, module.id, {"object_text": "Nested"}, parent_id=first.id)
    service.add_requirement(project.id, module.id, {"object_text": "Second"})

    service.import_csv(project.id, module.id, "Id,Parent,ObjectText\nA,,Imported A\n,A,Imported child\n,,Imported B\n")

    reqs = _by_text(service, project.id, module.id)
    roots = sorted((r for r in reqs.values() if r.parent_id is None), key=lambda r: r.order)
    assert [r.attributes["object_text"] for r in roots] == ["First", "Second", "Imported A", "Imported B"]
    assert [r.order for r in roots] == [0, 1, 2, 3]
    assert reqs["Imported child"].order == 0


def test_import_csv_roots_follow_existing_roots_after_a_removal(service):
    project = create_project(service)
    module = create_module(service, project.id)
    first = service.add_requirement(project.id, module.id, {"object_text": "First"})
    service.add_requirement(project.id, module.id, {"object_text": "Second"})
    service.add_requirement(project.id, module.id, {"object_text": "Third"})
    service.remove_requirement(project.id, module.id, first.id)

    service.import_csv(project.id, module.id, "ObjectText\nImported\n")

    reqs = _by_text(service, project.id, module.id)
    existing_orders = [reqs["Second"].order, reqs["Third"].order]
    assert reqs["Imported"].order > max(existing_orders)


def test_import_csv_persists_in_json_format_after_reload(service, tmp_path):
    project = create_project(service)
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    module = create_module(service, project.id)
    reference = service.add_requirement(project.id, module.id, {"object_text": "Reference"})

    service.import_csv(project.id, module.id, "Id,Parent,ObjectText,Safety\nP,,Zażółć 🚀,yes\nC,P,Child,no\n")

    reloaded = ProjectService(ProjectStore(tmp_path))
    reqs = {r.id: r for r in reloaded.list_requirements(project.id, module.id)}
    assert reqs["REQ-0002"].attributes == {"object_text": "Zażółć 🚀", "safety": True}
    assert (reqs["REQ-0003"].parent_id, reqs["REQ-0003"].attributes["safety"]) == ("REQ-0002", False)
    assert reloaded.get_project(project.id).id_counters == {"REQ": 3}

    req_dir = tmp_path / project.id / "modules" / module.id / "requirements"
    reference_keys = list(json.loads((req_dir / f"{reference.id}.json").read_text()))
    for requirement_id in ("REQ-0002", "REQ-0003"):
        text = (req_dir / f"{requirement_id}.json").read_text()
        data = json.loads(text)
        assert text == json.dumps(data, indent=2) + "\n"
        assert list(data) == reference_keys == ["id", "attributes", "parent_id", "order"]
    project_text = (tmp_path / project.id / "project.json").read_text()
    assert project_text == json.dumps(json.loads(project_text), indent=2) + "\n"


def test_import_csv_never_reuses_ids_after_failure_and_removal(service, tmp_path):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    assert service.import_csv(project.id, module.id, "ObjectText\na\nb\n")["requirement_ids"] == ["REQ-0001", "REQ-0002"]
    service.remove_requirement(project.id, module.id, "REQ-0002")

    with pytest.raises(ValueError):
        service.import_csv(project.id, module.id, "ObjectText,Priority\nc,1\nd,bad\n")
    assert service.get_project(project.id).id_counters == {"REQ": 2}

    reloaded = ProjectService(ProjectStore(tmp_path))
    result = reloaded.import_csv(project.id, module.id, "ObjectText\ne\n")
    assert result["requirement_ids"] == ["REQ-0003"]
    assert sorted(r.id for r in reloaded.list_requirements(project.id, module.id)) == ["REQ-0001", "REQ-0003"]


def test_import_csv_shares_id_counter_with_other_modules_using_same_prefix(service):
    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "REQ")
    mod_b = create_module(service, project.id, "Module B", "REQ")
    service.add_requirement(project.id, mod_a.id, {"object_text": "In A"})

    assert service.import_csv(project.id, mod_b.id, "ObjectText\nIn B\n")["requirement_ids"] == ["REQ-0002"]


def test_import_csv_strips_leading_bom(service):
    project = create_project(service)
    module = create_module(service, project.id)

    result = service.import_csv(project.id, module.id, "﻿Id,Parent,ObjectText\nA,,root\nB,A,child\n")

    assert result["imported"] == 2
    assert _by_text(service, project.id, module.id)["child"].parent_id == "REQ-0001"


def test_import_csv_reports_at_most_twenty_errors(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    module = create_module(service, project.id)

    csv_text = "ObjectText,Priority\n" + "".join(f"r{i},bad\n" for i in range(25))
    with pytest.raises(ValueError) as excinfo:
        service.import_csv(project.id, module.id, csv_text)

    assert len(excinfo.value.errors) == 25
    message = str(excinfo.value)
    assert message.startswith("CSV import failed; nothing was imported.\n")
    assert "Row 21:" in message and "Row 22:" not in message
    assert message.endswith("... and 5 more error(s)")


def test_import_csv_rejects_unterminated_quoted_cell(service):
    project = create_project(service)
    module = create_module(service, project.id)

    # The closing quote is missing, so the rest of the file would silently end
    # up in one cell. The file is malformed and must be rejected.
    message = _import_error(service, project.id, module.id, 'Id,ObjectText\n1,"open\n2,next\n')
    assert "Malformed CSV" in message


def test_import_csv_rejects_non_string_content(service):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(ValueError, match="must be a string"):
        service.import_csv(project.id, module.id, b"ObjectText\nx\n")


def _build_roundtrip_project(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    service.add_attribute_definition(
        project.id, "status", "Status", type="enum", choices=["Draft", "Approved — final, \"v2\""]
    )
    service.add_attribute_definition(project.id, "notes", "Notes, \"internal\"")
    source = create_module(service, project.id, "Source", "SRC")
    root = service.add_requirement(
        project.id,
        source.id,
        {
            "object_text": 'Zażółć gęślą jaźń — "quoted", comma; 日本語 🚀',
            "priority": -3,
            "safety": True,
            "status": "Approved — final, \"v2\"",
            "notes": "  padded  ",
        },
    )
    child = service.add_requirement(
        project.id,
        source.id,
        {"object_text": "Line 1\nLine 2\r\nLine 3", "priority": 0, "safety": False},
        parent_id=root.id,
    )
    service.add_requirement(
        project.id, source.id, {"object_text": "=SUM(A1)", "notes": "'leading quote"}, parent_id=child.id
    )
    service.add_requirement(project.id, source.id, {"object_text": "Back\\slash\ttab", "status": "Draft"})
    return project, source


def test_import_csv_roundtrips_export_with_special_and_unicode_characters(service):
    project, source = _build_roundtrip_project(service)
    exported = service.export_csv(project.id)
    target = create_module(service, project.id, "Target", "TGT")

    result = service.import_csv(project.id, target.id, exported)

    assert result["imported"] == 4 and result["ignored_columns"] == []
    originals = {r.id: r for r in service.list_requirements(project.id, source.id)}
    copies = {r.attributes["object_text"]: r for r in service.list_requirements(project.id, target.id)}
    copy_id_of = {orig_id: copies[o.attributes["object_text"]].id for orig_id, o in originals.items()}
    for original in originals.values():
        copy = copies[original.attributes["object_text"]]
        assert copy.attributes == original.attributes
        assert copy.parent_id == (copy_id_of[original.parent_id] if original.parent_id else None)
        assert copy.order == original.order
        assert copy.id.startswith("TGT-")


def test_import_csv_roundtrips_export_with_large_text_cell(service):
    project = create_project(service)
    source = create_module(service, project.id, "Source", "SRC")
    big_text = "x" * 200_000
    service.add_requirement(project.id, source.id, {"object_text": big_text})
    exported = service.export_csv(project.id)
    target = create_module(service, project.id, "Target", "TGT")

    service.import_csv(project.id, target.id, exported)

    assert [r.attributes["object_text"] for r in service.list_requirements(project.id, target.id)] == [big_text]


def test_import_csv_rejects_text_over_the_size_cap(service):
    from reqman.csv_import import MAX_CSV_IMPORT_BYTES

    project = create_project(service)
    module = create_module(service, project.id)
    header = "ObjectText\n"
    at_cap = header + "x" * (MAX_CSV_IMPORT_BYTES - len(header))

    with pytest.raises(ValueError, match="too large"):
        service.import_csv(project.id, module.id, at_cap + "x")
    assert service.list_requirements(project.id, module.id) == []

    service.import_csv(project.id, module.id, at_cap)
    assert len(service.list_requirements(project.id, module.id)) == 1


def test_import_csv_roundtrip_preserves_manual_sibling_order(service):
    project = create_project(service)
    source = create_module(service, project.id, "Source", "SRC")
    a = service.add_requirement(project.id, source.id, {"object_text": "A"})
    service.add_requirement(project.id, source.id, {"object_text": "B"})
    c = service.add_requirement(project.id, source.id, {"object_text": "C"})
    service.add_requirement(project.id, source.id, {"object_text": "A1"}, parent_id=a.id)
    a2 = service.add_requirement(project.id, source.id, {"object_text": "A2"}, parent_id=a.id)
    service.reorder_requirement(project.id, source.id, c.id, "up")
    service.reorder_requirement(project.id, source.id, c.id, "up")
    service.reorder_requirement(project.id, source.id, a2.id, "up")
    exported = service.export_csv(project.id)
    target = create_module(service, project.id, "Target", "TGT")

    service.import_csv(project.id, target.id, exported)

    def tree(module_id):
        reqs = service.list_requirements(project.id, module_id)
        text_of = {r.id: r.attributes["object_text"] for r in reqs}

        def children(parent_id):
            return [
                (text_of[r.id], children(r.id))
                for r in sorted((r for r in reqs if r.parent_id == parent_id), key=lambda r: r.order)
            ]

        return children(None)

    expected = [("C", []), ("A", [("A2", []), ("A1", [])]), ("B", [])]
    assert tree(source.id) == expected
    assert tree(target.id) == expected


# --- F1: attribute names -------------------------------------------------------


@pytest.mark.parametrize("name", ["Id", "Module", "Parent", "id", " PARENT ", "module"])
def test_add_attribute_rejects_reserved_export_column_names(service, name):
    project = create_project(service)
    with pytest.raises(ValueError, match="reserved"):
        service.add_attribute_definition(project.id, "custom", name)
    assert [a.id for a in service.get_project(project.id).attributes] == ["object_text"]


@pytest.mark.parametrize("name", ["ObjectText", "objecttext", "  OBJECTTEXT  ", "Satisfies", "satisfies "])
def test_add_attribute_rejects_duplicate_names_case_insensitively(service, name):
    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "A")
    mod_b = create_module(service, project.id, "Module B", "B")
    service.create_link_configuration(project.id, "Satisfies", mod_a.id, mod_b.id)
    with pytest.raises(ValueError, match="already used"):
        service.add_attribute_definition(project.id, "custom", name)


@pytest.mark.parametrize("name", ["", "   ", "\t"])
def test_add_attribute_rejects_empty_names(service, name):
    project = create_project(service)
    with pytest.raises(ValueError, match="must not be empty"):
        service.add_attribute_definition(project.id, "custom", name)


def test_add_attribute_stores_stripped_name(service, tmp_path):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "  Priority  ")
    reloaded = ProjectService(ProjectStore(tmp_path)).get_project(project.id)
    assert reloaded.attributes[-1].name == "Priority"


def test_add_attribute_does_not_rewrite_legacy_colliding_names(service, tmp_path):
    project = create_project(service)
    legacy = service.get_project(project.id)
    from reqman.models import AttributeDefinition

    legacy.attributes.append(AttributeDefinition(id="parent_attr", name="Parent"))
    service.store.save_project(legacy)

    service.add_attribute_definition(project.id, "priority", "Priority")

    names = [a.name for a in service.get_project(project.id).attributes]
    assert names == ["ObjectText", "Parent", "Priority"]


def _write_legacy_attributes(service, project_id, *definitions):
    from reqman.models import AttributeDefinition

    project = service.get_project(project_id)
    project.attributes.extend(AttributeDefinition(**d) for d in definitions)
    service.store.save_project(project)


def test_import_csv_legacy_attribute_named_like_reserved_column(service):
    project = create_project(service)
    _write_legacy_attributes(service, project.id, {"id": "parent_attr", "name": "Parent"})
    module = create_module(service, project.id)

    # A single "Parent" column is the reserved column, not the attribute.
    service.import_csv(project.id, module.id, "Id,Parent,ObjectText\nA,,root\nB,A,child\n")
    reqs = _by_text(service, project.id, module.id)
    assert reqs["child"].parent_id == reqs["root"].id
    assert "parent_attr" not in reqs["child"].attributes

    # The project's own export has the header twice: explain how to fix it.
    exported = service.export_csv(project.id)
    other = create_module(service, project.id, "Module Two", "OTH")
    message = _import_error(service, project.id, other.id, exported)
    assert "Duplicate column 'Parent'" in message
    assert "Attribute 'parent_attr' is also named 'Parent'" in message
    assert "parent_attr" in message

    # Heading the attribute's column with its id imports it.
    service.import_csv(project.id, other.id, "Id,Parent,ObjectText,parent_attr\nA,,root,value\n")
    assert _by_text(service, project.id, other.id)["root"].attributes["parent_attr"] == "value"


def test_import_csv_legacy_attributes_sharing_a_name(service):
    project = create_project(service)
    _write_legacy_attributes(
        service, project.id, {"id": "owner_a", "name": "Owner"}, {"id": "owner_b", "name": "Owner"}
    )
    module = create_module(service, project.id)

    message = _import_error(service, project.id, module.id, "ObjectText,Owner\nx,alice\n")
    assert "Column 'Owner' matches several attributes named 'Owner': ['owner_a', 'owner_b']" in message

    exported_header = service.export_csv(project.id).splitlines()[0]
    message = _import_error(service, project.id, module.id, exported_header + "\n,,,x,a,b\n")
    assert "Duplicate column 'Owner'" in message
    assert "['owner_a', 'owner_b']" in message

    service.import_csv(project.id, module.id, "ObjectText,owner_a,owner_b\nx,alice,bob\n")
    assert _by_text(service, project.id, module.id)["x"].attributes == {
        "object_text": "x", "owner_a": "alice", "owner_b": "bob",
    }


# --- F2: write atomicity -------------------------------------------------------


def _requirement_files(tmp_path, project_id, module_id):
    req_dir = tmp_path / project_id / "modules" / module_id / "requirements"
    return sorted(p.name for p in req_dir.iterdir())


def _fail_on_nth_save(monkeypatch, store, n):
    original = store.save_requirement
    calls = {"count": 0}

    def flaky(project_id, module_id, requirement):
        calls["count"] += 1
        if calls["count"] == n:
            raise OSError("disk full")
        original(project_id, module_id, requirement)

    monkeypatch.setattr(store, "save_requirement", flaky)
    return lambda: monkeypatch.setattr(store, "save_requirement", original)


@pytest.mark.parametrize("fail_on", [1, 2, 3])
def test_import_csv_rolls_back_when_a_write_fails(service, tmp_path, monkeypatch, fail_on):
    project = create_project(service)
    module = create_module(service, project.id)
    existing = service.add_requirement(project.id, module.id, {"object_text": "Existing"})
    files_before = _requirement_files(tmp_path, project.id, module.id)
    counters_before = service.get_project(project.id).id_counters

    restore = _fail_on_nth_save(monkeypatch, service.store, fail_on)
    with pytest.raises(OSError, match="disk full"):
        service.import_csv(project.id, module.id, "Id,Parent,ObjectText\nA,,a\nB,A,b\nC,,c\n")
    restore()

    assert _requirement_files(tmp_path, project.id, module.id) == files_before
    assert service.get_project(project.id).id_counters == counters_before

    result = service.import_csv(project.id, module.id, "ObjectText\nd\ne\n")
    assert existing.id not in result["requirement_ids"]
    assert len(set(result["requirement_ids"])) == 2
    assert len(service.list_requirements(project.id, module.id)) == 3


def test_import_csv_rolls_back_when_saving_the_project_fails(service, tmp_path, monkeypatch):
    project = create_project(service)
    module = create_module(service, project.id)
    files_before = _requirement_files(tmp_path, project.id, module.id)

    def broken(_project):
        raise OSError("read-only")

    monkeypatch.setattr(service.store, "save_project", broken)
    with pytest.raises(OSError, match="read-only"):
        service.import_csv(project.id, module.id, "ObjectText\na\nb\n")
    monkeypatch.undo()

    assert _requirement_files(tmp_path, project.id, module.id) == files_before
    assert service.get_project(project.id).id_counters == {}


@pytest.mark.parametrize("fail_on", [1, 2, 3])
def test_copy_requirement_rolls_back_when_a_write_fails(service, tmp_path, monkeypatch, fail_on):
    project = create_project(service)
    module = create_module(service, project.id)
    root = service.add_requirement(project.id, module.id, {"object_text": "Root"})
    child = service.add_requirement(project.id, module.id, {"object_text": "Child"}, parent_id=root.id)
    service.add_requirement(project.id, module.id, {"object_text": "Grandchild"}, parent_id=child.id)
    files_before = _requirement_files(tmp_path, project.id, module.id)
    counters_before = service.get_project(project.id).id_counters

    restore = _fail_on_nth_save(monkeypatch, service.store, fail_on)
    with pytest.raises(OSError, match="disk full"):
        service.copy_requirement(project.id, module.id, root.id, None)
    restore()

    assert _requirement_files(tmp_path, project.id, module.id) == files_before
    assert service.get_project(project.id).id_counters == counters_before

    copy = service.copy_requirement(project.id, module.id, root.id, None)
    assert copy.id == "REQ-0004"
    assert len(service.list_requirements(project.id, module.id)) == 6


def test_allocation_skips_ids_already_on_disk_after_a_crash(service, tmp_path):
    from reqman.models import Requirement

    project = create_project(service)
    module = create_module(service, project.id)
    other = create_module(service, project.id, "Module Two", "REQ")
    service.add_requirement(project.id, module.id, {"object_text": "First"})
    # Simulate a crash after a requirement file was written but before the
    # project (with its advanced counter) was saved.
    service.store.save_requirement(project.id, other.id, Requirement(id="REQ-0002", attributes={"object_text": "x"}))
    assert service.get_project(project.id).id_counters == {"REQ": 1}

    assert service.add_requirement(project.id, module.id, {"object_text": "Next"}).id == "REQ-0003"
    assert service.import_csv(project.id, module.id, "ObjectText\ny\n")["requirement_ids"] == ["REQ-0004"]


def test_store_writes_atomically_and_keeps_json_format(service, tmp_path, monkeypatch):
    import os

    project = create_project(service)
    module = create_module(service, project.id)
    requirement = service.add_requirement(project.id, module.id, {"object_text": "Original"})
    path = tmp_path / project.id / "modules" / module.id / "requirements" / f"{requirement.id}.json"
    original_text = path.read_text()
    assert original_text == json.dumps(json.loads(original_text), indent=2) + "\n"
    assert list(json.loads(original_text)) == ["id", "attributes", "parent_id", "order"]

    def failing_replace(src, dst):
        raise OSError("replace failed")

    monkeypatch.setattr(os, "replace", failing_replace)
    with pytest.raises(OSError, match="replace failed"):
        service.update_requirement(project.id, module.id, requirement.id, {"object_text": "Changed"})
    monkeypatch.undo()

    assert path.read_text() == original_text
    assert sorted(p.name for p in path.parent.iterdir()) == [path.name]


# --- F3: one defaults/validation path ------------------------------------------


def _typed_project(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    service.add_attribute_definition(project.id, "status", "Status", type="enum", choices=["Draft", "Done"])
    service.add_attribute_definition(project.id, "notes", "Notes", type="string")
    module = create_module(service, project.id)
    return project, module


@pytest.mark.parametrize(
    "attributes, expected",
    [
        ({"priority": "5"}, "Priority: '5' is not an integer"),
        ({"priority": 1.5}, "Priority: '1.5' is not an integer"),
        ({"priority": True}, "Priority: 'True' is not an integer"),
        ({"safety": "true"}, "Safety: 'true' is not a boolean"),
        ({"safety": 1}, "Safety: '1' is not a boolean"),
        ({"status": "draft"}, "Status: 'draft' is not one of ['Draft', 'Done']"),
        ({"status": ""}, "Status: '' is not one of ['Draft', 'Done']"),
        ({"notes": 5}, "Notes: expected a string, got int"),
        ({"object_text": ["x"]}, "ObjectText: expected a string, got list"),
    ],
)
def test_add_and_update_requirement_reject_wrong_value_types(service, attributes, expected):
    project, module = _typed_project(service)
    with pytest.raises(ValueError) as excinfo:
        service.add_requirement(project.id, module.id, {"object_text": "x", **attributes})
    assert expected in str(excinfo.value)
    assert service.list_requirements(project.id, module.id) == []

    requirement = service.add_requirement(project.id, module.id, {"object_text": "ok"})
    with pytest.raises(ValueError) as excinfo:
        service.update_requirement(project.id, module.id, requirement.id, attributes)
    assert expected in str(excinfo.value)
    assert service.get_requirement(project.id, module.id, requirement.id).attributes == {"object_text": "ok"}


def test_add_and_update_requirement_accept_valid_typed_values_and_null(service):
    project, module = _typed_project(service)
    requirement = service.add_requirement(
        project.id,
        module.id,
        {"object_text": "x", "priority": -3, "safety": False, "status": "Done", "notes": ""},
    )
    assert requirement.attributes["priority"] == -3

    updated = service.update_requirement(
        project.id, module.id, requirement.id, {"priority": None, "safety": None, "status": None, "notes": None}
    )
    assert updated.attributes["priority"] is None and updated.attributes["status"] is None


def test_add_requirement_excludes_system_attribute_defaults(service):
    from reqman.models import AttributeDefinition

    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "A")
    mod_b = create_module(service, project.id, "Module B", "B")
    config = service.create_link_configuration(project.id, "Satisfies", mod_a.id, mod_b.id)
    stored = service.get_project(project.id)
    stored.attributes = [
        AttributeDefinition(**{**a.to_dict(), "default": "> stale"}) if a.id == config.id else a
        for a in stored.attributes
    ]
    service.store.save_project(stored)

    requirement = service.add_requirement(project.id, mod_a.id, {"object_text": "x"})
    service.import_csv(project.id, mod_a.id, "ObjectText\ny\n")

    assert requirement.attributes == {"object_text": "x"}
    assert all(config.id not in r.attributes for r in service.list_requirements(project.id, mod_a.id))


def test_add_requirement_required_message_uses_attribute_name(service):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(ValueError, match="required attribute 'ObjectText' is empty"):
        service.add_requirement(project.id, module.id, {})


def test_import_csv_reports_enum_and_required_errors_for_every_row(service):
    project, module = _typed_project(service)
    service.add_attribute_definition(project.id, "owner", "Owner", required=True)

    message = _import_error(
        service, project.id, module.id, "ObjectText,Status,Owner\na,bad,\nb,Draft,\nc,worse,bob\n"
    )
    assert "Row 2: Status: 'bad' is not one of ['Draft', 'Done']" in message
    assert "Row 2: required attribute 'Owner' is empty" in message
    assert "Row 3: required attribute 'Owner' is empty" in message
    assert "Row 4: Status: 'worse' is not one of ['Draft', 'Done']" in message


# --- F4/F5: XLSX export ---------------------------------------------------------


def _xlsx_rows(service, project_id):
    import io

    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(service.export_xlsx(project_id))).active
    return sheet


def test_export_xlsx_never_writes_formulas(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", type="integer")
    service.add_attribute_definition(project.id, "safety", "Safety", type="boolean")
    module = create_module(service, project.id)
    service.add_requirement(project.id, module.id, {"object_text": "=1+1", "priority": 7, "safety": True})
    service.add_requirement(project.id, module.id, {"object_text": '=HYPERLINK("http://x")'})

    sheet = _xlsx_rows(service, project.id)
    texts = [sheet.cell(row=r, column=4) for r in (2, 3)]
    assert [(c.value, c.data_type) for c in texts] == [("=1+1", "s"), ('=HYPERLINK("http://x")', "s")]
    assert (sheet.cell(row=2, column=5).value, sheet.cell(row=2, column=5).data_type) == (7, "n")
    assert (sheet.cell(row=2, column=6).value, sheet.cell(row=2, column=6).data_type) == (True, "b")
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)


def test_export_xlsx_row_order_matches_csv_for_reordered_nested_project(service):
    import csv as csv_module
    import io

    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "A")
    mod_b = create_module(service, project.id, "Module B", "B")
    a = service.add_requirement(project.id, mod_a.id, {"object_text": "A"})
    service.add_requirement(project.id, mod_a.id, {"object_text": "B"})
    c = service.add_requirement(project.id, mod_a.id, {"object_text": "C"})
    service.add_requirement(project.id, mod_a.id, {"object_text": "A1"}, parent_id=a.id)
    a2 = service.add_requirement(project.id, mod_a.id, {"object_text": "A2"}, parent_id=a.id)
    service.add_requirement(project.id, mod_a.id, {"object_text": "A2x"}, parent_id=a2.id)
    service.reorder_requirement(project.id, mod_a.id, c.id, "up")
    service.reorder_requirement(project.id, mod_a.id, c.id, "up")
    service.reorder_requirement(project.id, mod_a.id, a2.id, "up")
    b1 = service.add_requirement(project.id, mod_b.id, {"object_text": "B1"})
    service.add_requirement(project.id, mod_b.id, {"object_text": "B2"})
    service.reorder_requirement(project.id, mod_b.id, b1.id, "down")

    csv_rows = list(csv_module.reader(io.StringIO(service.export_csv(project.id))))
    xlsx_rows = [list(row) for row in _xlsx_rows(service, project.id).iter_rows(values_only=True)]

    def normalize(rows):
        return [["" if v is None else str(v) for v in row] for row in rows]

    assert normalize(xlsx_rows) == normalize(csv_rows)
    assert [row[3] for row in csv_rows[1:]] == ["C", "A", "A2", "A2x", "A1", "B", "B2", "B1"]


# --- R1: link configuration names ----------------------------------------------


@pytest.mark.parametrize(
    "name, expected",
    [
        ("Parent", "reserved"),
        (" id ", "reserved"),
        ("MODULE", "reserved"),
        ("ObjectText", "already used by attribute 'object_text'"),
        ("priority ", "already used by attribute 'priority'"),
        ("satisfies", "already used by attribute 'satisfies'"),
        ("", "must not be empty"),
        ("   ", "must not be empty"),
    ],
)
def test_create_link_configuration_rejects_reserved_duplicate_and_empty_names(service, name, expected):
    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "A")
    mod_b = create_module(service, project.id, "Module B", "B")
    service.add_attribute_definition(project.id, "priority", "Priority")
    service.create_link_configuration(project.id, "Satisfies", mod_a.id, mod_b.id)
    before = service.get_project(project.id).to_dict()

    with pytest.raises(ValueError, match=expected) as excinfo:
        service.create_link_configuration(project.id, name, mod_a.id, mod_b.id)
    assert str(excinfo.value).startswith("Link name")
    assert service.get_project(project.id).to_dict() == before


def test_create_link_configuration_stores_stripped_name_and_export_roundtrips(service):
    project = create_project(service)
    mod_a = create_module(service, project.id, "Module A", "A")
    mod_b = create_module(service, project.id, "Module B", "B")
    config = service.create_link_configuration(project.id, "  Verifies  ", mod_a.id, mod_b.id)
    assert config.name == "Verifies"
    assert service.get_project(project.id).attributes[-1].name == "Verifies"

    source = service.add_requirement(project.id, mod_a.id, {"object_text": "src"})
    target = service.add_requirement(project.id, mod_b.id, {"object_text": "tgt"})
    service.link_requirements(project.id, config.id, source.id, target.id)
    exported = service.export_csv(project.id)
    assert exported.splitlines()[0] == "Id,Module,Parent,ObjectText,Verifies"
    header = exported.splitlines()[0]
    other = create_module(service, project.id, "Module C", "C")
    result = service.import_csv(project.id, other.id, f"{header}\n,,,copied,> {target.id}\n")
    assert result["ignored_columns"] == ["Verifies"]


# --- R2: attribute defaults and choices ----------------------------------------


@pytest.mark.parametrize(
    "kwargs, expected",
    [
        ({"type": "integer", "default": "5"}, "Invalid default value: Priority: '5' is not an integer"),
        ({"type": "integer", "default": True}, "Invalid default value: Priority: 'True' is not an integer"),
        ({"type": "boolean", "default": "true"}, "Invalid default value: Priority: 'true' is not a boolean"),
        ({"type": "string", "default": 3}, "Invalid default value: Priority: expected a string, got int"),
        ({"type": "text", "default": ["x"]}, "Invalid default value: Priority: expected a string, got list"),
        (
            {"type": "enum", "choices": ["Low", "High"], "default": "Medium"},
            "Invalid default value: Priority: 'Medium' is not one of ['Low', 'High']",
        ),
        ({"type": "enum"}, "Enum attribute 'Priority' needs a non-empty list of choices"),
        ({"type": "enum", "choices": []}, "Enum attribute 'Priority' needs a non-empty list of choices"),
        ({"type": "enum", "choices": ["Low", 3]}, "Enum attribute 'Priority' needs a non-empty list of choices"),
        ({"type": "string", "choices": ["a"]}, "Attribute 'Priority' of type 'string' cannot have choices"),
        ({"type": "integer", "choices": []}, "Attribute 'Priority' of type 'integer' cannot have choices"),
    ],
)
def test_add_attribute_rejects_invalid_default_and_choices(service, kwargs, expected):
    project = create_project(service)
    with pytest.raises(ValueError) as excinfo:
        service.add_attribute_definition(project.id, "priority", "Priority", **kwargs)
    assert str(excinfo.value) == expected
    assert [a.id for a in service.get_project(project.id).attributes] == ["object_text"]


@pytest.mark.parametrize(
    "kwargs",
    [
        {"type": "integer", "default": -2},
        {"type": "boolean", "default": False},
        {"type": "string", "default": ""},
        {"type": "text", "default": "x"},
        {"type": "enum", "choices": ["Low", "High"], "default": "High"},
        {"type": "enum", "choices": ["Low"]},
        {"type": "integer"},
    ],
)
def test_add_attribute_accepts_valid_default_and_choices(service, kwargs):
    project = create_project(service)
    service.add_attribute_definition(project.id, "priority", "Priority", **kwargs)
    module = create_module(service, project.id)
    requirement = service.add_requirement(project.id, module.id, {"object_text": "x"})
    assert requirement.attributes.get("priority") == kwargs.get("default")


def test_legacy_invalid_defaults_are_skipped_not_rewritten(service, tmp_path):
    project = create_project(service)
    _write_legacy_attributes(
        service,
        project.id,
        {"id": "priority", "name": "Priority", "type": "integer", "default": "high"},
        {"id": "status", "name": "Status", "type": "enum", "choices": ["Draft"], "default": "Gone"},
        {"id": "owner", "name": "Owner", "type": "string", "default": "team"},
    )
    project_file = tmp_path / project.id / "project.json"
    stored_before = json.loads(project_file.read_text())["attributes"]
    module = create_module(service, project.id)

    requirement = service.add_requirement(project.id, module.id, {"object_text": "x"})
    service.import_csv(project.id, module.id, "ObjectText\ny\n")

    expected = {"object_text": "x", "owner": "team"}
    assert requirement.attributes == expected
    assert _by_text(service, project.id, module.id)["y"].attributes == {"object_text": "y", "owner": "team"}
    assert json.loads(project_file.read_text())["attributes"] == stored_before


def test_legacy_invalid_default_of_required_attribute_counts_as_no_default(service):
    project = create_project(service)
    _write_legacy_attributes(
        service, project.id, {"id": "priority", "name": "Priority", "type": "integer", "required": True, "default": "x"}
    )
    module = create_module(service, project.id)
    with pytest.raises(ValueError, match="required attribute 'Priority' is empty"):
        service.add_requirement(project.id, module.id, {"object_text": "x"})
    assert service.add_requirement(project.id, module.id, {"object_text": "x", "priority": 1}).attributes == {
        "object_text": "x", "priority": 1,
    }


@pytest.mark.parametrize("type_", ["string", "text"])
@pytest.mark.parametrize("default", ["", "   ", "\n\t"])
def test_add_attribute_rejects_blank_default_of_required_attribute(service, type_, default):
    project = create_project(service)
    with pytest.raises(ValueError) as excinfo:
        service.add_attribute_definition(project.id, "owner", "Owner", type=type_, required=True, default=default)
    assert str(excinfo.value) == "Attribute 'Owner' is required, so its default must not be empty"
    assert [a.id for a in service.get_project(project.id).attributes] == ["object_text"]


def test_add_attribute_accepts_blank_default_of_optional_and_no_default_of_required_attribute(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "notes", "Notes", type="text", default="  ")
    service.add_attribute_definition(project.id, "owner", "Owner", required=True)
    assert [a.id for a in service.get_project(project.id).attributes] == ["object_text", "notes", "owner"]


@pytest.mark.parametrize("type_", ["Integer", "Enum", "str", "", "datetime", "Date"])
def test_add_attribute_rejects_unknown_type(service, type_):
    project = create_project(service)
    with pytest.raises(ValueError) as excinfo:
        service.add_attribute_definition(project.id, "priority", "Priority", type=type_, default="x")
    assert str(excinfo.value) == (
        f"Unknown attribute type '{type_}'; expected one of string, text, integer, boolean, enum, date"
    )
    assert [a.id for a in service.get_project(project.id).attributes] == ["object_text"]


def test_legacy_unknown_attribute_type_is_kept_and_not_rewritten(service, tmp_path):
    project = create_project(service)
    _write_legacy_attributes(service, project.id, {"id": "size", "name": "Size", "type": "Integer"})
    project_file = tmp_path / project.id / "project.json"
    stored_before = json.loads(project_file.read_text())["attributes"]
    module = create_module(service, project.id)

    requirement = service.add_requirement(project.id, module.id, {"object_text": "x", "size": "big"})
    assert requirement.attributes == {"object_text": "x", "size": "big"}
    assert service.get_project(project.id).attributes[-1].type == "Integer"
    assert json.loads(project_file.read_text())["attributes"] == stored_before


# --- R4: blank values for required attributes -----------------------------------


@pytest.mark.parametrize("value", [None, "", "   ", "\n\t"])
def test_add_requirement_rejects_blank_required_value(service, value):
    project = create_project(service)
    module = create_module(service, project.id)
    with pytest.raises(ValueError) as excinfo:
        service.add_requirement(project.id, module.id, {"object_text": value})
    assert str(excinfo.value) == "required attribute 'ObjectText' is empty"
    assert service.list_requirements(project.id, module.id) == []


@pytest.mark.parametrize("value", [None, "", "  "])
def test_add_requirement_explicit_blank_overrides_default_of_required_attribute(service, value):
    project = create_project(service)
    service.add_attribute_definition(project.id, "owner", "Owner", required=True, default="team")
    module = create_module(service, project.id)

    with pytest.raises(ValueError, match="required attribute 'Owner' is empty"):
        service.add_requirement(project.id, module.id, {"object_text": "x", "owner": value})
    assert service.add_requirement(project.id, module.id, {"object_text": "x"}).attributes["owner"] == "team"


@pytest.mark.parametrize("value", [None, "", "   "])
def test_update_requirement_rejects_blank_required_value(service, value):
    project = create_project(service)
    module = create_module(service, project.id)
    requirement = service.add_requirement(project.id, module.id, {"object_text": "keep"})

    with pytest.raises(ValueError) as excinfo:
        service.update_requirement(project.id, module.id, requirement.id, {"object_text": value})
    assert str(excinfo.value) == "required attribute 'ObjectText' is empty"
    assert service.get_requirement(project.id, module.id, requirement.id).attributes == {"object_text": "keep"}


def test_update_requirement_ignores_required_attributes_not_being_written(service):
    project = create_project(service)
    module = create_module(service, project.id)
    requirement = service.add_requirement(project.id, module.id, {"object_text": "keep"})
    service.add_attribute_definition(project.id, "owner", "Owner", required=True)
    service.add_attribute_definition(project.id, "notes", "Notes")

    updated = service.update_requirement(project.id, module.id, requirement.id, {"notes": ""})
    assert updated.attributes == {"object_text": "keep", "notes": ""}


def test_required_enum_with_invalid_value_reports_only_the_value_problem(service):
    project = create_project(service)
    service.add_attribute_definition(project.id, "status", "Status", type="enum", required=True, choices=["Draft"])
    module = create_module(service, project.id)
    with pytest.raises(ValueError) as excinfo:
        service.add_requirement(project.id, module.id, {"object_text": "x", "status": ""})
    assert str(excinfo.value) == "Status: '' is not one of ['Draft']"


@pytest.mark.parametrize(
    "type_, value, cell",
    [("boolean", False, "false"), ("integer", 0, "0")],
)
def test_required_false_and_zero_count_as_values(service, type_, value, cell):
    project = create_project(service)
    service.add_attribute_definition(project.id, "flag", "Flag", type=type_, required=True)
    module = create_module(service, project.id)

    created = service.add_requirement(project.id, module.id, {"object_text": "a", "flag": value})
    assert created.attributes == {"object_text": "a", "flag": value}

    other = service.add_requirement(
        project.id, module.id, {"object_text": "b", "flag": True if type_ == "boolean" else 1}
    )
    updated = service.update_requirement(project.id, module.id, other.id, {"flag": value})
    assert updated.attributes == {"object_text": "b", "flag": value}
    assert type(updated.attributes["flag"]) is type(value)

    service.import_csv(project.id, module.id, f"ObjectText,Flag\nc,{cell}\n")
    imported = _by_text(service, project.id, module.id)["c"].attributes
    assert imported == {"object_text": "c", "flag": value}
    assert type(imported["flag"]) is type(value)


# --- R5: durable writes ----------------------------------------------------------


def test_store_syncs_temporary_file_before_replace(service, monkeypatch):
    import os

    from reqman import store as store_module

    events: list[str] = []
    real_fsync, real_replace = os.fsync, os.replace

    def fsync(fd):
        events.append("sync")
        real_fsync(fd)

    def replace(src, dst):
        events.append("replace")
        real_replace(src, dst)

    monkeypatch.setattr(os, "fsync", fsync)
    monkeypatch.setattr(os, "replace", replace)
    fcntl = store_module.fcntl
    if fcntl is not None and hasattr(fcntl, "F_FULLFSYNC"):
        real_fcntl = fcntl.fcntl

        def full_fsync(fd, cmd, *args):
            if cmd == fcntl.F_FULLFSYNC:
                events.append("sync")
            return real_fcntl(fd, cmd, *args)

        monkeypatch.setattr(fcntl, "fcntl", full_fsync)
    create_project(service)
    assert events == ["sync", "replace"]


@pytest.mark.parametrize(
    "has_full_fsync, full_fsync_fails, expected",
    [
        (True, False, ["F_FULLFSYNC"]),
        (True, True, ["F_FULLFSYNC", "fsync"]),
        (False, False, ["fsync"]),
    ],
)
def test_store_sync_prefers_full_fsync_and_falls_back_to_fsync(
    monkeypatch, has_full_fsync, full_fsync_fails, expected
):
    import os
    import types

    from reqman import store as store_module

    calls: list[str] = []

    def fake_fcntl(fd, cmd):
        calls.append("F_FULLFSYNC")
        if full_fsync_fails:
            raise OSError("not supported")

    fake = types.SimpleNamespace(fcntl=fake_fcntl)
    if has_full_fsync:
        fake.F_FULLFSYNC = 51
    monkeypatch.setattr(store_module, "fcntl", fake)
    monkeypatch.setattr(os, "fsync", lambda fd: calls.append("fsync"))

    store_module._sync(0)
    assert calls == expected
