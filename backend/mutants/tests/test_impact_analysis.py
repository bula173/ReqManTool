"""Impact analysis: the transitive upstream/downstream walk from selected requirements.

See docs/ARCHITECTURE.md "Impact analysis". Like the other reports it is
read-only: no file is written and nothing is broadcast.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import csv
import importlib
import io
import json

import openpyxl
import pytest
from docx import Document
from fastapi.testclient import TestClient

from reqman.service import MAX_IMPACT_DEPTH, MAX_IMPACT_SEEDS, ProjectService
from reqman.store import ProjectStore


@pytest.fixture
def service(tmp_path):
    return ProjectService(ProjectStore(tmp_path))


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REQMAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(tmp_path / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app)


def _rpc(ws, method, params, request_id=1):
    ws.send_text(json.dumps({"jsonrpc": "2.0", "id": request_id, "method": method, "params": params}))
    return json.loads(ws.receive_text())


def _snapshot(directory: Path) -> dict[str, bytes]:
    return {str(p.relative_to(directory)): p.read_bytes() for p in sorted(directory.rglob("*")) if p.is_file()}


class Lab:
    """Base ← Derived ← Detail, plus a test case verifying the Derived module.

    ``Satisfies`` (Derived → Base) and ``Refines`` (Detail → Derived) point
    downstream → upstream like the RBC configurations, ``Verifies``
    (Tests → Derived) links a test case to a requirement.
    """

    def __init__(self, service: ProjectService):
        self.service = service
        project = service.create_project("Impact")
        self.pid = project.id
        service.add_attribute_definition(
            self.pid, "object_type", "Object Type", type="enum", choices=["Heading", "Requirement"],
        )
        service.add_attribute_definition(
            self.pid, "sil", "SIL", type="enum", choices=["SIL 1", "SIL 2", "SIL 4"],
        )
        service.add_attribute_definition(self.pid, "safety_related", "Safety Related", type="boolean")
        self.base = service.create_module(self.pid, "Base", "BAS").id
        self.derived = service.create_module(self.pid, "Derived", "DER").id
        self.detail = service.create_module(self.pid, "Detail", "DTL").id
        self.tests = service.create_module(self.pid, "Tests", "TC", type="test_cases").id
        self.satisfies = service.create_link_configuration(self.pid, "Satisfies", self.derived, self.base).id
        self.refines = service.create_link_configuration(self.pid, "Refines", self.detail, self.derived).id
        self.verifies = service.create_link_configuration(self.pid, "Verifies", self.tests, self.derived).id

    def add(self, module_id: str, text: str, **values) -> str:
        return self.service.add_requirement(self.pid, module_id, {"object_text": text, **values}).id

    def link(self, config_id: str, source_id: str, target_id: str) -> None:
        self.service.link_requirements(self.pid, config_id, source_id, target_id)


@pytest.fixture
def lab(service):
    lab = Lab(service)
    # A heading of the Derived module: left out of the analysis, like every report.
    lab.heading = lab.add(lab.derived, "Chapter", object_type="Heading")
    lab.b1 = lab.add(lab.base, "Base requirement", sil="SIL 4", safety_related=True)
    lab.d1 = lab.add(lab.derived, "Derived requirement", sil="SIL 2")
    lab.x1 = lab.add(lab.detail, "Detailed requirement")
    lab.t1 = lab.add(lab.tests, "Test case")
    lab.link(lab.satisfies, lab.d1, lab.b1)
    lab.link(lab.refines, lab.x1, lab.d1)
    lab.link(lab.refines, lab.x1, lab.heading)  # links to excluded objects are ignored
    lab.link(lab.verifies, lab.t1, lab.d1)
    lab.link(lab.verifies, lab.t1, lab.heading)
    return lab


def _rows(report):
    return [(row["direction"], row["distance"], row["via"], tuple(row["path"]), row["id"]) for row in report["results"]]


# --- the walk ---------------------------------------------------------------------------


def test_impact_upstream_walk_has_distance_and_path(lab, service):
    report = service.impact_analysis(lab.pid, [lab.x1], 1, "upstream")
    assert _rows(report) == [("upstream", 1, lab.x1, (lab.x1, lab.d1), lab.d1)]
    assert report["summary"] == {
        "seeds": 1, "reached": 1, "upstream": 1, "downstream": 0, "max_distance": 1,
    }

    report = service.impact_analysis(lab.pid, [lab.x1], 2, "upstream")
    assert _rows(report) == [
        ("upstream", 1, lab.x1, (lab.x1, lab.d1), lab.d1),
        ("upstream", 2, lab.d1, (lab.x1, lab.d1, lab.b1), lab.b1),
    ]
    assert report["summary"]["max_distance"] == 2


def test_impact_downstream_walk_reaches_the_test_case(lab, service):
    report = service.impact_analysis(lab.pid, [lab.b1], 2, "downstream")
    assert _rows(report) == [
        ("downstream", 1, lab.b1, (lab.b1, lab.d1), lab.d1),
        # Distance 2, in module order: Detail before Tests.
        ("downstream", 2, lab.d1, (lab.b1, lab.d1, lab.x1), lab.x1),
        ("downstream", 2, lab.d1, (lab.b1, lab.d1, lab.t1), lab.t1),
    ]
    assert report["summary"] == {
        "seeds": 1, "reached": 3, "upstream": 0, "downstream": 3, "max_distance": 2,
    }


def test_impact_both_directions_and_start_order(lab, service):
    report = service.impact_analysis(lab.pid, [lab.d1], 5, "both")
    # Nothing is upstream of Base, and a walk never starts from a result.
    assert _rows(report) == [
        ("upstream", 1, lab.d1, (lab.d1, lab.b1), lab.b1),
        ("downstream", 1, lab.d1, (lab.d1, lab.x1), lab.x1),
        ("downstream", 1, lab.d1, (lab.d1, lab.t1), lab.t1),
    ]
    assert [seed["id"] for seed in report["seeds"]] == [lab.d1]
    assert report["direction"] == "both" and report["depth"] == 5


def test_impact_never_walks_back_to_a_start(lab, service):
    # A cycle: the base requirement also links back to the derived one.
    cycle = service.create_link_configuration(lab.pid, "Refined by", lab.base, lab.derived).id
    service.link_requirements(lab.pid, cycle, lab.b1, lab.d1)
    report = service.impact_analysis(lab.pid, [lab.b1], MAX_IMPACT_DEPTH, "both")
    # Upstream stops at the seed, downstream continues past it to Detail and the test case.
    assert [row["id"] for row in report["results"]] == [lab.d1, lab.d1, lab.x1, lab.t1]
    assert lab.b1 not in {row["id"] for row in report["results"]}


def test_impact_leaves_out_headings(lab, service):
    report = service.impact_analysis(lab.pid, [lab.x1], MAX_IMPACT_DEPTH, "both")
    assert lab.heading not in {row["id"] for row in report["results"]}
    # The Detail -> heading link is ignored, so the heading is never a step either.
    assert [row["id"] for row in report["results"]] == [lab.d1, lab.b1]


def test_impact_lists_sil_safety_related_and_linked_tests(lab, service):
    report = service.impact_analysis(lab.pid, [lab.x1], 2, "upstream")
    by_id = {row["id"]: row for row in report["results"]}
    assert by_id[lab.d1]["sil"] == "SIL 2"
    assert by_id[lab.d1]["safety_related"] is None
    assert by_id[lab.d1]["tests"] == [{"id": lab.t1, "module_id": lab.tests, "text": "Test case"}]
    assert by_id[lab.b1]["sil"] == "SIL 4"
    assert by_id[lab.b1]["safety_related"] is True
    assert by_id[lab.b1]["tests"] == []

    seeds = {seed["id"]: seed for seed in report["seeds"]}
    assert seeds[lab.x1]["sil"] is None and seeds[lab.x1]["tests"] == []

    modules = {module["id"]: module for module in report["modules"]}
    assert set(modules) == {lab.detail, lab.derived, lab.base}
    assert modules[lab.derived]["sil_attribute"] == {"id": "sil", "name": "SIL"}
    assert modules[lab.base]["safety_related_attribute"] == {"id": "safety_related", "name": "Safety Related"}
    assert modules[lab.detail]["safety_related_attribute"] == {"id": "safety_related", "name": "Safety Related"}
    assert report["data_source"]["history"] in ("off", "ready", "no_repo")


def test_impact_without_sil_or_safety_related_attributes(service):
    project = service.create_project("Plain")
    first = service.create_module(project.id, "One", "ONE").id
    second = service.create_module(project.id, "Two", "TWO").id
    config = service.create_link_configuration(project.id, "Derives", second, first).id
    a = service.add_requirement(project.id, first, {"object_text": "A"}).id
    b = service.add_requirement(project.id, second, {"object_text": "B"}).id
    service.link_requirements(project.id, config, b, a)
    report = service.impact_analysis(project.id, [b], 3, "both")
    assert report["results"][0]["sil"] is None
    assert report["results"][0]["safety_related"] is None
    assert report["results"][0]["tests"] == []
    modules = {module["id"]: module for module in report["modules"]}
    assert modules[second]["sil_attribute"] is None
    assert modules[second]["safety_related_attribute"] is None


# --- validation --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "requirement_ids, message",
    [
        ([], "Pick at least one starting requirement"),
        (None, "Pick at least one starting requirement"),
        ("x", "Pick at least one starting requirement"),
        (["x", 1], "Every starting requirement must be a requirement id"),
        ([1, 2], "Every starting requirement must be a requirement id"),
        (["   "], "Every starting requirement must be a requirement id"),
        (
            [f"REQ-{i}" for i in range(MAX_IMPACT_SEEDS + 1)],
            f"An impact analysis starts from at most {MAX_IMPACT_SEEDS} requirements",
        ),
    ],
)
def test_impact_rejects_the_starting_requirements(lab, service, requirement_ids, message):
    with pytest.raises(ValueError) as excinfo:
        service.impact_analysis(lab.pid, requirement_ids)
    assert str(excinfo.value) == message


@pytest.mark.parametrize(
    "depth, message",
    [
        (0, "The depth must be from 1 to 10, got 0"),
        (-1, "The depth must be from 1 to 10, got -1"),
        (11, "The depth must be from 1 to 10, got 11"),
        ("abc", "Invalid depth 'abc': expected a number from 1 to 10"),
        ("", "Invalid depth '': expected a number from 1 to 10"),
        (None, "Invalid depth 'None': expected a number from 1 to 10"),
        (True, "Invalid depth 'True': expected a number from 1 to 10"),
        (2.5, "Invalid depth '2.5': expected a number from 1 to 10"),
    ],
)
def test_impact_rejects_the_depth(lab, service, depth, message):
    with pytest.raises(ValueError) as excinfo:
        service.impact_analysis(lab.pid, [lab.x1], depth)
    assert str(excinfo.value) == message


@pytest.mark.parametrize(
    "direction, message",
    [
        ("sideways", "Invalid direction 'sideways': expected one of upstream, downstream, both"),
        ("Upstream", "Invalid direction 'Upstream': expected one of upstream, downstream, both"),
        (None, "Invalid direction 'None': expected one of upstream, downstream, both"),
        (3, "Invalid direction '3': expected one of upstream, downstream, both"),
        (["upstream"], "Invalid direction '['upstream']': expected one of upstream, downstream, both"),
    ],
)
def test_impact_rejects_the_direction(lab, service, direction, message):
    with pytest.raises(ValueError) as excinfo:
        service.impact_analysis(lab.pid, [lab.x1], 1, direction)
    assert str(excinfo.value) == message


def test_impact_accepts_a_numeric_depth_and_deduplicates_the_starts(lab, service):
    report = service.impact_analysis(lab.pid, [f" {lab.x1} ", lab.x1], "1", "upstream")
    assert report["depth"] == 1
    assert [seed["id"] for seed in report["seeds"]] == [lab.x1]


def test_impact_unknown_start_is_not_found(lab, service):
    with pytest.raises(FileNotFoundError, match=f"Requirement 'NOPE-9' not found in project '{lab.pid}'"):
        service.impact_analysis(lab.pid, [lab.x1, "NOPE-9"])


def test_impact_heading_start_is_refused(lab, service):
    with pytest.raises(ValueError) as excinfo:
        service.impact_analysis(lab.pid, [lab.heading])
    assert str(excinfo.value) == (
        f"Requirement '{lab.heading}' is left out of reports (its Object Type is a heading, "
        "information, definition, title, comment or note) and cannot start an impact analysis"
    )


def test_impact_writes_nothing(lab, tmp_path):
    before = _snapshot(tmp_path / lab.pid)
    service = ProjectService(ProjectStore(tmp_path))
    service.impact_analysis(lab.pid, [lab.x1], MAX_IMPACT_DEPTH, "both")
    assert _snapshot(tmp_path / lab.pid) == before


# --- REST / WebSocket ----------------------------------------------------------------------


def _api_lab(client, tmp_path):
    service = ProjectService(ProjectStore(tmp_path))
    project = service.create_project("Impact")
    pid = project.id
    first = service.create_module(pid, "First", "ONE").id
    second = service.create_module(pid, "Two", "TWO").id
    config = service.create_link_configuration(pid, "Derives", second, first).id
    a = service.add_requirement(pid, first, {"object_text": "Base"}).id
    b = service.add_requirement(pid, second, {"object_text": "Derived"}).id
    service.link_requirements(pid, config, b, a)
    return pid, a, b


def test_impact_rest_ws_parity_and_no_broadcast(client, tmp_path):
    pid, a, b = _api_lab(client, tmp_path)
    before = _snapshot(tmp_path / pid)
    url = f"/projects/{pid}/reports/impact?requirement={b}&depth=3&direction=both"
    with client.websocket_connect("/ws") as ws, client.websocket_connect("/ws") as watcher:
        rest = client.get(url)
        assert rest.status_code == 200, rest.text
        assert rest.json()["results"][0]["id"] == a
        response = _rpc(ws, "report.impact", {"project_id": pid, "requirement_ids": [b], "depth": "3"})
        assert response["result"] == rest.json()

        # Nothing was broadcast: the watcher's next event is a later, distinguishable mutation.
        client.post("/projects", json={"name": "Other"})
        assert json.loads(watcher.receive_text()) == {"event": "changed", "project_id": "other"}
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": "other"}
    assert _snapshot(tmp_path / pid) == before


@pytest.mark.parametrize(
    "method, params, url, status, code",
    [
        ("report.impact", {"requirement_ids": []}, "/reports/impact", 400, -32602),
        ("report.impact", {"requirement_ids": ["nope"]}, "/reports/impact?requirement=nope", 404, -32001),
        ("report.impact", {"requirement_ids": ["x"], "depth": 0}, "/reports/impact?requirement=x&depth=0", 400, -32602),
        (
            "report.impact",
            {"requirement_ids": ["x"], "direction": "sideways"},
            "/reports/impact?requirement=x&direction=sideways",
            400,
            -32602,
        ),
        ("report.impact", {"requirement_ids": ["x"], "depth": "deep"}, "/reports/impact?requirement=x&depth=deep", 400, -32602),
    ],
)
def test_impact_errors_rest_and_ws(client, tmp_path, method, params, url, status, code):
    pid, _, _ = _api_lab(client, tmp_path)
    resp = client.get(f"/projects/{pid}{url}")
    assert resp.status_code == status, resp.text
    with client.websocket_connect("/ws") as ws:
        response = _rpc(ws, method, {"project_id": pid, **params})
    assert response["error"]["code"] == code
    assert response["error"]["message"] == resp.json()["detail"]


def test_impact_on_missing_project_is_404(client):
    assert client.get("/projects/missing/reports/impact?requirement=x").status_code == 404


def test_impact_is_a_client_route(client):
    """A page load of the report route gets index.html; its exports stay API paths."""
    from reqman.app import SPA_ROUTE

    assert SPA_ROUTE.match("/projects/p/reports/impact")
    assert SPA_ROUTE.match("/projects/p/reports/impact/export/csv") is None
    assert SPA_ROUTE.match("/projects/p/reports/impact/export/xlsx") is None
    assert SPA_ROUTE.match("/projects/p/reports/impact/export/docx") is None


# --- exports -------------------------------------------------------------------------------


def test_impact_export_headers_and_content(client, tmp_path):
    pid, a, b = _api_lab(client, tmp_path)
    query = f"requirement={b}&depth=3&direction=both"

    resp = client.get(f"/projects/{pid}/reports/impact/export/csv?{query}")
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.headers["content-disposition"] == f'attachment; filename="{pid}-impact.csv"'
    rows = list(csv.reader(io.StringIO(resp.text)))
    assert rows[0] == [
        "Direction", "Distance", "Requirement", "Module", "Text",
        "SIL", "Safety Related", "Linked tests", "Path", "Change reference",
    ]
    assert rows[1] == ["Starting requirement", "0", b, "Two", "Derived", "", "", "", b, ""]
    assert rows[2] == ["upstream 1", "1", a, "First", "Base", "", "", "", f"{b} → {a}", ""]
    assert rows[2][-1] == ""  # the change reference is left to fill in

    resp = client.get(f"/projects/{pid}/reports/impact/export/xlsx?{query}")
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Summary", "Requirements"]
    summary = [[c.value for c in row] for row in workbook["Summary"].iter_rows()]
    assert ["Direction", "both"] in [row[:2] for row in summary]
    assert ["Depth", 3] in [row[:2] for row in summary]
    assert ["Data", "Current data (no git history)"] in [row[:2] for row in summary]
    requirements = [[c.value for c in row] for row in workbook["Requirements"].iter_rows()]
    assert requirements[0][-1] == "Change reference"
    assert requirements[2][-1] in (None, "")  # empty until the user fills it in
    assert requirements[2][2] == a

    resp = client.get(f"/projects/{pid}/reports/impact/export/docx?{query}")
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    document = Document(io.BytesIO(resp.content))
    assert document.paragraphs[0].text == "Impact analysis (both)"
    assert any(p.text.startswith("Data: Current data") for p in document.paragraphs)
    affected = document.tables[-1]
    assert affected.rows[0].cells[-1].text == "Change reference"
    assert affected.rows[1].cells[-1].text == ""
    assert affected.rows[1].cells[2].text == a


def test_impact_export_wrong_format_is_400(client, tmp_path):
    pid, _, b = _api_lab(client, tmp_path)
    resp = client.get(f"/projects/{pid}/reports/impact/export/pdf?requirement={b}")
    assert resp.status_code == 400
    assert resp.json()["detail"] == "Unknown impact analysis export format 'pdf'; expected csv, xlsx or docx"
