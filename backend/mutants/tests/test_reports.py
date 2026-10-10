"""Reports: traceability matrix, verification report, test execution report, verification status.

See docs/ARCHITECTURE.md "Reports" and "Verification status". The reports are
read-only: they never write a file and never broadcast.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import csv
import importlib
import io
import json

import openpyxl
import pytest
from docx import Document
from fastapi.testclient import TestClient

from reqman.service import MAX_TRACE_CHAIN, REPORT_EXCLUDED_OBJECT_TYPES, ProjectService
from reqman.store import ProjectStore

RBC = "builtin:rbc_template"


@pytest.fixture
def service(tmp_path):
    return ProjectService(ProjectStore(tmp_path))


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("REQMAN_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REQMAN_CONFIG_FILE", str(tmp_path / "config.json"))
    import reqman.app as app_module

    importlib.reload(app_module)
    return TestClient(app_module.app)


def _rpc(ws, method, params, request_id=1):
    ws.send_text(json.dumps({"jsonrpc": "2.0", "id": request_id, "method": method, "params": params}))
    return json.loads(ws.receive_text())


def _snapshot(directory: Path) -> dict[str, bytes]:
    return {str(p.relative_to(directory)): p.read_bytes() for p in sorted(directory.rglob("*")) if p.is_file()}


class Lab:
    """A project with a requirements, a test cases and a test execution module.

    Attribute names are deliberately not the RBC ones: the reports find them
    by role ("Verdict" = result, "Run Date" = execution date, "Tester" =
    executed by, "Build" = version, "Log" = evidence, "Kind" is not "Object
    Type" so the object type attribute is called "Object Type").
    """

    def __init__(self, service: ProjectService, name: str = "Lab"):
        self.service = service
        project = service.create_project(name)
        self.pid = project.id
        service.add_attribute_definition(
            self.pid, "object_type", "Object Type", type="enum",
            choices=["Heading", "Requirement", "Information", "Test Case", "Note"],
        )
        service.add_attribute_definition(
            self.pid, "verdict", "Verdict", type="enum", choices=["Pass", "Fail", "Blocked", "Not Run", "N/A"],
            module_types=["test_execution"],
        )
        service.add_attribute_definition(self.pid, "run_date", "Run Date", type="date", module_types=["test_execution"])
        service.add_attribute_definition(self.pid, "tester", "Tester", module_types=["test_execution"])
        service.add_attribute_definition(self.pid, "build", "Build", module_types=["test_execution"])
        service.add_attribute_definition(self.pid, "log", "Log", module_types=["test_execution"])
        service.add_attribute_definition(
            self.pid, "tc_result", "Test Result", type="enum", choices=["Not Run", "Passed", "Failed", "Blocked"],
            module_types=["test_cases"],
        )
        self.reqs = service.create_module(self.pid, "Reqs", "REQ").id
        self.tests = service.create_module(self.pid, "Tests", "TC", type="test_cases").id
        self.execs = service.create_module(self.pid, "Exec", "EX", type="test_execution").id
        self.verifies = service.create_link_configuration(self.pid, "Verifies", self.tests, self.reqs).id
        self.executes = service.create_link_configuration(self.pid, "Executes", self.execs, self.tests).id

    def add(self, module_id: str, text: str, parent_id: str | None = None, **values) -> str:
        return self.service.add_requirement(self.pid, module_id, {"object_text": text, **values}, parent_id=parent_id).id

    def verify(self, test_id: str, requirement_id: str) -> None:
        self.service.link_requirements(self.pid, self.verifies, test_id, requirement_id)

    def execute(self, record_id: str, test_id: str) -> None:
        self.service.link_requirements(self.pid, self.executes, record_id, test_id)

    def record(self, test_id: str | None, verdict: str | None, run_date: str | None = None, module_id: str | None = None, **values) -> str:
        attrs = {k: v for k, v in {"verdict": verdict, "run_date": run_date, **values}.items() if v is not None}
        record_id = self.add(module_id or self.execs, f"run of {test_id}", **attrs)
        if test_id is not None:
            if module_id is None or module_id == self.execs:
                self.execute(record_id, test_id)
        return record_id


# --- traceability matrix: chain validation --------------------------------------------


@pytest.mark.parametrize("config_ids", [[], None, "verifies", [1], ["verifies", None]])
def test_traceability_requires_a_list_of_configuration_ids(service, config_ids):
    lab = Lab(service)
    with pytest.raises(ValueError) as excinfo:
        service.traceability_matrix(lab.pid, config_ids)
    assert str(excinfo.value) == "Pick at least one link configuration"


def test_traceability_refuses_duplicate_configuration(service):
    lab = Lab(service)
    with pytest.raises(ValueError) as excinfo:
        service.traceability_matrix(lab.pid, [lab.verifies, lab.executes, lab.verifies])
    assert str(excinfo.value) == "A link configuration appears twice in the chain"


def test_traceability_refuses_chain_longer_than_limit(service):
    lab = Lab(service)
    assert MAX_TRACE_CHAIN == 10
    with pytest.raises(ValueError) as excinfo:
        service.traceability_matrix(lab.pid, [f"c{i}" for i in range(MAX_TRACE_CHAIN + 1)])
    assert str(excinfo.value) == "A traceability chain has at most 10 link configurations"


def test_traceability_unknown_configuration_is_not_found(service):
    lab = Lab(service)
    with pytest.raises(FileNotFoundError, match="Link configuration 'nope' not found in project 'lab'"):
        service.traceability_matrix(lab.pid, [lab.verifies, "nope"])
    with pytest.raises(FileNotFoundError):
        service.traceability_matrix("missing", [lab.verifies])


def test_traceability_refuses_a_step_that_does_not_continue_the_chain(service):
    lab = Lab(service)
    other = service.create_module(lab.pid, "Other", "OTH").id
    derives = service.create_link_configuration(lab.pid, "Derives", other, lab.reqs).id
    # verifies: Reqs -> Tests; executes: -> Exec; derives joins Reqs, not Exec.
    with pytest.raises(ValueError) as excinfo:
        service.traceability_matrix(lab.pid, [lab.verifies, lab.executes, derives])
    assert str(excinfo.value) == (
        "Link configuration 'Derives' does not continue the chain: it does not join module 'Exec'"
    )


def test_traceability_chain_starts_at_target_or_reversed_start(service):
    lab = Lab(service)
    forward = service.traceability_matrix(lab.pid, [lab.verifies, lab.executes])
    assert [m["id"] for m in forward["modules"]] == [lab.reqs, lab.tests, lab.execs]
    # The second configuration joins only the first one's source: start at the source.
    reversed_ = service.traceability_matrix(lab.pid, [lab.executes, lab.verifies])
    assert [m["id"] for m in reversed_["modules"]] == [lab.execs, lab.tests, lab.reqs]
    single = service.traceability_matrix(lab.pid, [lab.executes])
    assert [m["id"] for m in single["modules"]] == [lab.tests, lab.execs]


def test_traceability_refuses_cross_project_configuration(service):
    lab = Lab(service)
    other = service.create_project("Upstream")
    upstream = service.create_module(other.id, "Up", "UP").id
    service.set_project_references(lab.pid, [other.id])
    cross = service.create_link_configuration(lab.pid, "Refines", lab.reqs, upstream, target_project_id=other.id).id
    with pytest.raises(ValueError) as excinfo:
        service.traceability_matrix(lab.pid, [cross])
    assert str(excinfo.value) == (
        "Link configuration 'Refines' links to another project; cross-project configurations cannot be part "
        "of a traceability matrix yet"
    )


# --- traceability matrix: content -------------------------------------------------------


def _traced_lab(service):
    lab = Lab(service)
    heading = lab.add(lab.reqs, "Chapter", object_type="Heading")
    r1 = lab.add(lab.reqs, "First", parent_id=heading)
    r2 = lab.add(lab.reqs, "Second", parent_id=heading)
    r3 = lab.add(lab.reqs, "Third")
    t1, t2, t3 = (lab.add(lab.tests, f"Test {i}") for i in (1, 2, 3))
    lab.verify(t1, r1)
    lab.verify(t2, r1)
    lab.verify(t2, r2)
    lab.verify(t3, heading)  # links to excluded objects are ignored
    e1 = lab.record(t1, "Pass", "2026-01-01")
    e2 = lab.record(t2, "Fail", "2026-01-02")
    e3 = lab.record(None, "Pass")  # not linked to a test case
    return lab, dict(heading=heading, r1=r1, r2=r2, r3=r3, t1=t1, t2=t2, t3=t3, e1=e1, e2=e2, e3=e3)


def test_traceability_grid_coverage_orphans_trace_and_unverified(service):
    lab, ids = _traced_lab(service)
    report = service.traceability_matrix(lab.pid, [lab.verifies, lab.executes])

    reqs, tests, execs = report["modules"]
    assert [it["id"] for it in reqs["items"]] == [ids["r1"], ids["r2"], ids["r3"]]
    assert reqs["excluded"] == 1
    assert [it["level"] for it in reqs["items"]] == [1, 1, 0]
    assert reqs["type"] == "requirements" and tests["type"] == "test_cases" and execs["type"] == "test_execution"

    step1, step2 = report["steps"]
    assert step1["config"]["id"] == lab.verifies
    assert (step1["from_module_id"], step1["to_module_id"]) == (lab.reqs, lab.tests)
    assert step1["links"] == [[ids["r1"], ids["t1"]], [ids["r1"], ids["t2"]], [ids["r2"], ids["t2"]]]
    assert (step1["total"], step1["covered"], step1["covered_percent"]) == (3, 2, 66.7)
    assert step1["uncovered"] == [ids["r3"]]
    assert step1["orphans"] == [ids["t3"]]
    assert step2["links"] == [[ids["t1"], ids["e1"]], [ids["t2"], ids["e2"]]]
    assert step2["uncovered"] == [ids["t3"]]
    assert step2["orphans"] == [ids["e3"]]

    assert report["trace"] == [
        {"id": ids["r1"], "reached": [[ids["t1"], ids["t2"]], [ids["e1"], ids["e2"]]], "complete": True},
        {"id": ids["r2"], "reached": [[ids["t2"]], [ids["e2"]]], "complete": True},
        {"id": ids["r3"], "reached": [[], []], "complete": False},
    ]
    assert report["summary"] == {
        "total": 3, "covered": 2, "covered_percent": 66.7, "traced": 2, "traced_percent": 66.7,
    }
    assert report["unverified"] == [
        {"module_id": lab.reqs, "test_module_id": lab.tests, "total": 3, "ids": [ids["r3"]]},
    ]


@pytest.mark.parametrize("kind", sorted(REPORT_EXCLUDED_OBJECT_TYPES))
def test_report_excluded_object_types_compare_like_names(service, kind):
    project = service.create_project("P")
    service.add_attribute_definition(project.id, "kind", "Object-Type", type="string")
    reqs = service.create_module(project.id, "Reqs", "REQ").id
    tests = service.create_module(project.id, "Tests", "TC", type="test_cases").id
    config = service.create_link_configuration(project.id, "Verifies", tests, reqs).id
    service.add_requirement(project.id, reqs, {"object_text": "x", "kind": f"  {kind.upper()} "})
    kept = service.add_requirement(project.id, reqs, {"object_text": "y", "kind": "Requirement"}).id
    report = service.traceability_matrix(project.id, [config])
    assert [it["id"] for it in report["modules"][0]["items"]] == [kept]
    assert report["modules"][0]["excluded"] == 1


def test_traceability_item_text_is_plain_and_collapsed(service):
    lab = Lab(service)
    rid = lab.add(
        lab.reqs,
        '<div xmlns="http://www.w3.org/1999/xhtml"><p>Rich <strong>text</strong></p><ul><li>item</li></ul></div>',
    )
    report = service.traceability_matrix(lab.pid, [lab.verifies])
    text = report["modules"][0]["items"][0]["text"]
    assert report["modules"][0]["items"][0]["id"] == rid
    assert "<" not in text and "Rich text" in text and "item" in text
    assert "\n" not in text


def test_traceability_writes_nothing(service, tmp_path):
    lab, _ = _traced_lab(service)
    before = _snapshot(tmp_path)
    service.traceability_matrix(lab.pid, [lab.verifies, lab.executes])
    service.verification_report(lab.pid, lab.reqs)
    service.execution_report(lab.pid, lab.execs)
    service.verification_statuses(lab.pid, lab.reqs)
    service.export_traceability(lab.pid, [lab.verifies], "xlsx")
    service.export_verification_report(lab.pid, lab.reqs, "docx")
    service.export_execution_report(lab.pid, lab.execs, "xlsx")
    assert _snapshot(tmp_path) == before


# --- attribute roles ---------------------------------------------------------------------


def test_execution_report_finds_role_attributes_by_non_rbc_names(service):
    lab = Lab(service)
    report = service.execution_report(lab.pid, lab.execs)
    assert report["attributes"] == {
        "result": {"id": "verdict", "name": "Verdict"},
        "execution_date": {"id": "run_date", "name": "Run Date"},
        "executed_by": {"id": "tester", "name": "Tester"},
        "version": {"id": "build", "name": "Build"},
        "evidence": {"id": "log", "name": "Log"},
    }


def test_role_fallbacks_only_date_and_contains_words(service):
    project = service.create_project("P")
    for attr_id, name, type_ in (
        ("outcome_x", "Outcome", "enum"),
        ("when", "When", "date"),
        ("fw", "Firmware Version Label", "string"),
        ("proof", "Test Evidence Link", "string"),
        ("who", "Executor", "string"),
    ):
        kwargs = {"choices": ["OK", "NOK"]} if type_ == "enum" else {}
        service.add_attribute_definition(project.id, attr_id, name, type=type_, module_types=["test_execution"], **kwargs)
    execs = service.create_module(project.id, "Exec", "EX", type="test_execution").id
    report = service.execution_report(project.id, execs)
    assert {role: (a["id"] if a else None) for role, a in report["attributes"].items()} == {
        "result": "outcome_x", "execution_date": "when", "executed_by": "who", "version": "fw", "evidence": "proof",
    }


def test_two_unnamed_date_attributes_give_no_execution_date_role(service):
    project = service.create_project("P")
    service.add_attribute_definition(project.id, "a", "Start", type="date")
    service.add_attribute_definition(project.id, "b", "End", type="date")
    execs = service.create_module(project.id, "Exec", "EX", type="test_execution").id
    assert service.execution_report(project.id, execs)["attributes"]["execution_date"] is None


def test_role_ignores_attributes_not_applying_to_the_module_and_wrong_types(service):
    project = service.create_project("P")
    service.add_attribute_definition(project.id, "result", "Result", type="integer")
    service.add_attribute_definition(project.id, "verdict", "Verdict", type="string", module_types=["test_cases"])
    execs = service.create_module(project.id, "Exec", "EX", type="test_execution").id
    assert service.execution_report(project.id, execs)["attributes"]["result"] is None


# --- verification report --------------------------------------------------------------


def _verification_lab(service):
    lab = Lab(service)
    lab.add(lab.reqs, "Chapter", object_type="Heading")
    r = {i: lab.add(lab.reqs, f"R{i}") for i in range(1, 7)}
    t = {i: lab.add(lab.tests, f"T{i}") for i in range(1, 6)}
    service.update_requirement(lab.pid, lab.tests, t[3], {"tc_result": "Passed"})
    service.update_requirement(lab.pid, lab.tests, t[4], {"tc_result": "Passed"})
    lab.verify(t[1], r[1])
    lab.verify(t[1], r[2])
    lab.verify(t[2], r[2])
    lab.verify(t[3], r[3])
    lab.verify(t[4], r[4])
    lab.verify(t[5], r[5])
    lab.record(t[1], "Pass", "2026-02-01", tester="Ann", build="1.0")
    lab.record(t[2], "Fail", "2026-02-02")
    lab.record(t[4], "Not Run")  # planned only: no fallback to the test case's own result
    lab.record(t[5], "Blocked", "2026-02-03")
    return lab, r, t


def test_verification_report_statuses_totals_and_fallback(service):
    lab, r, t = _verification_lab(service)
    report = service.verification_report(lab.pid, lab.reqs)
    status = {row["id"]: row["status"] for row in report["requirements"]}
    assert status == {
        r[1]: "verified", r[2]: "failed", r[3]: "verified", r[4]: "not_run", r[5]: "not_run", r[6]: "not_covered",
    }
    assert report["totals"] == {
        "total": 6, "verified": 2, "failed": 1, "not_run": 2, "not_covered": 1, "excluded": 1,
    }
    assert report["percent"] == {"verified": 33.3, "failed": 16.7, "not_run": 33.3, "not_covered": 16.7}
    rows = {row["id"]: row for row in report["requirements"]}
    t1 = rows[r[1]]["tests"][0]
    assert (t1["id"], t1["result"], t1["source"], t1["executions"]) == (t[1], "passed", "execution", 1)
    assert t1["latest"]["executed_by"] == "Ann" and t1["latest"]["version"] == "1.0"
    assert "_key" not in t1["latest"] and "_order" not in t1
    fallback = rows[r[3]]["tests"][0]
    assert (fallback["result"], fallback["source"], fallback["executions"], fallback["latest"]) == (
        "passed", "test_case", 0, None,
    )
    planned = rows[r[4]]["tests"][0]
    assert (planned["result"], planned["source"], planned["executions"], planned["latest"]) == (
        "not_run", None, 1, None,
    )
    assert [x["id"] for x in rows[r[2]]["tests"]] == [t[1], t[2]]
    assert report["module"] == {"id": lab.reqs, "name": "Reqs", "type": "requirements"}
    assert report["test_modules"] == [{"id": lab.tests, "name": "Tests", "type": "test_cases"}]
    assert report["execution_modules"][0]["id"] == lab.execs
    assert report["execution_modules"][0]["test_module_id"] == lab.tests


@pytest.mark.parametrize(
    "results, expected",
    [
        (["Pass", "Pass"], "verified"),
        (["Pass", "Fail"], "failed"),
        (["Blocked", "Fail"], "failed"),
        (["Pass", "Blocked"], "not_run"),
        (["Pass", None], "not_run"),
        (["N/A"], "not_run"),
    ],
)
def test_verification_status_precedence(service, results, expected):
    lab = Lab(service)
    requirement = lab.add(lab.reqs, "R")
    for index, verdict in enumerate(results):
        test = lab.add(lab.tests, f"T{index}")
        lab.verify(test, requirement)
        if verdict is not None:
            lab.record(test, verdict, "2026-01-01")
    assert service.verification_statuses(lab.pid, lab.reqs)["statuses"] == {requirement: expected}


@pytest.mark.parametrize("own, expected", [("Failed", "failed"), ("Blocked", "not_run"), ("Not Run", "not_run"), (None, "not_run")])
def test_test_case_own_result_fallback(service, own, expected):
    lab = Lab(service)
    requirement = lab.add(lab.reqs, "R")
    values = {"tc_result": own} if own else {}
    test = lab.add(lab.tests, "T", **values)
    lab.verify(test, requirement)
    report = service.verification_report(lab.pid, lab.reqs)
    assert report["requirements"][0]["status"] == expected
    assert report["requirements"][0]["tests"][0]["source"] == ("test_case" if own in ("Failed", "Blocked") else None)


@pytest.mark.parametrize(
    "value, category",
    [("Pass", "passed"), ("PASSED", "passed"), (" ok ", "passed"), ("Success", "passed"),
     ("Fail", "failed"), ("NOK", "failed"), ("error", "failed"), ("Failure", "failed"),
     ("blocked", "blocked"), ("Not Run", None), ("Skipped", None), ("weird", None)],
)
def test_result_words_compare_like_names(service, value, category):
    project = service.create_project("P")
    service.add_attribute_definition(project.id, "result", "Result", type="string")
    execs = service.create_module(project.id, "Exec", "EX", type="test_execution").id
    service.add_requirement(project.id, execs, {"object_text": "x", "result": value})
    report = service.execution_report(project.id, execs)
    expected = category or "not_run"
    assert report["records"][0]["result"] == expected
    assert report["records"][0]["result_value"] == value


# --- latest execution ----------------------------------------------------------------------


def test_latest_execution_greatest_date_wins_over_tree_order(service):
    lab = Lab(service)
    requirement = lab.add(lab.reqs, "R")
    test = lab.add(lab.tests, "T")
    lab.verify(test, requirement)
    newer = lab.record(test, "Pass", "2026-03-02")
    lab.record(test, "Fail", "2026-03-01")  # later in tree order, earlier date
    lab.record(test, "Not Run", "2026-12-31")  # planned: never the latest
    report = service.verification_report(lab.pid, lab.reqs)
    latest = report["requirements"][0]["tests"][0]["latest"]
    assert latest["id"] == newer
    assert report["requirements"][0]["status"] == "verified"


def test_latest_execution_record_without_date_sorts_before_dated(service):
    lab = Lab(service)
    requirement = lab.add(lab.reqs, "R")
    test = lab.add(lab.tests, "T")
    lab.verify(test, requirement)
    dated = lab.record(test, "Fail", "2020-01-01")
    lab.record(test, "Pass")  # later in tree order but undated
    assert service.verification_report(lab.pid, lab.reqs)["requirements"][0]["tests"][0]["latest"]["id"] == dated


def test_latest_execution_invalid_stored_date_counts_as_no_date(service, tmp_path):
    lab = Lab(service)
    requirement = lab.add(lab.reqs, "R")
    test = lab.add(lab.tests, "T")
    lab.verify(test, requirement)
    dated = lab.record(test, "Fail", "2020-01-01")
    broken = lab.record(test, "Pass", "2020-01-02")
    path = tmp_path / lab.pid / "modules" / lab.execs / "requirements" / f"{broken}.json"
    data = json.loads(path.read_text())
    data["attributes"]["run_date"] = "9999-99-99"
    path.write_text(json.dumps(data, indent=2) + "\n")
    test_row = service.verification_report(lab.pid, lab.reqs)["requirements"][0]["tests"][0]
    assert test_row["latest"]["id"] == dated
    records = {r["id"]: r for r in service.execution_report(lab.pid, lab.execs)["records"]}
    assert records[broken]["execution_date"] is None


def test_latest_execution_same_date_later_tree_position_wins(service):
    lab = Lab(service)
    requirement = lab.add(lab.reqs, "R")
    test = lab.add(lab.tests, "T")
    lab.verify(test, requirement)
    lab.record(test, "Pass", "2026-01-01")
    last = lab.record(test, "Fail", "2026-01-01")
    report = service.verification_report(lab.pid, lab.reqs)
    assert report["requirements"][0]["tests"][0]["latest"]["id"] == last
    assert report["requirements"][0]["status"] == "failed"
    # Reordering puts the passing record last: it becomes the latest.
    service.reorder_requirement(lab.pid, lab.execs, last, "up")
    report = service.verification_report(lab.pid, lab.reqs)
    assert report["requirements"][0]["status"] == "verified"


def test_latest_execution_same_date_later_module_in_module_order_wins(service):
    lab = Lab(service)
    second = service.create_module(lab.pid, "Exec 2", "EXB", type="test_execution").id
    executes2 = service.create_link_configuration(lab.pid, "Executes 2", second, lab.tests).id
    requirement = lab.add(lab.reqs, "R")
    test = lab.add(lab.tests, "T")
    lab.verify(test, requirement)
    in_second = lab.add(second, "run", verdict="Fail", run_date="2026-01-01")
    service.link_requirements(lab.pid, executes2, in_second, test)
    lab.record(test, "Pass", "2026-01-01")
    assert service.verification_statuses(lab.pid, lab.reqs)["statuses"] == {requirement: "failed"}
    service.reorder_module(lab.pid, second, "up")  # now before Exec
    assert [m.id for m in service.list_modules(lab.pid)].index(second) < [
        m.id for m in service.list_modules(lab.pid)
    ].index(lab.execs)
    assert service.verification_statuses(lab.pid, lab.reqs)["statuses"] == {requirement: "verified"}


# --- execution report --------------------------------------------------------------------


def _execution_lab(service):
    lab = Lab(service)
    r1, r2 = lab.add(lab.reqs, "R1"), lab.add(lab.reqs, "R2")
    lab.add(lab.tests, "Section", object_type="Heading")
    t = {i: lab.add(lab.tests, f"T{i}") for i in range(1, 6)}
    lab.verify(t[2], r1)
    lab.verify(t[2], r2)
    lab.verify(t[4], r2)
    e = {
        1: lab.record(t[1], "Pass", "2026-01-02", tester="Ann", build="1.0", log="log1"),
        2: lab.record(t[2], "Fail", "2026-01-05", tester="Bob", build="1.1", log="log2"),
        3: lab.record(t[3], "Not Run", "2026-01-09"),
        4: lab.record(t[4], "Blocked", None, tester="Ann"),
        5: lab.record(None, "Pass", "2026-01-01", tester="Cy"),
    }
    return lab, (r1, r2), t, e


def test_execution_report_counts_problems_unlinked_versions_executors_and_dates(service):
    lab, (r1, r2), t, e = _execution_lab(service)
    report = service.execution_report(lab.pid, lab.execs)
    assert report["module"] == {"id": lab.execs, "name": "Exec", "type": "test_execution"}
    assert report["test_modules"] == [{"id": lab.tests, "name": "Tests", "type": "test_cases"}]
    assert report["record_counts"] == {
        "total": 5, "passed": 2, "failed": 1, "blocked": 1, "not_run": 1, "not_applicable": 0, "excluded": 0,
    }
    status = {tc["id"]: tc["status"] for tc in report["test_cases"]}
    assert status == {t[1]: "passed", t[2]: "failed", t[3]: "not_run", t[4]: "blocked", t[5]: "not_planned"}
    assert report["test_counts"] == {
        "total": 5, "planned": 4, "executed": 3, "executed_percent": 75.0, "pass_percent": 33.3,
        "passed": 1, "failed": 1, "blocked": 1, "not_run": 1, "not_planned": 1,
    }
    problems = {p["id"]: p for p in report["problems"]}
    assert list(problems) == [t[2], t[4]]
    assert [x["id"] for x in problems[t[2]]["requirements"]] == [r1, r2]
    assert problems[t[2]]["requirements"][0] == {"id": r1, "module_id": lab.reqs, "text": "R1"}
    assert [x["id"] for x in problems[t[4]]["requirements"]] == [r2]
    assert report["unlinked_records"] == [e[5]]
    assert report["versions"] == ["1.0", "1.1"]
    assert report["executors"] == ["Ann", "Bob", "Cy"]
    # Executed records only: the Not Run record's 2026-01-09 is outside the range.
    assert report["date_range"] == ["2026-01-01", "2026-01-05"]
    records = {r["id"]: r for r in report["records"]}
    assert records[e[1]]["test_case_ids"] == [t[1]]
    assert records[e[1]]["evidence"] == "log1"
    assert records[e[5]]["test_case_ids"] == []
    assert all("_key" not in r for r in report["records"])


def test_execution_report_empty_module(service):
    lab = Lab(service)
    report = service.execution_report(lab.pid, lab.execs)
    assert report["test_counts"]["executed_percent"] == 0.0
    assert report["test_counts"]["pass_percent"] == 0.0
    assert report["date_range"] is None
    assert report["versions"] == [] and report["executors"] == [] and report["problems"] == []


def test_execution_report_only_counts_its_own_records(service):
    lab = Lab(service)
    second = service.create_module(lab.pid, "Exec 2", "EXB", type="test_execution").id
    executes2 = service.create_link_configuration(lab.pid, "Executes 2", second, lab.tests).id
    test = lab.add(lab.tests, "T")
    lab.record(test, "Fail", "2026-05-01")
    other = lab.add(second, "run", verdict="Pass", run_date="2026-06-01")
    service.link_requirements(lab.pid, executes2, other, test)
    assert service.execution_report(lab.pid, lab.execs)["test_cases"][0]["status"] == "failed"
    assert service.execution_report(lab.pid, second)["test_cases"][0]["status"] == "passed"


# --- wrong module types -------------------------------------------------------------------


def test_reports_refuse_wrong_module_types(service):
    lab = Lab(service)
    with pytest.raises(ValueError) as excinfo:
        service.verification_report(lab.pid, lab.tests)
    assert str(excinfo.value) == (
        "Module 'Tests' is a test_cases module; the verification report is for a requirements module"
    )
    with pytest.raises(ValueError) as excinfo:
        service.verification_statuses(lab.pid, lab.execs)
    assert "the verification report is for a requirements module" in str(excinfo.value)
    with pytest.raises(ValueError) as excinfo:
        service.execution_report(lab.pid, lab.reqs)
    assert str(excinfo.value) == (
        "Module 'Reqs' is a requirements module; the test execution report is for a test_execution module"
    )
    with pytest.raises(FileNotFoundError):
        service.verification_report(lab.pid, "missing")
    with pytest.raises(FileNotFoundError):
        service.execution_report("missing", lab.execs)


# --- RBC template smoke --------------------------------------------------------------


def test_rbc_template_reports_leave_out_skeleton_headings(service):
    project = service.create_project("R", template_id=RBC)
    report = service.verification_report(project.id, "syrs")
    assert report["totals"]["total"] == 0 or all(r["status"] == "not_covered" for r in report["requirements"])
    assert report["test_modules"] == [{"id": "systest", "name": "SysTest", "type": "test_cases"}]
    execution = service.execution_report(project.id, "swtestexec")
    assert execution["attributes"] == {
        "result": {"id": "result", "name": "Result"},
        "execution_date": {"id": "execution_date", "name": "Execution Date"},
        "executed_by": {"id": "executed_by", "name": "Executed By"},
        "version": {"id": "version_under_test", "name": "Build/SW Version Under Test"},
        "evidence": {"id": "evidence_ref", "name": "Evidence/Log Ref"},
    }
    chain = ["satisfies-strs", "satisfies-syrs", "verifies-swrs", "executes-swtest"]
    matrix = service.traceability_matrix(project.id, chain)
    assert [m["id"] for m in matrix["modules"]] == ["strs", "syrs", "swrs", "swtest", "swtestexec"]
    matrix = service.traceability_matrix(project.id, ["verifies-syrs", "satisfies-strs"])
    assert [m["id"] for m in matrix["modules"]] == ["systest", "syrs", "strs"]


# --- data source ----------------------------------------------------------------------


def test_report_data_source_without_history_and_with_lock(service):
    lab = Lab(service)
    report = service.verification_report(lab.pid, lab.reqs)
    assert report["data_source"] == {
        "history": "off", "commit": None, "pending_changes": 0, "baselines": [], "latest_baseline": None,
        "locked": None,
    }
    service.lock_project(lab.pid, "Release 1")
    source = service.execution_report(lab.pid, lab.execs)["data_source"]
    assert source["locked"]["reason"] == "Release 1"
    assert "data_source" not in service.traceability_matrix(lab.pid, [lab.verifies])
    # Reports still work on a locked project.
    assert service.verification_statuses(lab.pid, lab.reqs)["module_id"] == lab.reqs


# --- REST / WebSocket ------------------------------------------------------------------


def _api_lab(client, tmp_path):
    service = ProjectService(ProjectStore(tmp_path))
    lab, ids = _traced_lab(service)
    return lab, ids


def test_reports_rest_ws_parity_and_no_broadcast(client, tmp_path):
    lab, ids = _api_lab(client, tmp_path)
    pid = lab.pid
    before = _snapshot(tmp_path / pid)
    with client.websocket_connect("/ws") as ws, client.websocket_connect("/ws") as watcher:
        rest = client.get(f"/projects/{pid}/reports/traceability?config={lab.verifies}&config={lab.executes}")
        assert rest.status_code == 200, rest.text
        ws_trace = _rpc(ws, "report.traceability", {"project_id": pid, "config_ids": [lab.verifies, lab.executes]})
        assert ws_trace["result"] == rest.json()

        rest = client.get(f"/projects/{pid}/modules/{lab.reqs}/reports/verification")
        assert rest.status_code == 200, rest.text
        assert _rpc(ws, "report.verification", {"project_id": pid, "module_id": lab.reqs}, 2)["result"] == rest.json()

        rest = client.get(f"/projects/{pid}/modules/{lab.execs}/reports/execution")
        assert rest.status_code == 200, rest.text
        assert _rpc(ws, "report.execution", {"project_id": pid, "module_id": lab.execs}, 3)["result"] == rest.json()

        rest = client.get(f"/projects/{pid}/modules/{lab.reqs}/verification-status")
        assert rest.status_code == 200, rest.text
        assert rest.json() == {
            "project_id": pid, "module_id": lab.reqs,
            "statuses": {ids["r1"]: "failed", ids["r2"]: "failed", ids["r3"]: "not_covered"},
        }
        assert _rpc(ws, "report.verification_status", {"project_id": pid, "module_id": lab.reqs}, 4)["result"] == rest.json()

        for fmt, url in (
            ("csv", f"/projects/{pid}/reports/traceability/export/csv?config={lab.verifies}"),
            ("xlsx", f"/projects/{pid}/modules/{lab.reqs}/reports/verification/export/xlsx"),
            ("docx", f"/projects/{pid}/modules/{lab.execs}/reports/execution/export/docx"),
        ):
            assert client.get(url).status_code == 200, (fmt, url)

        # Nothing was broadcast: the watcher's next event is a later, distinguishable mutation.
        client.post("/projects", json={"name": "Other"})
        assert json.loads(watcher.receive_text()) == {"event": "changed", "project_id": "other"}
        assert json.loads(ws.receive_text()) == {"event": "changed", "project_id": "other"}
    assert _snapshot(tmp_path / pid) == before


@pytest.mark.parametrize(
    "method, params, url, status, code",
    [
        ("report.traceability", {"config_ids": []}, "/reports/traceability", 400, -32602),
        ("report.traceability", {"config_ids": ["nope"]}, "/reports/traceability?config=nope", 404, -32001),
        ("report.verification", {"module_id": "tests"}, "/modules/tests/reports/verification", 400, -32602),
        ("report.verification", {"module_id": "missing"}, "/modules/missing/reports/verification", 404, -32001),
        ("report.execution", {"module_id": "reqs"}, "/modules/reqs/reports/execution", 400, -32602),
        ("report.execution", {"module_id": "missing"}, "/modules/missing/reports/execution", 404, -32001),
        ("report.verification_status", {"module_id": "exec"}, "/modules/exec/verification-status", 400, -32602),
        ("report.verification_status", {"module_id": "tests"}, "/modules/tests/verification-status", 400, -32602),
        ("report.verification_status", {"module_id": "missing"}, "/modules/missing/verification-status", 404, -32001),
    ],
)
def test_report_errors_rest_and_ws(client, tmp_path, method, params, url, status, code):
    lab, _ = _api_lab(client, tmp_path)
    resp = client.get(f"/projects/{lab.pid}{url}")
    assert resp.status_code == status, resp.text
    with client.websocket_connect("/ws") as ws:
        response = _rpc(ws, method, {"project_id": lab.pid, **params})
    assert response["error"]["code"] == code
    assert response["error"]["message"] == resp.json()["detail"]


def test_report_on_missing_project_is_404(client):
    assert client.get("/projects/missing/reports/traceability?config=x").status_code == 404
    assert client.get("/projects/missing/modules/m/reports/verification").status_code == 404
    assert client.get("/projects/missing/modules/m/verification-status").status_code == 404


# --- exports ------------------------------------------------------------------------------


def test_report_export_headers_and_content(client, tmp_path):
    lab, ids = _api_lab(client, tmp_path)
    pid = lab.pid
    resp = client.get(f"/projects/{pid}/reports/traceability/export/csv?config={lab.verifies}&config={lab.executes}")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert resp.headers["content-disposition"] == f'attachment; filename="{pid}-traceability.csv"'
    rows = list(csv.reader(io.StringIO(resp.text)))
    assert rows[0] == ["Reqs Id", "Reqs Text", "Tests", "Exec", "Covered", "Traced to Exec"]
    assert rows[1] == [ids["r1"], "First", f"{ids['t1']}; {ids['t2']}", f"{ids['e1']}; {ids['e2']}", "yes", "yes"]
    assert rows[3] == [ids["r3"], "Third", "", "", "no", "no"]

    resp = client.get(f"/projects/{pid}/reports/traceability/export/xlsx?config={lab.verifies}&config={lab.executes}")
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.headers["content-disposition"] == f'attachment; filename="{pid}-traceability.xlsx"'
    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Summary", "Trace", "Matrix 1", "Matrix 2", "Gaps"]
    gaps = [[c.value for c in row] for row in workbook["Gaps"].iter_rows()]
    assert ["Uncovered", 1, "Reqs", ids["r3"], "Third"] in gaps
    assert ["Orphan", 1, "Tests", ids["t3"], "Test 3"] in gaps
    assert ["Unverified", None, "Reqs", ids["r3"], "Third"] in gaps or ["Unverified", "", "Reqs", ids["r3"], "Third"] in gaps

    resp = client.get(f"/projects/{pid}/modules/{lab.reqs}/reports/verification/export/xlsx")
    assert resp.headers["content-disposition"] == f'attachment; filename="{pid}-{lab.reqs}-verification.xlsx"'
    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Summary", "Requirements"]
    summary = [[c.value for c in row] for row in workbook["Summary"].iter_rows()]
    assert ["Data", "Current data (no git history)"] in [row[:2] for row in summary]

    resp = client.get(f"/projects/{pid}/modules/{lab.reqs}/reports/verification/export/docx")
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    document = Document(io.BytesIO(resp.content))
    assert document.paragraphs[0].text == "Verification report: Reqs"
    assert any(p.text == "Data: Current data (no git history)" for p in document.paragraphs)

    resp = client.get(f"/projects/{pid}/modules/{lab.execs}/reports/execution/export/xlsx")
    assert resp.headers["content-disposition"] == f'attachment; filename="{pid}-{lab.execs}-execution.xlsx"'
    workbook = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert workbook.sheetnames == ["Summary", "Test cases", "Failed and blocked", "Records"]
    resp = client.get(f"/projects/{pid}/modules/{lab.execs}/reports/execution/export/docx")
    assert Document(io.BytesIO(resp.content)).paragraphs[0].text == "Test execution report: Exec"


@pytest.mark.parametrize(
    "url, message",
    [
        ("/reports/traceability/export/pdf?config=verifies", "Unknown traceability export format 'pdf'; expected csv or xlsx"),
        ("/reports/traceability/export/docx?config=verifies", "Unknown traceability export format 'docx'; expected csv or xlsx"),
        ("/modules/reqs/reports/verification/export/csv", "Unknown report export format 'csv'; expected docx or xlsx"),
        ("/modules/exec/reports/execution/export/pdf", "Unknown report export format 'pdf'; expected docx or xlsx"),
    ],
)
def test_report_export_wrong_format_is_400(client, tmp_path, url, message):
    lab, _ = _api_lab(client, tmp_path)
    resp = client.get(f"/projects/{lab.pid}{url}")
    assert resp.status_code == 400
    assert resp.json()["detail"] == message


def test_report_exports_wrong_module_type_is_400(client, tmp_path):
    lab, _ = _api_lab(client, tmp_path)
    assert client.get(f"/projects/{lab.pid}/modules/{lab.execs}/reports/verification/export/xlsx").status_code == 400
    assert client.get(f"/projects/{lab.pid}/modules/{lab.reqs}/reports/execution/export/docx").status_code == 400


@pytest.mark.parametrize("text", ["=SUM(A1:A2)", "+1+1", "-2", "@cmd", "=HYPERLINK(\"http://x\")"])
def test_traceability_csv_escapes_formula_prefixes(service, text):
    lab = Lab(service)
    rid = lab.add(lab.reqs, text)
    out = service.export_traceability(lab.pid, [lab.verifies], "csv")
    rows = list(csv.reader(io.StringIO(out)))
    assert rows[1] == [rid, "'" + text, "", "no"]
    workbook = openpyxl.load_workbook(io.BytesIO(service.export_traceability(lab.pid, [lab.verifies], "xlsx")))
    trace = [[c for c in row] for row in workbook["Trace"].iter_rows(min_row=2)]
    cell = trace[0][1]
    assert cell.value == text
    assert cell.data_type == "s"


def test_traceability_csv_keeps_safe_text_and_special_characters(service):
    lab = Lab(service)
    rid = lab.add(lab.reqs, 'Quote " comma, ünïcödé; semi')
    rows = list(csv.reader(io.StringIO(service.export_traceability(lab.pid, [lab.verifies], "csv"))))
    assert rows[1][:2] == [rid, 'Quote " comma, ünïcödé; semi']
